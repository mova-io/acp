"""Approval → verified bytes → durable artifact → Release eligibility.

Storage and the full assessment engine are local doubles. The PPTX writer and
first-party missing-title detector are real; these prove title presence, not the
semantic quality of an automatically generated title or full WCAG conformance.
Storage failure must leave approvals retryable and successful writes must refresh provenance.
"""
from __future__ import annotations

import hashlib
import io
import sys

import pytest
from hitl_viewed import approve_bound

SID = "approved-release-fixture"
FILE = "review.pptx"
TITLE = "Quarterly accessibility review"


class MemoryBlob:
    def __init__(self, data):
        self.data = data
        self.uploads = 0
        self.fail_upload = False

    def download_remediated(self, *_args):
        return self.data

    def upload_remediated(self, _owner, _sid, _file, data, _mime):
        if self.fail_upload:
            raise OSError("fixture: storage unavailable")
        self.data = data
        self.uploads += 1
        return "https://fixture.invalid/corrected-v2"

    # The approved writer publishes to a digest-scoped immutable object and moves the pointer
    # at commit (handlers._apply_approved_values); this fake serves whatever was stored last.
    def upload_immutable_retry(self, owner, sid, f, data, mime):
        return self.upload_remediated(owner, sid, f, data, mime)


@pytest.fixture
def progression(isolated_store, monkeypatch, tmp_path):
    import core
    import handlers
    from office_structure import pptx_checks
    from pptx import Presentation
    from proposals import Verification

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5])
    output = io.BytesIO()
    prs.save(output)
    original = output.getvalue()
    blob = MemoryBlob(original)
    store = isolated_store
    store.init_scan_run(SID, "local", 1, "2026-09-08T00:00:00Z", "rubric", "hash")
    store.save_file_result(SID, {
        "file": FILE, "engine": "office", "status": "pass", "score": 60,
        "compliant": 0, "skipped_rules": 0,
        "issues": [{"ruleId": "PPTX_TITLE_EMPTY", "wcag": "2.4.6 Headings and Labels",
                    "severity": "MODERATE", "location": "Slide 1"}],
    }, "2026-09-08T00:00:00Z")
    store.record_remediation(SID, FILE, drive_write_url="https://fixture.invalid/drive-copy", blob_url="https://fixture.invalid/corrected-v1",
                             corrected_sha256=hashlib.sha256(original).hexdigest(),
                             corrected_bytes=len(original))
    execution = store.enqueue_stage_batch(
        SID, "remediate", "remediate_file", [{"scan_id": SID, "file": FILE}],
        snapshot_id=SID, request_fingerprint="fixture-request")
    blob.batch_id = execution["batch_id"]
    store.seed_finding_dispositions(SID, blob.batch_id)
    item = store.enqueue_proposals(SID, FILE, "2.4.6", [{
        "locator": "slide 1", "before": "", "proposed_value": "",
        "rationale": "Human supplied title", "source": "reviewer",
    }], rule_name="Headings and Labels")
    approve_bound(store, item, [TITLE])

    def verify(data, filename, **kwargs):
        path = tmp_path / filename
        path.write_bytes(data)
        missing = any(f["ruleId"] == "PPTX_TITLE_EMPTY" for f in pptx_checks(path))
        return Verification(True, {"2.4.6"} if missing else set())

    assert verify(original, FILE).residual == {"2.4.6"}
    monkeypatch.setattr(core, "store", store)
    monkeypatch.setitem(sys.modules, "blob", blob)
    monkeypatch.setattr(handlers, "_verify_residual", verify)
    return store, blob, item, handlers, verify


def run(fixture):
    fixture[3]._apply_approved_values({"scan_id": SID, "file": FILE}, {})


def eligible(store):
    row = store.get_file_record(SID, FILE)
    return bool(row["compliant"] and row["remediated_at"])


def test_approval_alone_does_not_enable_release(progression):
    store, blob, item, _, verify = progression
    assert store.count_unapplied_approved_values(SID, FILE) == 1
    assert not store.mark_file_compliant_if_reviewed(SID, FILE)
    assert not eligible(store)
    assert not verify(blob.data, FILE).cleared({"2.4.6"})


def test_supported_approved_fix_reaches_release_after_write_and_verification(progression):
    from pptx import Presentation
    store, blob, item, _, verify = progression
    run(progression)
    assert blob.uploads == 1
    assert Presentation(io.BytesIO(blob.data)).slides[0].shapes.title.text == TITLE
    assert verify(blob.data, FILE).cleared({"2.4.6"})
    assert store.get_hitl_item(item)["applied"]
    assert store.count_unapplied_approved_values(SID, FILE) == 0
    assert eligible(store)
    rows = store.list_finding_dispositions(SID, blob.batch_id)
    assert len(rows) == 1 and rows[0]["disposition"] == "resolved_verified"
    run(progression)
    assert blob.uploads == 1  # successful retry is idempotent


@pytest.mark.parametrize("result", ["still_failing", "engine_unavailable"])
def test_failed_verification_preserves_pending_work(progression, monkeypatch, result):
    from proposals import Verification
    store, blob, item, handlers, _ = progression
    original = blob.data
    verification = (Verification(True, {"2.4.6"}) if result == "still_failing"
                    else Verification(False, (), "fixture engine unavailable"))
    monkeypatch.setattr(handlers, "_verify_residual", lambda *_, **kwargs: verification)
    run(progression)
    assert blob.data == original and blob.uploads == 0
    assert not store.get_hitl_item(item)["applied"]
    assert store.count_unapplied_approved_values(SID, FILE) == 1
    assert not eligible(store)


def test_unrelated_pending_review_legitimately_blocks_release(progression):
    store, blob, _, _, _ = progression
    store.enqueue_proposals(SID, FILE, "1.3.3", [{
        "locator": "slide 1", "before": "the red item", "proposed_value": "",
        "rationale": "Needs human meaning", "source": "reviewer",
    }], rule_name="Sensory Characteristics")
    run(progression)
    assert blob.uploads == 1
    assert store.count_unapplied_approved_values(SID, FILE) == 0
    assert not eligible(store)


def test_upload_failure_keeps_approval_retryable(progression):
    store, blob, item, _, _ = progression
    original = blob.data
    blob.fail_upload = True
    with pytest.raises(OSError, match="storage unavailable"):
        run(progression)
    assert blob.data == original and not eligible(store)
    applied_after_failure = store.get_hitl_item(item)["applied"]
    pending_after_failure = store.count_unapplied_approved_values(SID, FILE)
    disposition_after_failure = store.list_finding_dispositions(SID, blob.batch_id)[0]["disposition"]
    blob.fail_upload = False
    run(progression)
    assert (not applied_after_failure and pending_after_failure == 1
            and disposition_after_failure != "resolved_verified"
            and blob.uploads == 1 and eligible(store)), {
                "applied_after_failure": applied_after_failure,
                "pending_after_failure": pending_after_failure,
                "disposition_after_failure": disposition_after_failure,
                "retry_uploads": blob.uploads, "release_eligible": eligible(store),
            }


def test_approved_artifact_provenance_matches_durable_bytes(progression):
    store, blob, _, _, _ = progression
    run(progression)
    with store._db.cursor() as cur:
        store._db.execute(cur, "SELECT * FROM file_records WHERE scan_id=%s AND file=%s", (SID, FILE))
        record = store._db.fetchone(cur)
    assert record["corrected_sha256"] == hashlib.sha256(blob.data).hexdigest()
    assert record["corrected_bytes"] == len(blob.data)
    assert record["blob_url"] == "https://fixture.invalid/corrected-v2"
    assert record["drive_write_url"] == "https://fixture.invalid/drive-copy"


def test_storage_without_a_durable_url_cannot_credit_approval(progression, monkeypatch):
    store, blob, item, _, _ = progression
    monkeypatch.setattr(blob, "upload_remediated", lambda *_: None)
    with pytest.raises(RuntimeError, match="durable storage is unavailable"):
        run(progression)
    assert not store.get_hitl_item(item)["applied"]
    assert store.count_unapplied_approved_values(SID, FILE) == 1
    assert store.list_finding_dispositions(SID, blob.batch_id)[0]["disposition"] != "resolved_verified"
    assert not eligible(store)


@pytest.mark.parametrize("method", ["record_remediation", "mark_row_applied",
                                     "mark_file_compliant_if_reviewed"])
def test_database_failure_rolls_back_credit_and_retry_completes(progression, monkeypatch, method):
    store, blob, item, _, verify = progression
    original_method = getattr(store, method)
    previous_record = store.get_file_record(SID, FILE)

    def fail_after_write(*args, **kwargs):
        original_method(*args, **kwargs)
        raise RuntimeError("fixture database commit failure")

    with monkeypatch.context() as patch:
        patch.setattr(store, method, fail_after_write)
        with pytest.raises(RuntimeError, match="database commit failure"):
            run(progression)
    assert blob.uploads == 1 and verify(blob.data, FILE).cleared({"2.4.6"})
    assert not store.get_hitl_item(item)["applied"]
    assert store.count_unapplied_approved_values(SID, FILE) == 1
    assert store.list_finding_dispositions(SID, blob.batch_id)[0]["disposition"] != "resolved_verified"
    assert store.get_file_record(SID, FILE)["remediated_at"] == previous_record["remediated_at"]
    assert not eligible(store)
    run(progression)
    assert blob.uploads == 2 and eligible(store)
    assert store.get_hitl_item(item)["applied"]
