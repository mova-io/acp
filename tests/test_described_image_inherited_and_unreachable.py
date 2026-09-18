"""ADR 0055's FIRST-listed motivating case, which the lane did not actually serve.

The ADR opens its argument with an image "referenced by a layout or master … The reviewer sees a
card they cannot action." #1742 shipped the lane and #1764 extended it to docx and xlsx, and that
exact case still wedged the file — permanently, and silently:

  * `ocr._ooxml_images` walks the ZIP NAMELIST, so a picture on a slideLayout or slideMaster
    raises 1.4.5 and gets an `image N` review card like any other.
  * `formats.office.images.ALT_TARGETS` listed only `ppt/slides/slideN.xml`, so the SAME image was
    invisible to the 1.1.1 detector and unreachable by `apply_alt` — the first of the two failure
    directions that module's own docstring warns about.
  * So `resolve_media_locators` returned {}, nothing was written, and
    `_apply_one_value_kind` returned at `if not applied:` BEFORE logging anything. No
    `apply_outcome`, so no card; `approved`, so not in the pending inbox either. The reviewer got
    a file that never publishes and no explanation anywhere.

This module proves both halves of the fix: the inherited parts are now REACHED, and an image that
still cannot be reached is now VISIBLE rather than silent.

WHY THE FIXTURES USE ZIP SURGERY. python-pptx has no `LayoutShapes.add_picture` — there is no API
for putting a picture on a layout or a master — so the deck is built normally and the `<p:pic>`
element plus its relationship are MOVED onto the inherited part. Only those two entries change;
every other part is copied verbatim, so this is the real package shape rather than a mock of it.

NO PROVES_LANES DECLARATION, and the filename is part of that: `('pptx','1.1.1')` is claimed by
tests/test_remediation_verified_pptx_alt.py. This is a new PATH into that lane, not a new
capability, so it stays outside the `test_remediation_verified_*` namespace that
tests/test_capability_assisted_contract.py globs — the same call #1742 and #1764 made.
"""
from __future__ import annotations

import functools
import glob
import io
import re
import sys
import tempfile
import zipfile
from pathlib import Path

import pytest
from hitl_viewed import viewed_fields

ACP = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ACP / "api"))

pytest.importorskip("pptx")
pytest.importorskip("PIL")

LINES = ("Benefits at a glance",
         "Medical dental and vision cover",
         "Enrollment closes on Friday")
BODY = "Body copy that must survive the write."
DESC = ("A benefits summary card in the brand's display face, listing medical, dental and vision "
        "cover, and noting that enrollment closes on Friday.")


def _ocr_ready() -> bool:
    import ocr
    return ocr.is_available()


needs_ocr = pytest.mark.skipif(not _ocr_ready(),
                               reason="tesseract/pytesseract unavailable — 1.4.5 cannot fire")


def _font():
    from PIL import ImageFont
    for pat in ("/usr/share/fonts/**/DejaVuSans.ttf", "/usr/share/fonts/**/*.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf"):
        hits = glob.glob(pat, recursive=True)
        if hits:
            return ImageFont.truetype(hits[0], 34)
    return ImageFont.load_default()


@functools.lru_cache(maxsize=None)
def _png() -> Path:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (760, 260), "white")
    d = ImageDraw.Draw(im)
    font = _font()
    for i, line in enumerate(LINES):
        d.text((24, 24 + i * 70), line, fill="black", font=font)
    p = Path(tempfile.mkdtemp()) / "text.png"
    im.save(p)
    return p


@functools.lru_cache(maxsize=None)
def _plain_deck() -> bytes:
    """A deck with the picture on the SLIDE, as python-pptx builds it."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Open enrolment"
    s.shapes.add_picture(str(_png()), Inches(0.5), Inches(2), Inches(4), Inches(1.4))
    box = s.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
    box.text_frame.text = BODY
    out = Path(tempfile.mkdtemp()) / "src.pptx"
    prs.save(out)
    return out.read_bytes()


def _move_picture_to(host_part: str) -> bytes:
    """The plain deck with its picture MOVED off the slide onto `host_part`.

    The picture element and its image relationship are both relocated, so the media part is
    referenced by the inherited part and by nothing else — which is exactly the package a real
    deck has when a logo or a banner is placed on the layout rather than on each slide.
    """
    with zipfile.ZipFile(io.BytesIO(_plain_deck())) as z:
        entries = {n: z.read(n) for n in z.namelist()}

    slide = entries["ppt/slides/slide1.xml"].decode()
    rels = entries["ppt/slides/_rels/slide1.xml.rels"].decode()
    pic = re.search(r"<p:pic\b.*?</p:pic>", slide, re.S).group(0)
    img_rel = re.search(r'<Relationship[^>]*Type="[^"]*/image"[^>]*/>', rels).group(0)
    old_rid = re.search(r'Id="([^"]+)"', img_rel).group(1)
    new_rid = "rIdInherited"

    entries["ppt/slides/slide1.xml"] = slide.replace(pic, "").encode()
    entries["ppt/slides/_rels/slide1.xml.rels"] = rels.replace(img_rel, "").encode()

    host = entries[host_part].decode()
    entries[host_part] = host.replace(
        "</p:spTree>",
        pic.replace(f'r:embed="{old_rid}"', f'r:embed="{new_rid}"') + "</p:spTree>").encode()

    hdir, hname = host_part.rsplit("/", 1)
    hrels_path = f"{hdir}/_rels/{hname}.rels"
    hrels = entries[hrels_path].decode()
    entries[hrels_path] = hrels.replace(
        "</Relationships>",
        f'<Relationship Id="{new_rid}" Type="http://schemas.openxmlformats.org/officeDocument/'
        f'2006/relationships/image" Target="../media/image1.png"/></Relationships>').encode()

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, b in entries.items():
            zout.writestr(n, b)
    return buf.getvalue()


def _layout_of_slide1() -> str:
    with zipfile.ZipFile(io.BytesIO(_plain_deck())) as z:
        rels = z.read("ppt/slides/_rels/slide1.xml.rels").decode()
    return "ppt/" + re.search(r'Target="\.\./(slideLayouts/slideLayout\d+\.xml)"', rels).group(1)


@functools.lru_cache(maxsize=None)
def _layout_deck() -> bytes:
    return _move_picture_to(_layout_of_slide1())


@functools.lru_cache(maxsize=None)
def _master_deck() -> bytes:
    return _move_picture_to("ppt/slideMasters/slideMaster1.xml")


@functools.lru_cache(maxsize=None)
def _twice_on_one_slide() -> bytes:
    """ONE picture placed TWICE on ONE slide — one media part, one rId, two <p:cNvPr>."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Open enrolment"
    s.shapes.add_picture(str(_png()), Inches(0.5), Inches(2), Inches(3.5), Inches(1.2))
    s.shapes.add_picture(str(_png()), Inches(4.5), Inches(2), Inches(3.5), Inches(1.2))
    out = Path(tempfile.mkdtemp()) / "twice.pptx"
    prs.save(out)
    return out.read_bytes()


@functools.lru_cache(maxsize=None)
def _orphan_deck() -> bytes:
    """A raster in ppt/media that NO part references — still carded by 1.4.5, reachable by nothing.

    The residual wedge after the ALT_TARGETS widening, and what the visibility fix is for.
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_plain_deck())) as zin, \
         zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info.filename))
        zout.writestr("ppt/media/image9.png", _png().read_bytes())
    return buf.getvalue()


def _spill(data: bytes, name: str) -> Path:
    p = Path(tempfile.mkdtemp()) / name
    p.write_bytes(data)
    return p


@functools.lru_cache(maxsize=None)
def _assess_cached(data: bytes, name: str) -> frozenset:
    from assessment_policy import _extract_sc
    from scanner import analyse_and_assess
    with tempfile.TemporaryDirectory() as d:
        (Path(d) / name).write_bytes(data)
        fd, _ = analyse_and_assess(Path(d), name, detect_pii=False)
    return frozenset(sc for i in (fd or {}).get("issues", [])
                     if (sc := _extract_sc(i.get("wcag", ""))))


def _assess(data: bytes, name: str = "deck.pptx") -> set[str]:
    """The SCs a REAL assessment reports. Keyed on the bytes — an OCR scan per call would cost
    minutes of the backend suite re-deriving an answer that cannot have changed."""
    return set(_assess_cached(data, name))


@functools.lru_cache(maxsize=None)
def _locators(data: bytes, name: str = "deck.pptx") -> list[str]:
    """The locators the REAL 1.4.5 proposer mints — never hand-written."""
    import proposals
    p = _spill(data, name)
    return [x["locator"] for x in proposals.propose_images_of_text(p, ".pptx")]


def _descrs(data: bytes, part: str) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read(part).decode("utf-8")
    return re.findall(r'<p:cNvPr[^>]*\bdescr="([^"]*)"', xml)


def _media(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [n for n in z.namelist() if n.startswith("ppt/media/")]


class _Blob:
    """The only thing patched here. Stores bytes verbatim; decides nothing."""

    def __init__(self, data: bytes):
        self.data, self.uploads = data, []

    def download_remediated(self, owner, sid, f):
        return self.data

    def upload_remediated(self, owner, sid, f, data, mime):
        self.data = data
        self.uploads.append(f)
        return "http://b/2"


@pytest.fixture()
def store(monkeypatch):
    import store as store_mod
    monkeypatch.setattr(store_mod, "_SQLITE_PATH", Path(tempfile.mkdtemp()) / "inh.db")
    return store_mod.Store()


SID = "described-inherited"


def _seed(store, name: str, locators: list[str]) -> str:
    store.init_scan_run(SID, "drive", 1, "2026-09-07T00:00:00Z", "rubric", "hash")
    store.save_file_result(SID, {
        "file": name, "engine": "office", "status": "pass", "score": 60, "compliant": 0,
        "skipped_rules": 0, "drive_file_id": "d1",
        "issues": [{"ruleId": "OCR_IMAGE_OF_TEXT", "wcag": "1.4.5 Images of Text",
                    "severity": "SERIOUS", "detail": f"embedded {loc} contains readable text"}
                   for loc in locators],
    }, "2026-09-07T00:00:00Z")
    store.record_remediation(SID, name, drive_write_url="http://d/1", blob_url="http://b/1")
    return store.enqueue_proposals(SID, name, "1.4.5", [
        {"locator": loc, "before": "text baked into an image",
         "proposed_value": "the OCR transcript", "rationale": "r", "source": "OCR"}
        for loc in locators], rule_name="Images of Text")


def _decide(store, item_id, monkeypatch, values):
    """Through the PRODUCTION route, not by poking the store."""
    import core
    from routes.hitl import HitlUpdate, hitl_update
    monkeypatch.setattr(core, "store", store)
    return hitl_update(item_id, HitlUpdate(status="approved", **viewed_fields(item_id),
                                           resolution=store.DESCRIBED_RESOLUTION,
                                           approved_values=values), None)


def _run_lane(monkeypatch, store, blob, name):
    """The production handler, with the re-scan UNPATCHED."""
    import core
    import handlers
    monkeypatch.setattr(core, "store", store)
    monkeypatch.setitem(sys.modules, "blob", blob)
    handlers._apply_approved_values({"scan_id": SID, "file": name}, {})


# ------------------------------------------------------------------ the fixtures are honest

@needs_ocr
@pytest.mark.parametrize("builder,part", [
    (_layout_deck, "layout"), (_master_deck, "ppt/slideMasters/slideMaster1.xml")])
def test_the_inherited_picture_really_left_the_slide(builder, part):
    """Without this the round trips below are vacuous — they would be describing a slide picture."""
    data = builder()
    assert _media(data) == ["ppt/media/image1.png"]
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert "<p:pic" not in z.read("ppt/slides/slide1.xml").decode()
    assert "1.4.5" in _assess(data)          # ocr walks the namelist, so it is still carded
    assert _locators(data) == ["image 1"]


# ------------------------------------------------------------------ the parts are now reached

@needs_ocr
@pytest.mark.parametrize("builder", [_layout_deck, _master_deck])
def test_an_inherited_picture_is_now_seen_by_1_1_1_and_reachable(builder):
    """Both halves move together, which is the whole reason ALT_TARGETS is ONE table.

    Before the widening the 1.1.1 detector reported nothing for this image and the resolver
    returned {} — detector blind, applier unable, and the 1.4.5 card pointing at both.
    """
    import apply_alt
    import apply_office_image_of_text as aoit
    from formats.office.images import package_images
    data = builder()

    assert "1.1.1" in _assess(data), "the undescribed inherited picture is a real 1.1.1 finding"
    found = package_images(_spill(data, "deck.pptx"))
    assert len(found) == 1
    part = found[0]["part"]
    assert re.match(r"ppt/(slideLayouts|slideMasters)/", part)
    assert apply_alt.tag_for_part(part) == "p:cNvPr"
    assert aoit.resolve_media_locators(data, ["image 1"], "pptx") == {"image 1": [found[0]["locator"]]}


@needs_ocr
@pytest.mark.parametrize("builder", [_layout_deck, _master_deck])
def test_the_inherited_picture_round_trips(store, monkeypatch, builder):
    """decide → write → REAL re-scan → credit, on the ADR's own first-listed motivating case."""
    data = builder()
    item_id = _seed(store, "deck.pptx", _locators(data))
    _decide(store, item_id, monkeypatch, [DESC])

    rows = {r["rule_id"]: r for r in store.list_hitl_queue(scan_id=SID)}
    described = rows[f"1.1.1{store.DESCRIBED_RULE_SUFFIX}"]
    assert described["status"] == "approved" and not described.get("applied")
    assert store.count_unapplied_approved_values(SID, "deck.pptx") == 1

    blob = _Blob(data)
    _run_lane(monkeypatch, store, blob, "deck.pptx")

    from formats.office.images import package_images
    part = package_images(_spill(data, "deck.pptx"))[0]["part"]
    assert _descrs(blob.data, part) == [DESC]
    assert _media(blob.data) == _media(data)              # the picture STAYED
    assert "1.1.1" not in _assess(blob.data)              # a real re-scan agrees
    assert "1.4.5" in _assess(blob.data)                  # and it is still an image of text
    assert store.get_hitl_item(described["id"])["applied"]
    assert store.count_unapplied_approved_values(SID, "deck.pptx") == 0


@needs_ocr
def test_one_picture_twice_on_one_slide_describes_both(store, monkeypatch):
    """The second hole the slides-only walk had, and it needed no new code to close.

    One slide showing the same picture twice shares ONE relationship id, so the old per-(slide,
    rId) locator reached only the first <p:cNvPr>. The shared walk addresses PLACEMENTS, so both
    are described and 1.1.1 can actually clear.
    """
    data = _twice_on_one_slide()
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        rids = re.findall(r'r:embed="([^"]+)"', z.read("ppt/slides/slide1.xml").decode())
    assert len(rids) == 2 and len(set(rids)) == 1         # two placements, ONE rId
    assert _media(data) == ["ppt/media/image1.png"]

    item_id = _seed(store, "twice.pptx", _locators(data, "twice.pptx"))
    _decide(store, item_id, monkeypatch, [DESC])
    blob = _Blob(data)
    _run_lane(monkeypatch, store, blob, "twice.pptx")

    assert _descrs(blob.data, "ppt/slides/slide1.xml") == [DESC, DESC]
    assert "1.1.1" not in _assess(blob.data, "twice.pptx")
    assert store.count_unapplied_approved_values(SID, "twice.pptx") == 0


@needs_ocr
def test_a_deck_with_no_inherited_pictures_gains_nothing():
    """The blast-radius pin. A blank deck's layouts and masters carry dozens of <p:cNvPr> —
    placeholders, titles, footers — and none of them is an image. The `p:pic` wrapper is what
    keeps the widening from inventing findings, so it is asserted rather than assumed."""
    from pptx import Presentation
    from formats.office.images import package_images
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5])
    blank = _spill(_b := _save(prs), "blank.pptx")
    assert package_images(blank) == []
    with zipfile.ZipFile(io.BytesIO(_b)) as z:
        inherited = [n for n in z.namelist()
                     if re.match(r"ppt/(slideLayouts/slideLayout|slideMasters/slideMaster)\d+", n)]
        cnvpr = sum(len(re.findall(r"<p:cNvPr\b", z.read(n).decode())) for n in inherited)
    assert len(inherited) >= 10 and cnvpr >= 50, "the fixture must actually have elements to skip"


def _save(prs) -> bytes:
    out = Path(tempfile.mkdtemp()) / "blank.pptx"
    prs.save(out)
    return out.read_bytes()


# ------------------------------------------------------------------ the residual wedge is VISIBLE

@needs_ocr
def test_an_unreachable_image_now_tells_the_reviewer_why(store, monkeypatch):
    """The half that matters even when the reach cannot be extended.

    A raster no part references is still carded by 1.4.5 and still reachable by nothing, so the
    description cannot be written and the file cannot certify — that is honest. What was NOT
    honest was the silence: `_apply_one_value_kind` returned at `if not applied:` before logging,
    so nothing rendered a card and the row sat `approved` outside the pending inbox. The reviewer
    saw a file that never published and no reason anywhere.

    Now the lane logs the same apply.unverified shape its other refusals use, and apply_outcome
    reads it back as NOTHING_WRITTEN — a distinct outcome, because the reviewer's next move is
    different: there is nothing to re-run, the image cannot carry alt text at all.
    """
    import apply_outcome
    data = _orphan_deck()
    assert len(_media(data)) == 2
    locs = _locators(data, "orphan.pptx")
    assert "image 2" in locs, "the orphan raster must really be carded"

    item_id = _seed(store, "orphan.pptx", ["image 2"])
    _decide(store, item_id, monkeypatch, [DESC])
    blob = _Blob(data)
    _run_lane(monkeypatch, store, blob, "orphan.pptx")

    assert blob.uploads == [] and blob.data == data      # nothing written, nothing published
    rows = [r for r in store.list_hitl_queue(scan_id=SID)
            if r["rule_id"].startswith("1.1.1")]
    assert rows and not rows[0].get("applied")
    assert store.count_unapplied_approved_values(SID, "orphan.pptx") == 1
    assert store.mark_file_compliant_if_reviewed(SID, "orphan.pptx") is False

    # THE NEW PART: the reviewer is told, instead of being left to guess.
    outcome = apply_outcome.apply_outcome_for(rows[0], store.list_decisions(scan_id=SID))
    assert outcome and outcome["outcome"] == apply_outcome.NOTHING_WRITTEN
    assert outcome["criteria"] == ["1.1.1"], "the card must name the criterion this row owes"


def test_the_nothing_written_shape_is_parsed_from_the_wording_the_lane_writes():
    """Pure-parse half, so the shape cannot drift from the producer without failing here.

    The detail string is copied from handlers._apply_one_value_kind rather than paraphrased —
    a paraphrase would let the two diverge and this test would go on passing.
    """
    import apply_outcome
    detail = ("wrote no description value(s) for ['1.1.1']: all 1 approved locator(s) reach no "
              "image in this document. Credit withheld; the approved value is kept for retry")
    parsed = apply_outcome.parse_unverified(detail)
    assert parsed == {"outcome": apply_outcome.NOTHING_WRITTEN,
                      "criteria": ["1.1.1"], "reason": ""}
    # And it must not be confused with either older shape.
    assert apply_outcome.parse_unverified(
        "wrote 2 description value(s) but ['1.1.1'] still fails on re-scan"
    )["outcome"] == apply_outcome.STILL_FAILING
    assert apply_outcome.parse_unverified(
        "wrote 2 description value(s) but could not verify ['1.1.1']: engine missing. "
        "Credit withheld; the approved value is kept for retry"
    )["outcome"] == apply_outcome.COULD_NOT_VERIFY
