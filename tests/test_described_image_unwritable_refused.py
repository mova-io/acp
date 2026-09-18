"""A described decision is refused for an image no writer can reach — the last ADR 0055 wedge.

`ocr._ooxml_images` walks the whole ZIP namelist; the appliers reach only the parts in
`formats/office/images.ALT_TARGETS`. #1767 widened that to layouts and masters, which closed the
case ADR 0055 opens with. What it could not close is the residual: a media part that NO
alt-bearing part references — a Word footnote image, a VML sheet graphic, a raster left in
`ppt/media` by an editor that dropped its placement.

Such an image is still carded by 1.4.5, and describing it recorded an obligation nothing could
meet: the `1.1.1/described` row stayed approved-and-unapplied forever,
`count_unapplied_approved_values` counted it forever, and the file could never certify.
Re-running the apply job was not self-healing. #1767 made that VISIBLE (the reviewer gets a
NOTHING_WRITTEN card instead of silence). This makes it IMPOSSIBLE TO ACCEPT.

WHERE THE CHECK LIVES, AND WHY NOT IN THE ROUTE. Reachability needs the package, and reading the
document inside the review request is what #1742 deliberately avoided: 'image N' is a media
INDEX, so resolving it against a copy that is not the one it was minted from can name a DIFFERENT
PICTURE. `handlers._mark_describable` therefore computes it at PROPOSE time, where the bytes are
already open, and stamps `describable` on each proposal. The flag then describes exactly the
bytes the reviewer's card was minted from.

That is safe across the original -> remediated hop because `remediate_office` rebuilds the
package from `z.namelist()` and writes every entry back, deleting no media — checked, not
assumed. The tests below drive the REAL proposer and the REAL route, so the flag is production's
answer rather than one this file invented.

ABSENT MEANS UNKNOWN, AND UNKNOWN DOES NOT REFUSE. Rows enqueued before the flag existed carry no
`describable`, and refusing them would break decisions that work today; they keep the older
behaviour, which #1767 at least made visible. `test_a_row_from_before_the_flag_is_not_refused`
pins that, because a migration that quietly starts rejecting old rows is the failure mode this
choice exists to avoid.
"""
from __future__ import annotations

import functools
import glob
import io
import sys
import tempfile
import zipfile
from pathlib import Path

import pytest
from hitl_viewed import viewed_fields

ACP = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ACP / "api"))

pytest.importorskip("pptx")
pytest.importorskip("PIL")

SID = "s-unwritable"
FILE = "orphan.pptx"
DESC = "A benefits summary card listing medical, dental and vision cover."
LINES = ("Benefits at a glance", "Medical dental and vision cover", "Enrollment closes Friday")


def _ocr_ready() -> bool:
    import ocr
    return ocr.is_available()


needs_ocr = pytest.mark.skipif(not _ocr_ready(),
                               reason="tesseract/pytesseract unavailable — 1.4.5 cannot fire")


def _font():
    from PIL import ImageFont
    for pat in ("/usr/share/fonts/**/DejaVuSans.ttf", "/usr/share/fonts/**/*.ttf"):
        hits = glob.glob(pat, recursive=True)
        if hits:
            return ImageFont.truetype(hits[0], 34)
    return ImageFont.load_default()


@functools.lru_cache(maxsize=None)
def _png() -> Path:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (760, 260), "white")
    d = ImageDraw.Draw(im)
    f = _font()
    for i, line in enumerate(LINES):
        d.text((24, 24 + i * 70), line, fill="black", font=f)
    p = Path(tempfile.mkdtemp()) / "text.png"
    im.save(p)
    return p


@functools.lru_cache(maxsize=None)
def _deck_with_orphan() -> bytes:
    """One picture a slide really shows, plus one raster NOTHING references.

    Both are carded by 1.4.5 — `ocr._ooxml_images` reads the namelist — so the deck exercises the
    distinction rather than only the unhappy half: the first image must stay describable, or a
    test that only ever refuses would pass against a check that refuses everything.
    """
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Open enrolment"
    s.shapes.add_picture(str(_png()), Inches(1), Inches(2), Inches(4), Inches(1.4))
    out = Path(tempfile.mkdtemp()) / FILE
    prs.save(out)

    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(out.read_bytes())) as zin, \
         zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info.filename))
        zout.writestr("ppt/media/image9.png", _png().read_bytes())   # referenced by nothing
    return buf.getvalue()


@pytest.fixture()
def store(monkeypatch):
    import store as store_mod
    monkeypatch.setattr(store_mod, "_SQLITE_PATH", Path(tempfile.mkdtemp()) / "unw.db")
    return store_mod.Store()


def _proposals(data: bytes) -> list[dict]:
    """What the REAL proposer mints, stamped by the REAL production marker."""
    import handlers
    import proposals as prop
    p = Path(tempfile.mkdtemp()) / FILE
    p.write_bytes(data)
    return handlers._mark_describable(prop.propose_images_of_text(p, ".pptx"), data, FILE, SID)


def _seed(store, props: list[dict]) -> str:
    store.init_scan_run(SID, "drive", 1, "2026-09-07T00:00:00Z", "rubric", "hash")
    store.save_file_result(SID, {
        "file": FILE, "engine": "office", "status": "pass", "score": 60, "compliant": 0,
        "skipped_rules": 0, "drive_file_id": "d1",
        "issues": [{"ruleId": "OCR_IMAGE_OF_TEXT", "wcag": "1.4.5 Images of Text",
                    "severity": "SERIOUS", "detail": "embedded image contains readable text"}],
    }, "2026-09-07T00:00:00Z")
    store.record_remediation(SID, FILE, drive_write_url="http://d/1", blob_url="http://b/1")
    return store.enqueue_proposals(SID, FILE, "1.4.5", props, rule_name="Images of Text")


def _decide(store, item_id, monkeypatch, values):
    """Through the PRODUCTION route."""
    import core
    from routes.hitl import HitlUpdate, hitl_update
    monkeypatch.setattr(core, "store", store)
    return hitl_update(item_id, HitlUpdate(status="approved", **viewed_fields(item_id),
                                           resolution=store.DESCRIBED_RESOLUTION,
                                           approved_values=values), None)


# --------------------------------------------------------------- the flag is production's answer

@needs_ocr
def test_the_proposer_marks_the_orphan_undescribable_and_the_real_picture_describable():
    """Both halves, from the real proposer — the fixture is only honest if it carries both."""
    props = _proposals(_deck_with_orphan())
    by_loc = {p["locator"]: p.get("describable") for p in props}
    assert len(by_loc) == 2, f"both rasters must be carded, got {by_loc}"
    shown, orphan = sorted(by_loc)          # 'image 1' is the placed one, 'image 2' the orphan
    assert by_loc[shown] is True, "a picture a slide actually shows must stay describable"
    assert by_loc[orphan] is False, "a raster nothing references must be marked unwritable"


@needs_ocr
def test_the_flag_agrees_with_the_resolver_that_would_do_the_writing():
    """The flag is not a second opinion. It is `resolve_media_locators` asked in advance, so the
    check and the write cannot disagree about what is addressable — the drift this repo has lost
    days to. Asserted against the resolver itself rather than against a remembered answer."""
    from apply_office_image_of_text import resolve_media_locators
    data = _deck_with_orphan()
    props = _proposals(data)
    resolved = resolve_media_locators(data, [p["locator"] for p in props], "pptx")
    for p in props:
        assert p["describable"] == (p["locator"] in resolved), p["locator"]


# --------------------------------------------------------------- the decision is refused

@needs_ocr
def test_describing_the_unreachable_image_is_refused_and_changes_nothing(store, monkeypatch):
    """The fix. Refused at the boundary, with the row left exactly as unresolved as it was —
    because the old failure was not the error, it was what the acceptance left behind."""
    from fastapi import HTTPException
    props = _proposals(_deck_with_orphan())
    orphan = next(p for p in props if p["describable"] is False)
    item_id = _seed(store, props)
    values = ["" if p is not orphan else DESC for p in props]

    with pytest.raises(HTTPException) as e:
        _decide(store, item_id, monkeypatch, values)
    assert e.value.status_code == 422
    assert orphan["locator"] in str(e.value.detail)
    assert "another way" in str(e.value.detail)          # tells them what to do instead

    row = store.get_hitl_item(item_id)
    assert row["status"] == "pending" and not (row.get("resolution") or "").strip()
    assert [r["rule_id"] for r in store.list_hitl_queue(scan_id=SID)] == ["1.4.5"]
    assert store.mark_file_compliant_if_reviewed(SID, FILE) is False


@needs_ocr
def test_describing_only_the_reachable_image_is_accepted(store, monkeypatch):
    """The control, and the reason the refusal is per-image rather than per-row. A deck with one
    unwritable raster must not become undescribable altogether."""
    props = _proposals(_deck_with_orphan())
    shown = next(p for p in props if p["describable"] is True)
    item_id = _seed(store, props)
    _decide(store, item_id, monkeypatch, [DESC if p is shown else "" for p in props])

    owed = store.approved_alt_values(SID, FILE)
    assert owed == {shown["locator"]: DESC}
    assert store.count_unapplied_approved_values(SID, FILE) == 1


@needs_ocr
def test_a_row_from_before_the_flag_is_not_refused(store, monkeypatch):
    """UNKNOWN is not FALSE. A row enqueued before handlers._mark_describable existed carries no
    `describable`, and must keep working: a migration that quietly starts rejecting decisions
    reviewers were making yesterday is worse than the wedge it set out to close, and #1767
    already made that wedge visible rather than silent."""
    props = [{k: v for k, v in p.items() if k != "describable"}
             for p in _proposals(_deck_with_orphan())]
    assert all("describable" not in p for p in props)
    item_id = _seed(store, props)
    _decide(store, item_id, monkeypatch, [DESC] + [""] * (len(props) - 1))
    assert store.count_unapplied_approved_values(SID, FILE) == 1
