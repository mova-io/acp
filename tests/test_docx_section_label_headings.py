"""Body-size SECTION LABELS in docx — "ACTIVITY:", "BATHING:" — as headings (1.3.1).

The reported document: short bold, underlined, UPPERCASE paragraphs at body size, each followed
by numbered instructions. They are the document's sections, and they were invisible to ACP: the
only heading signal was size (office_structure.looks_like_pseudo_heading, ≥14pt, or ≥13pt bold),
read from direct run XML only, and docx_checks stopped at the first hit. So a 12pt label never
fired, bold coming from a style never counted, and a document with ten candidates reported one.

These tests build REAL saved .docx packages (zip + OOXML on disk; one set through python-docx's
own template) and run the real detector and the real remediator over them. No user document is
involved: the screenshot that prompted this proves the layout, not the file's metadata, so the
fixtures reproduce the layout in each of the ways Word can encode it (direct formatting,
docDefaults size, character style, paragraph style).

Contract pinned here (other code depends on it): office_structure.docx_heading_candidates is the
ONE list both sides use; a candidate is "strong" (auto-promoted) or "review" (flagged, never
restyled); docx_checks emits one DOCX_PSEUDO_HEADING per candidate at `location` word:p:N, the
promoter's own locator numbering.
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

ACP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ACP / "api"))

from lxml import etree  # noqa: E402

import office_structure as osx      # noqa: E402
import remediate_office as R        # noqa: E402

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{" + W_NS + "}"

# ── fixture builder: a complete, Word-openable package written to disk ─────────────────────────

_CT = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>
<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
</Types>"""
_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
</Relationships>"""
_DOC_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/>
</Relationships>"""
# Title and language already set, so remediation's only work is the structure under test.
_CORE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Discharge instructions</dc:title><dc:language>en-US</dc:language></cp:coreProperties>"""
_NUMBERING = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="{W_NS}">
<w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl></w:abstractNum>
<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>
</w:numbering>"""


def _styles(extra: str = "", default_sz: int | None = 22) -> str:
    dd = f'<w:rPrDefault><w:rPr><w:sz w:val="{default_sz}"/></w:rPr></w:rPrDefault>' if default_sz else ""
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:styles xmlns:w="{W_NS}"><w:docDefaults>{dd}</w:docDefaults>'
            f'<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
            f'<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/>'
            f'<w:rPr><w:sz w:val="56"/></w:rPr></w:style>'
            f"{extra}</w:styles>")


def _rpr(*, b=False, u=False, caps=False, sz=None, rstyle=None, b_off=False) -> str:
    inner = ((f'<w:rStyle w:val="{rstyle}"/>' if rstyle else "")
             + ("<w:b/>" if b else "") + ('<w:b w:val="0"/>' if b_off else "")
             + ("<w:caps/>" if caps else "")
             + (f'<w:sz w:val="{sz}"/>' if sz else "")
             + ('<w:u w:val="single"/>' if u else ""))
    return f"<w:rPr>{inner}</w:rPr>" if inner else ""


def _p(text: str = "", *, pstyle=None, num=None, runs=None, **fmt) -> str:
    ppr = ((f'<w:pStyle w:val="{pstyle}"/>' if pstyle else "")
           + (f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num}"/></w:numPr>' if num else ""))
    ppr = f"<w:pPr>{ppr}</w:pPr>" if ppr else ""
    if runs is None:
        runs = [(text, fmt)]
    body = "".join(f'<w:r>{_rpr(**f)}<w:t xml:space="preserve">{t}</w:t></w:r>' for t, f in runs)
    return f"<w:p>{ppr}{body}</w:p>"


def label(text: str, **over) -> str:
    """The reported label: 12pt, bold, underlined, typed in capitals, body-styled."""
    fmt = {"b": True, "u": True, "sz": 24, **over}
    return _p(text, **fmt)


def item(text: str) -> str:
    return _p(text, num=1)


def body(text: str) -> str:
    return _p(text)


TITLE = _p("Discharge Instructions", b=True, sz=32)          # 16pt bold: the size path's H1
INTRO = body("Please read these instructions carefully before you leave the hospital today.")
CLOSE = body("Call your doctor if you have any questions about these instructions.")
ACTIVITY = [label("ACTIVITY:"), item("Walk for ten minutes three times a day."),
            item("Do not lift anything heavier than five pounds."),
            item("Rest when you feel tired.")]
BATHING = [label("BATHING:"), item("You may shower 48 hours after surgery."),
           item("Do not soak in a bath until your wound has healed.")]


def write_docx(tmp: Path, paragraphs, *, styles: str | None = None, name="doc.docx") -> Path:
    doc = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<w:document xmlns:w="{W_NS}"><w:body>{"".join(paragraphs)}'
           f'<w:sectPr><w:pgSz w:w="12240" w:h="15840"/></w:sectPr></w:body></w:document>')
    path = tmp / name
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _CT)
        z.writestr("_rels/.rels", _RELS)
        z.writestr("docProps/core.xml", _CORE)
        z.writestr("word/_rels/document.xml.rels", _DOC_RELS)
        z.writestr("word/document.xml", doc)
        z.writestr("word/styles.xml", styles if styles is not None else _styles())
        z.writestr("word/numbering.xml", _NUMBERING)
    return path


def part(path: Path, name: str) -> bytes:
    with zipfile.ZipFile(path) as z:
        return z.read(name)


def candidates(path: Path):
    return osx.docx_heading_candidates(part(path, "word/document.xml"),
                                       part(path, "word/styles.xml"),
                                       part(path, "word/numbering.xml"))


def pseudo_findings(path: Path):
    return [f for f in osx.docx_checks(path) if f["ruleId"] == "DOCX_PSEUDO_HEADING"]


def outline(path: Path) -> list[tuple[str, str]]:
    """[(text, pStyle)] for every paragraph with a pStyle, in document order."""
    root = etree.fromstring(part(path, "word/document.xml"))
    out = []
    for p in root.iter(W + "p"):
        st = p.find(f"{W}pPr/{W}pStyle")
        if st is not None:
            out.append(("".join(t.text or "" for t in p.iter(W + "t")), st.get(W + "val")))
    return out


def by_text(path: Path) -> dict[str, str]:
    return dict(outline(path))


# ── STEP 1: the miss, reproduced ────────────────────────────────────────────────────────────────

def test_the_size_predicate_alone_cannot_see_a_12pt_section_label():
    """Why this path exists: at 12pt bold the absolute size question says no, and so does the
    relative one — the label is the SAME size as the text around it. Lowering the floor would
    make every emphasised sentence a heading; the section-label path uses different evidence."""
    assert osx.looks_like_pseudo_heading("ACTIVITY:", bold=True, max_half_pt=24,
                                         styled_heading=False) is False
    assert osx.heading_signal("ACTIVITY:", bold=True, max_half_pt=24, body_half_pt=22,
                              styled_heading=False) is None
    # …and with no run size at all (inherited from docDefaults) the size path reads 0.
    assert osx.heading_signal("ACTIVITY:", bold=True, max_half_pt=0, body_half_pt=22,
                              styled_heading=False) is None


# ── the reported layout, end to end ─────────────────────────────────────────────────────────────

def _reported(tmp_path, **kw) -> Path:
    return write_docx(tmp_path, [TITLE, INTRO, *ACTIVITY, *BATHING, CLOSE], **kw)


def test_each_label_is_its_own_finding_with_a_distinct_location(tmp_path):
    src = _reported(tmp_path)
    labels = [f for f in pseudo_findings(src) if f.get("basis") == "section-label"]
    assert [f["location"] for f in labels] == ["word:p:3", "word:p:7"]
    assert all(f["severity"] == "MODERATE" and f["signal"] == "strong" for f in labels)
    assert "ACTIVITY:" in labels[0]["detail"] and "BATHING:" in labels[1]["detail"]
    assert "Paragraph 3" in labels[0]["detail"] and "Heading 2" in labels[0]["detail"]
    # Never collapsed: no two findings are the same record.
    allf = pseudo_findings(src)
    assert len({(f["detail"], f["location"]) for f in allf}) == len(allf) == 3   # + the title


def test_candidates_carry_the_contract_fields(tmp_path):
    cs = candidates(_reported(tmp_path))
    assert [(c.paragraph_index, c.locator, c.text, c.signal, c.basis, c.level) for c in cs] == [
        (1, "word:p:1", "Discharge Instructions", "strong", "size", 1),
        (3, "word:p:3", "ACTIVITY:", "strong", "section-label", 2),
        (7, "word:p:7", "BATHING:", "strong", "section-label", 2),
    ]
    act = cs[1]
    assert act.formatting == "bold, underlined, all caps, 12pt" and act.peers == 2
    assert act.reason
    with pytest.raises(Exception):
        act.level = 3                                         # frozen


def test_remediation_promotes_labels_to_h2_under_the_title_and_rescan_clears(tmp_path):
    src = _reported(tmp_path)
    diffs: list = []
    fixed, applied, skipped = R.remediate_office(src, ai_enabled=False, diffs=diffs)
    assert fixed is not None and any("pseudo-heading" in a for a in applied)
    assert outline(fixed) == [("Discharge Instructions", "Heading1"),
                              ("ACTIVITY:", "Heading2"), ("BATHING:", "Heading2")]
    assert pseudo_findings(Path(fixed)) == []
    assert [f for f in osx.docx_checks(Path(fixed)) if f["ruleId"] == "DOCX_HEADING_SKIP"] == []
    # Lock-step: exactly the strong candidates were promoted, each recorded at its own locator.
    promoted = [d["locator"] for d in diffs if d["rule_id"] == "1.3.1" and "promoted to Heading" in d["after"]]
    assert promoted == [c.locator for c in candidates(src) if c.signal == "strong"]


def _canon_without_pstyle(xml: bytes) -> list[bytes]:
    root = etree.fromstring(xml)
    out = []
    for p in root.iter(W + "p"):
        ppr = p.find(W + "pPr")
        if ppr is not None:
            st = ppr.find(W + "pStyle")
            if st is not None:
                ppr.remove(st)
            if len(ppr) == 0 and not ppr.attrib:
                p.remove(ppr)
        out.append(etree.tostring(p, method="c14n"))
    return out


def test_promotion_changes_pstyle_and_nothing_else(tmp_path):
    src = _reported(tmp_path)
    fixed, _applied, _skipped = R.remediate_office(src, ai_enabled=False)
    # Every paragraph — runs, text, run formatting, numbering references — identical once the
    # pStyle is set aside, and the numbering part is byte-for-byte the same.
    assert (_canon_without_pstyle(part(fixed, "word/document.xml"))
            == _canon_without_pstyle(part(src, "word/document.xml")))
    assert part(fixed, "word/numbering.xml") == part(src, "word/numbering.xml")


def test_the_referenced_heading_styles_exist_and_do_not_restyle_the_text(tmp_path):
    """Word writes only the styles a document uses; a pStyle naming an undefined style falls
    back to Normal and is NOT in the outline. The added definitions carry an outline level and no
    run formatting, so the label keeps exactly its look."""
    src = _reported(tmp_path)
    fixed, _applied, _skipped = R.remediate_office(src, ai_enabled=False)
    styles = etree.fromstring(part(fixed, "word/styles.xml"))
    for lvl in (1, 2):
        st = styles.find(f'{W}style[@{W}styleId="Heading{lvl}"]')
        assert st is not None, f"Heading{lvl} referenced but not defined"
        assert st.find(f"{W}pPr/{W}outlineLvl").get(W + "val") == str(lvl - 1)
        assert st.find(f"{W}rPr") is None
        assert st.find(f"{W}basedOn").get(W + "val") == "Normal"
    # …and the promoted document reads back through the shared resolver as headings.
    assert candidates(Path(fixed)) == []


def test_second_remediation_is_a_no_op(tmp_path):
    fixed, _applied, _skipped = R.remediate_office(_reported(tmp_path), ai_enabled=False)
    again, applied2, _skipped2 = R.remediate_office(Path(fixed), ai_enabled=False)
    assert not any("pseudo-heading" in a or "Heading 1" in a for a in applied2), applied2
    if again is not None:
        assert part(again, "word/document.xml") == part(fixed, "word/document.xml")
        assert part(again, "word/styles.xml") == part(fixed, "word/styles.xml")


def test_labels_nest_one_below_the_heading_they_sit_under(tmp_path):
    """Level is relative to the outline in document order: under each real Heading 2 the labels
    are Heading 3. The outline pass then has nothing to renumber and no level is skipped."""
    styles = _styles(extra="".join(
        f'<w:style w:type="paragraph" w:styleId="Heading{n}"><w:name w:val="heading {n}"/>'
        f'<w:basedOn w:val="Normal"/><w:pPr><w:outlineLvl w:val="{n - 1}"/></w:pPr></w:style>'
        for n in (1, 2)))
    src = write_docx(tmp_path, [
        _p("Discharge Instructions", pstyle="Heading1"), INTRO,
        _p("Recovery at home", pstyle="Heading2"), *ACTIVITY, *BATHING,
        _p("Follow-up", pstyle="Heading2"), label("APPOINTMENTS:"), item("See Dr Rivera in 2 weeks."),
        label("QUESTIONS:"), item("Call the ward."), CLOSE], styles=styles)
    got = {c.text: c.level for c in candidates(src) if c.signal == "strong"}
    assert got == {"ACTIVITY:": 3, "BATHING:": 3, "APPOINTMENTS:": 3, "QUESTIONS:": 3}
    fixed, _a, _s = R.remediate_office(src, ai_enabled=False)
    assert [v for _t, v in outline(fixed)] == ["Heading1", "Heading2", "Heading3", "Heading3",
                                               "Heading2", "Heading3", "Heading3"]
    assert pseudo_findings(Path(fixed)) == []


def test_short_caps_labels_with_a_colon_are_peers_but_wordmarks_are_not(tmp_path):
    """≤4-character all-caps text is normally a wordmark ("ACME") and stays furniture. A colon
    turns it into a label introducing what follows ("DIET:"), so it counts as a peer."""
    src = write_docx(tmp_path, [TITLE, INTRO, *ACTIVITY, label("DIET:"),
                                item("Eat small meals."), *BATHING,
                                _p("ACME", b=True, u=True, sz=24), CLOSE])
    labels = [(c.text, c.signal) for c in candidates(src) if c.basis == "section-label"]
    assert labels == [("ACTIVITY:", "strong"), ("DIET:", "strong"), ("BATHING:", "strong")]


def test_a_real_word_template_document(tmp_path):
    """Through python-docx's own template: a real Heading 1, labels formatted with run
    properties, list items styled "List Number". The labels land one level below the title."""
    docx = pytest.importorskip("docx")
    from docx.shared import Pt
    d = docx.Document()
    d.add_heading("Wound Care Instructions", level=1)
    d.add_paragraph("Follow these steps until your follow-up appointment.")
    for name, steps in (("ACTIVITY:", ["Walk daily.", "Avoid heavy lifting."]),
                        ("BATHING:", ["Shower after 48 hours.", "Pat the wound dry."]),
                        ("WHEN TO CALL:", ["Fever over 101F.", "Redness or swelling."])):
        run = d.add_paragraph().add_run(name)
        run.bold, run.underline, run.font.size = True, True, Pt(12)
        for s in steps:
            d.add_paragraph(s, style="List Number")
    src = tmp_path / "template.docx"
    d.save(src)

    assert [f["signal"] for f in pseudo_findings(src)] == ["strong"] * 3
    fixed, _applied, _skipped = R.remediate_office(src, ai_enabled=False)
    got = by_text(Path(fixed))
    assert got["Wound Care Instructions"] == "Heading1"
    assert [got[t] for t in ("ACTIVITY:", "BATHING:", "WHEN TO CALL:")] == ["Heading2"] * 3
    assert pseudo_findings(Path(fixed)) == []


# ── inherited formatting ────────────────────────────────────────────────────────────────────────

def test_size_inherited_from_doc_defaults(tmp_path):
    """No w:sz on any run — every size comes from docDefaults (12pt). Labels still group as
    peers; the (real) Heading 1 gives them a parent."""
    styles = _styles(default_sz=24, extra=(
        '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/>'
        '<w:basedOn w:val="Normal"/><w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style>'))
    src = write_docx(tmp_path, [_p("Discharge Instructions", pstyle="Heading1"), INTRO,
                                label("ACTIVITY:", sz=None), item("Walk daily."),
                                label("BATHING:", sz=None), item("Shower after two days."), CLOSE],
                     styles=styles)
    cs = candidates(src)
    assert [(c.text, c.signal, c.level, c.formatting) for c in cs] == [
        ("ACTIVITY:", "strong", 2, "bold, underlined, all caps, 12pt"),
        ("BATHING:", "strong", 2, "bold, underlined, all caps, 12pt")]
    fixed, _a, _s = R.remediate_office(src, ai_enabled=False)
    assert by_text(Path(fixed))["BATHING:"] == "Heading2"


def test_bold_from_a_character_style(tmp_path):
    """Bold and underline come from a character style (w:rStyle) — run-level formatting, so it
    survives the paragraph getting a heading style, and the label is auto-promoted."""
    styles = _styles(extra=(
        '<w:style w:type="character" w:styleId="LabelChar"><w:name w:val="Label Char"/>'
        '<w:rPr><w:b/><w:u w:val="single"/></w:rPr></w:style>'))
    src = write_docx(tmp_path, [TITLE, INTRO,
                                label("ACTIVITY:", b=False, u=False, rstyle="LabelChar"),
                                item("Walk daily."),
                                label("BATHING:", b=False, u=False, rstyle="LabelChar"),
                                item("Shower after two days."), CLOSE], styles=styles)
    assert [(c.text, c.signal, c.level) for c in candidates(src) if c.basis == "section-label"] == [
        ("ACTIVITY:", "strong", 2), ("BATHING:", "strong", 2)]
    fixed, _a, _s = R.remediate_office(src, ai_enabled=False)
    assert by_text(Path(fixed))["ACTIVITY:"] == "Heading2"
    assert pseudo_findings(Path(fixed)) == []


def test_bold_from_a_custom_paragraph_style_is_detected_but_left_for_review(tmp_path):
    """Bold comes from a custom paragraph style based on Normal. Detected — but promoting it would
    REPLACE that style with the heading style and drop the bold, so it is a review item, and the
    remediator leaves it alone and says why."""
    styles = _styles(extra=(
        '<w:style w:type="paragraph" w:styleId="SectionLabel"><w:name w:val="Section Label"/>'
        '<w:basedOn w:val="Normal"/><w:rPr><w:b/><w:u w:val="single"/></w:rPr></w:style>'))
    src = write_docx(tmp_path, [TITLE, INTRO,
                                label("ACTIVITY:", b=False, u=False, pstyle="SectionLabel"),
                                item("Walk daily."),
                                label("BATHING:", b=False, u=False, pstyle="SectionLabel"),
                                item("Shower after two days."), CLOSE], styles=styles)
    labels = [c for c in candidates(src) if c.basis == "section-label"]
    assert [(c.text, c.signal, c.reason) for c in labels] == [
        ("ACTIVITY:", "review", osx.REASON_LABEL_PARA_STYLE),
        ("BATHING:", "review", osx.REASON_LABEL_PARA_STYLE)]
    assert all(f["severity"] == "REVIEW" for f in pseudo_findings(src) if f["basis"] == "section-label")
    fixed, _a, skipped = R.remediate_office(src, ai_enabled=False)
    assert by_text(Path(fixed))["ACTIVITY:"] == "SectionLabel"
    assert any("word:p:3" in s and "paragraph style" in s for s in skipped), skipped


def test_bold_switched_off_is_not_bold(tmp_path):
    """w:b w:val="0" on the run overrides a bold character style: not emphasised, not a label."""
    styles = _styles(extra=(
        '<w:style w:type="character" w:styleId="LabelChar"><w:name w:val="Label Char"/>'
        '<w:rPr><w:b/></w:rPr></w:style>'))
    src = write_docx(tmp_path, [TITLE, INTRO,
                                label("ACTIVITY:", b=False, b_off=True, rstyle="LabelChar"),
                                item("Walk daily."),
                                label("BATHING:", b=False, b_off=True, rstyle="LabelChar"),
                                item("Shower."), CLOSE], styles=styles)
    assert [c for c in candidates(src) if c.basis == "section-label"] == []


# ── already structured: never flagged, never re-levelled ────────────────────────────────────────

def test_existing_custom_heading_styles_and_outline_levels_are_recognised(tmp_path):
    styles = _styles(extra=(
        '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/>'
        '<w:basedOn w:val="Normal"/><w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/>'
        '<w:basedOn w:val="Normal"/><w:pPr><w:outlineLvl w:val="1"/></w:pPr></w:style>'
        # a custom style based on a heading style
        '<w:style w:type="paragraph" w:styleId="ClinicSection"><w:name w:val="Clinic Section"/>'
        '<w:basedOn w:val="Heading2"/><w:rPr><w:b/><w:u w:val="single"/></w:rPr></w:style>'
        # a localised built-in: the id is translated, the name is not
        '<w:style w:type="paragraph" w:styleId="berschrift2"><w:name w:val="heading 2"/>'
        '<w:basedOn w:val="Normal"/></w:style>'))
    paras = [_p("Discharge Instructions", pstyle="Heading1", b=True, sz=32), INTRO,
             label("ACTIVITY:", pstyle="ClinicSection"), item("Walk daily."),
             label("BATHING:", pstyle="berschrift2"), item("Shower after two days."),
             "<w:p><w:pPr><w:outlineLvl w:val=\"1\"/></w:pPr><w:r><w:rPr><w:b/><w:u w:val=\"single\"/>"
             "<w:sz w:val=\"24\"/></w:rPr><w:t>DIET:</w:t></w:r></w:p>", item("Eat small meals."),
             _p("Heading 2 spelled with a space", pstyle="Heading 2", b=True, sz=40), CLOSE]
    src = write_docx(tmp_path, paras, styles=styles)
    assert candidates(src) == []
    assert pseudo_findings(src) == []
    before = outline(src)
    fixed, _a, _s = R.remediate_office(src, ai_enabled=False)
    assert outline(fixed or src) == before


def test_title_style_is_structure_not_a_pseudo_heading(tmp_path):
    """A Title-styled paragraph — even with a huge direct size — is already marked."""
    src = write_docx(tmp_path, [_p("Discharge Instructions", pstyle="Title", b=True, sz=56),
                                INTRO, CLOSE])
    assert candidates(src) == []


# ── negatives ───────────────────────────────────────────────────────────────────────────────────

def test_title_only_document_has_no_section_labels(tmp_path):
    src = write_docx(tmp_path, [TITLE, INTRO, CLOSE])
    assert [c.basis for c in candidates(src)] == ["size"]


def test_a_lone_label_is_review_not_promoted(tmp_path):
    src = write_docx(tmp_path, [TITLE, INTRO, *ACTIVITY, CLOSE])
    (lone,) = [c for c in candidates(src) if c.basis == "section-label"]
    assert (lone.signal, lone.level, lone.reason) == ("review", None, osx.REASON_LABEL_LONE)
    (f,) = [f for f in pseudo_findings(src) if f["basis"] == "section-label"]
    assert f["severity"] == "REVIEW" and f["location"] == "word:p:3"
    fixed, _applied, skipped = R.remediate_office(src, ai_enabled=False)
    assert "ACTIVITY:" not in by_text(Path(fixed))
    assert any("[word:p:3]" in s and "may be section headings" in s for s in skipped), skipped
    # the honest consequence: the review finding is still there after remediation
    assert [f["location"] for f in pseudo_findings(Path(fixed))] == ["word:p:3"]


# ── Title-styled documents: the Title becomes the level-1 anchor ─────────────────────────────────
#
# The likeliest real layout: a first paragraph in Word's Title style, then the labels, no Heading 1
# anywhere. Title is not an outline level, so on their own the labels could be neither H2 (the
# outline pass would lift the first alone to H1) nor H1 (it would demote all but one). The shared
# decision makes the Title the anchor: a DIRECT outlineLvl 0 on it (pStyle untouched), labels H2.

def title_para(style: str = "Title") -> str:
    return (f'<w:p><w:pPr><w:pStyle w:val="{style}"/><w:spacing w:after="240"/>'
            f'<w:jc w:val="center"/><w:rPr><w:b/></w:rPr></w:pPr>'
            f'<w:r><w:rPr><w:b/><w:color w:val="1F3864"/></w:rPr>'
            f"<w:t>Discharge Instructions</w:t></w:r></w:p>")


_DOC_TITLE_STYLE = ('<w:style w:type="paragraph" w:styleId="DocTitle"><w:name w:val="Doc Title"/>'
                    '<w:basedOn w:val="Title"/></w:style>')


def _first_para(xml: bytes):
    return next(etree.fromstring(xml).iter(W + "p"))


@pytest.mark.parametrize("style", ["Title", "DocTitle"])       # by id, and a style based on Title
def test_title_style_document_anchors_on_the_title_and_auto_fixes(tmp_path, style):
    src = write_docx(tmp_path, [title_para(style), INTRO, *ACTIVITY, *BATHING, CLOSE],
                     styles=_styles(extra=_DOC_TITLE_STYLE))
    labels = [c for c in candidates(src) if c.basis == "section-label"]
    assert [(c.text, c.signal, c.level, c.anchor_index, c.anchor_locator) for c in labels] == [
        ("ACTIVITY:", "strong", 2, 1, "word:p:1"), ("BATHING:", "strong", 2, 1, "word:p:1")]
    details = [f["detail"] for f in pseudo_findings(src)]
    assert len(details) == 2 and all("Auto-fix: mark it as Heading 2" in d for d in details)
    assert all("Title paragraph 1 is used as the document's level-1 heading" in d for d in details)

    diffs: list = []
    fixed, _applied, _skipped = R.remediate_office(src, ai_enabled=False, diffs=diffs)
    fixed = Path(fixed)
    # The Title: same pStyle, same runs, same pPr — plus exactly one outlineLvl 0, placed where
    # the schema requires it (before pPr/rPr), so Word opens it.
    before = _first_para(part(src, "word/document.xml"))
    after = _first_para(part(fixed, "word/document.xml"))
    ppr = after.find(W + "pPr")
    ols = ppr.findall(W + "outlineLvl")
    assert len(ols) == 1 and ols[0].get(W + "val") == "0"
    assert ppr.find(W + "pStyle").get(W + "val") == style
    assert [el.tag.split("}")[1] for el in ppr] == ["pStyle", "spacing", "jc", "outlineLvl", "rPr"]
    ppr.remove(ols[0])
    assert etree.tostring(after, method="c14n") == etree.tostring(before, method="c14n")
    # The saved outline: the normaliser kept the Title as the one H1 and the labels at H2.
    assert outline(fixed) == [("Discharge Instructions", style),
                              ("ACTIVITY:", "Heading2"), ("BATHING:", "Heading2")]
    assert "word:p:1" in {d.get("locator") for d in diffs}
    assert not any("Heading 1" in d["after"] for d in diffs)
    # Re-scan: the Title now reads as a heading, the labels as headings, nothing left to flag.
    root = etree.fromstring(part(fixed, "word/document.xml"))
    styles = osx._DocxStyles(part(fixed, "word/styles.xml"))
    kinds = [osx._outline_kind(p, styles) for p in root.iter(W + "p")]
    assert kinds[0] == ("heading", 1) and kinds[2] == ("heading", 2) and kinds[6] == ("heading", 2)
    assert pseudo_findings(fixed) == []
    assert [f for f in osx.docx_checks(fixed) if f["ruleId"] == "DOCX_HEADING_SKIP"] == []


def test_title_anchor_remediation_is_idempotent(tmp_path):
    src = write_docx(tmp_path, [title_para(), INTRO, *ACTIVITY, *BATHING, CLOSE])
    fixed, _a, _s = R.remediate_office(src, ai_enabled=False)
    again, applied2, _s2 = R.remediate_office(Path(fixed), ai_enabled=False)
    assert not any("pseudo-heading" in a or "Heading 1" in a for a in applied2), applied2
    if again is not None:
        assert part(again, "word/document.xml") == part(fixed, "word/document.xml")
        assert part(again, "word/styles.xml") == part(fixed, "word/styles.xml")


def test_a_document_with_its_own_heading_1_never_anchors_on_the_title(tmp_path):
    """Under a real Heading 1 the labels nest normally and the Title is left exactly as it was;
    with the Heading 1 AFTER the labels, anchoring would make two H1s, so they stay review."""
    h1 = _p("Overview", pstyle="Heading1")
    below = write_docx(tmp_path, [title_para(), h1, INTRO, *ACTIVITY, *BATHING, CLOSE],
                       name="below.docx")
    assert [(c.level, c.anchor_index) for c in candidates(below)
            if c.basis == "section-label"] == [(2, None), (2, None)]
    fixed, _a, _s = R.remediate_office(below, ai_enabled=False)
    assert (etree.tostring(_first_para(part(fixed, "word/document.xml")), method="c14n")
            == etree.tostring(_first_para(part(below, "word/document.xml")), method="c14n"))

    later = write_docx(tmp_path, [title_para(), INTRO, *ACTIVITY, *BATHING, h1, CLOSE],
                       name="later.docx")
    assert [(c.signal, c.reason) for c in candidates(later) if c.basis == "section-label"] == [
        ("review", osx.REASON_LABEL_NO_PARENT)] * 2
    fixed, _a, _s = R.remediate_office(later, ai_enabled=False)
    first = _first_para(part(Path(fixed or later), "word/document.xml"))
    assert first.find(f"{W}pPr/{W}outlineLvl") is None


def test_a_title_after_the_labels_is_no_anchor(tmp_path):
    src = write_docx(tmp_path, [INTRO, *ACTIVITY, *BATHING, title_para(), CLOSE])
    labels = [c for c in candidates(src) if c.basis == "section-label"]
    assert [(c.signal, c.reason, c.anchor_index) for c in labels] == [
        ("review", osx.REASON_LABEL_NO_PARENT, None)] * 2
    fixed, _a, _s = R.remediate_office(src, ai_enabled=False)
    out = Path(fixed or src)
    assert "ACTIVITY:" not in by_text(out)
    assert b"outlineLvl" not in part(out, "word/document.xml")


def test_no_title_and_no_heading_is_review(tmp_path):
    src = write_docx(tmp_path, [INTRO, *ACTIVITY, *BATHING, CLOSE])
    labels = [c for c in candidates(src) if c.basis == "section-label"]
    assert [(c.signal, c.reason) for c in labels] == [("review", osx.REASON_LABEL_NO_PARENT)] * 2
    assert all(f["severity"] == "REVIEW" for f in pseudo_findings(src))
    fixed, _a, skipped = R.remediate_office(src, ai_enabled=False)
    assert "ACTIVITY:" not in by_text(Path(fixed or src))
    assert any("[word:p:2, word:p:6]" in s for s in skipped), skipped


def test_inline_bold_label_followed_by_plain_text(tmp_path):
    src = write_docx(tmp_path, [TITLE, INTRO] + [
        _p(runs=[(f"{k}:", {"b": True, "u": True}), (f" {v}", {})])
        for k, v in (("PHONE", "555-0100"), ("DOCTOR", "Dr Rivera"), ("CLINIC", "North wing"))]
        + [CLOSE])
    assert [c for c in candidates(src) if c.basis == "section-label"] == []


def test_bold_contact_block(tmp_path):
    src = write_docx(tmp_path, [
        TITLE,
        label("SUNRISE HEALTH CLINIC"), label("42 HARBOR ROAD, SPRINGFIELD"),
        label("TEL: 555-0100"), label("EMAIL: CARE@SUNRISE.ORG"),
        INTRO, CLOSE])
    assert [c for c in candidates(src) if c.basis == "section-label"] == []


def test_bold_sentences_list_items_and_captions(tmp_path):
    src = write_docx(tmp_path, [
        TITLE, INTRO,
        label("TAKE ALL MEDICINE WITH FOOD."), body("It protects your stomach."),
        label("DO NOT DRIVE TODAY!"), body("The anaesthetic takes a day to wear off."),
        _p("WALK DAILY:", num=1, b=True, u=True, sz=24), item("Ten minutes at a time."),
        _p("STAIRS:", num=1, b=True, u=True, sz=24), item("One step at a time."),
        label("FIGURE 1: WOUND DRESSING"), body("The dressing covers the whole incision."),
        label("TABLE 2: DOSES"), body("Doses are listed by weight."),
        label("DEAR PATIENT,"), body("Thank you for choosing our clinic."),
        CLOSE])
    assert [c.text for c in candidates(src) if c.basis == "section-label"] == []


def test_labels_inside_table_cells_are_not_section_headings(tmp_path):
    """A bold label in a cell is the table's row or column heading — table structure (1.3.1 via
    header rows), not the document outline. Excluded outright, never promoted."""
    def cell(*paras):
        return f"<w:tc>{''.join(paras)}</w:tc>"
    tbl = ("<w:tbl><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid>"
           f"<w:tr>{cell(label('ACTIVITY:'))}{cell(body('Walk daily.'))}</w:tr>"
           f"<w:tr>{cell(label('BATHING:'))}{cell(body('Shower after two days.'))}</w:tr>"
           f"<w:tr>{cell(label('DIET:'), body('Small meals.'))}{cell(body('Soft food.'))}</w:tr>"
           "</w:tbl>")
    src = write_docx(tmp_path, [TITLE, INTRO, tbl, CLOSE])
    assert [c for c in candidates(src) if c.basis == "section-label"] == []


def test_paragraph_numbering_matches_the_promoters_locators_with_tables_before(tmp_path):
    """Indices count EVERY w:p in document order, table-cell paragraphs included — the promoter's
    `word:p:N` numbering — so a finding's location names the paragraph the fix edits."""
    tbl = ("<w:tbl><w:tblGrid><w:gridCol/></w:tblGrid>"
           f"<w:tr><w:tc>{body('Patient name')}</w:tc></w:tr>"
           f"<w:tr><w:tc>{body('Room 12')}</w:tc></w:tr></w:tbl>")
    src = write_docx(tmp_path, [TITLE, tbl, INTRO, *ACTIVITY, *BATHING, CLOSE])
    labels = [c for c in candidates(src) if c.basis == "section-label"]
    assert [c.locator for c in labels] == ["word:p:5", "word:p:9"]
    diffs: list = []
    R.remediate_office(src, ai_enabled=False, diffs=diffs)
    assert {"word:p:5", "word:p:9"} <= {d.get("locator") for d in diffs}


def test_namespace_less_parts_still_parse(tmp_path):
    """Hand-built fixtures elsewhere in this suite omit xmlns declarations; the shared resolver
    binds them rather than silently reporting nothing."""
    doc = ("<w:document><w:body>"
           + "".join([TITLE, INTRO, *ACTIVITY, *BATHING, CLOSE]).replace(' xml:space="preserve"', "")
           + "</w:body></w:document>")
    labels = [c for c in osx.docx_heading_candidates(doc) if c.basis == "section-label"]
    assert [c.locator for c in labels] == ["word:p:3", "word:p:7"]
