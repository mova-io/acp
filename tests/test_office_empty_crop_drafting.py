"""An empty or all-zero `a:srcRect` is not a crop — but only where the placement is vouched.

Word writes `<a:srcRect/>` for an inline picture whose crop handles were never moved (or
were reset), and production documents carry it. The shape reproduced on a saved DOCX, by
read-only inspection: one valid PNG, an unnamed inline drawing
(`<wp:docPr id="1" name="" descr="" title=""/>`), `<a:blip r:embed="rIdN"/>`, an
`<a:srcRect/>` with NO attributes and an `<a:xfrm>` with no attributes.

`has_word_crop` treated ANY srcRect as a crop while `visible_word_image` refuses all-zero
coordinates, so `visible_word_relationship` answered "cropped, no safe pixels" and
`remediate_office._image_bytes_for` returned None for an image with nothing hidden at all:
no vision request, no alt proposal, no preview, no 1.4.5 images-of-text draft, and a retry
reporting no usable draft.

The zero inset alone is NOT enough to release the raster, and that is what most of this
module pins. A picture may carry a zero-inset srcRect and still be flipped, rotated, tiled,
shared between drawings or doubly cropped — placements whose rendered pixels cannot be
predicted from the raster. Those keep supplying no bytes at `visible_word_relationship` and
`_image_bytes_for`, and make no vision call, exactly as before the fix.
"""
import io
import re
import zipfile

import pytest
from PIL import Image
from docx import Document
from lxml import etree as ET

from office_visible_image import NS, has_word_crop, visible_word_image, visible_word_relationship

W_DOC = 'word/document.xml'

# Placements that carry a srcRect this module will not reason about. Every one of them is
# built with an EMPTY or zero-inset srcRect where it can be, so the case under test is
# "the insets are harmless but the placement is not" — the regression this file exists for.
UNSUPPORTED = [
    ('flipped horizontally', dict(xfrm={'flipH': '1'})),
    ('flipped vertically', dict(xfrm={'flipV': '1'})),
    ('rotated', dict(xfrm={'rot': '5400000'})),
    ('tiled fill', dict(tile=True)),
    ('shared between two drawings', dict(shared=True)),
    ('two srcRects on one fill', dict(extra_rect=True)),
    ('unknown attribute', dict(srcrect={'foo': '1'})),
    ('zero inset beside an unknown attribute', dict(srcrect={'l': '0', 'foo': '0'})),
    ('non-integer inset', dict(srcrect={'t': 'oops'})),
    ('fractional inset', dict(srcrect={'t': '0.0'})),
    ('negative inset', dict(srcrect={'l': '-1'})),
    ('out-of-range inset', dict(srcrect={'l': '100000'})),
    ('nothing left visible', dict(srcrect={'l': '50000', 'r': '50000'})),
]
UNSUPPORTED_IDS = [label for label, _ in UNSUPPORTED]
UNSUPPORTED_OPTIONS = [options for _, options in UNSUPPORTED]


def _raster(colour):
    """A PNG at the production image's dimensions, with content a vision model could describe.

    Deliberately not near-solid and not icon-sized: `remediate_office`'s decorative inference
    consults the pixels when an image has no faithful alt source, and a flat block is exactly
    what it proposes erasing — which would take the test off the vision path it is about.
    """
    image = Image.new('RGB', (800, 646), 'white')
    for y in range(120):
        for x in range(800):
            image.putpixel((x, y), colour)
    for y in range(300, 646):
        for x in range(800):
            image.putpixel((x, y), (0, 0, 0) if (x // 40 + y // 40) % 2 else colour)
    out = io.BytesIO()
    image.save(out, format='PNG')
    return out.getvalue()


def _document(srcrect={}, *, name='', xfrm=None, tile=False, pictures=1,
              shared=False, extra_rect=False):
    """The reproduced production placement: unnamed inline drawings, one srcRect each.

    srcrect=None omits the element entirely, which is the baseline a zero-inset srcRect has
    to behave identically to. `pictures` gives each drawing its OWN raster; `shared` instead
    adds a second drawing over IDENTICAL bytes, which python-docx deduplicates into one media
    part and one relationship — the shared placement this module refuses.
    """
    doc = Document()
    for i in range(pictures):
        doc.add_picture(io.BytesIO(_raster((255 - 40 * i, 20 * i, 30 * i))))
    if shared:
        doc.add_picture(io.BytesIO(_raster((255, 0, 0))))
    saved = io.BytesIO()
    doc.save(saved)
    changed = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(saved.getvalue())) as source, \
            zipfile.ZipFile(changed, 'w') as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == W_DOC:
                root = ET.fromstring(data)
                for index, docpr in enumerate(root.xpath('.//wp:docPr', namespaces=NS)):
                    docpr.attrib.clear()
                    docpr.set('id', str(index + 1))
                    docpr.set('name', name)
                    docpr.set('descr', '')
                    docpr.set('title', '')
                for blip in root.xpath('.//a:blip', namespaces=NS):
                    fill = blip.getparent()
                    if tile:
                        ET.SubElement(fill, '{' + NS['a'] + '}tile')
                    for _ in range(2 if extra_rect else 1):
                        if srcrect is None:
                            continue
                        rectangle = ET.SubElement(fill, '{' + NS['a'] + '}srcRect')
                        for key, value in srcrect.items():
                            rectangle.set(key, value)
                for transform in root.xpath('.//a:xfrm', namespaces=NS):
                    for key in ('rot', 'flipH', 'flipV'):
                        transform.attrib.pop(key, None)
                    for key, value in (xfrm or {}).items():
                        transform.set(key, value)
                data = ET.tostring(root)
            target.writestr(entry, data)
    return changed.getvalue()


def _entries(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def _rids(data):
    """Every inline drawing's r:embed id, in document order."""
    root = ET.fromstring(_entries(data)[W_DOC])
    return [blip.get('{' + NS['r'] + '}embed')
            for blip in root.xpath('.//wp:inline//a:blip', namespaces=NS)]


def _descr_by_rid(data):
    """{r:embed id: the docPr descr describing that image}, read from the tree.

    Resolved through the drawing's own structure rather than through `apply_alt`'s
    element-adjacency regex, so this cannot agree with the writer by sharing its mistake.
    """
    root = ET.fromstring(_entries(data)[W_DOC])
    out = {}
    for inline in root.xpath('.//wp:inline', namespaces=NS):
        docpr = inline.xpath('./wp:docPr', namespaces=NS)[0]
        blip = inline.xpath('.//a:blip', namespaces=NS)[0]
        out[blip.get('{' + NS['r'] + '}embed')] = docpr.get('descr')
    return out


def _lookup(data, index=0):
    """`remediate_office._image_bytes_for` for one drawing: (r:embed id, bytes) or None.

    This is the call that actually feeds the vision model, which is why the unsupported
    cases are asserted here and not only on `visible_word_image`.
    """
    import remediate_office
    entries = _entries(data)
    xml = entries[W_DOC].decode('utf-8')
    match = list(re.finditer(r'<wp:docPr\b([^>]*?)(/?)>', xml))[index]
    return remediate_office._image_bytes_for(xml, match, 'wp:docPr', None, entries, W_DOC)


def _answers(data):
    """The four answers that decide whether any pixels reach a model, for one drawing."""
    found = _lookup(data)
    return {'has_word_crop': has_word_crop(data, 'image 1'),
            'visible_word_image': visible_word_image(data, 'image 1'),
            'visible_word_relationship': visible_word_relationship(
                _entries(data), W_DOC, _rids(data)[0]),
            'lookup_bytes': None if found is None else found[1]}


def test_fixture_carries_the_reproduced_production_shape():
    root = ET.fromstring(_entries(_document())[W_DOC])
    assert [dict(e.attrib) for e in root.xpath('.//wp:docPr', namespaces=NS)] == [
        {'id': '1', 'name': '', 'descr': '', 'title': ''}]
    assert [dict(e.attrib) for e in root.xpath('.//a:srcRect', namespaces=NS)] == [{}]
    assert [dict(e.attrib) for e in root.xpath('.//a:xfrm', namespaces=NS)] == [{}]
    assert len(root.xpath('.//wp:inline//a:blip', namespaces=NS)) == 1
    with Image.open(io.BytesIO(_entries(_document())['word/media/image1.png'])) as image:
        assert image.size == (800, 646)


@pytest.mark.parametrize('srcrect', [{}, {'l': '0', 't': '0', 'r': '0', 'b': '0'},
                                     {'b': '0'}, {'l': '0', 'r': '0'}])
def test_a_vouched_srcrect_that_hides_nothing_yields_the_whole_raster(srcrect):
    data = _document(srcrect)
    raw = _entries(data)['word/media/image1.png']
    answers = _answers(data)
    assert answers['has_word_crop'] is False
    # No visible region distinct from the raster, so there is no crop identity to bind —
    # the same answer a placement with no srcRect at all gets.
    assert answers['visible_word_image'] is None
    assert answers['visible_word_relationship'] == (False, None)
    assert answers['lookup_bytes'] == raw        # byte-identical: the reader sees every pixel


def test_a_zero_inset_srcrect_answers_exactly_as_no_srcrect_at_all():
    # The baseline this fix aims at, and the only comparison that licenses using the raster:
    # a plain, unique, untransformed inline picture. Nothing here says a TRANSFORMED picture
    # may be treated that way — see test_unsupported_placements_* for that.
    assert _answers(_document({})) == _answers(_document(None))
    assert _answers(_document({}))['lookup_bytes'] == _entries(_document({}))['word/media/image1.png']


@pytest.mark.parametrize('options', UNSUPPORTED_OPTIONS, ids=UNSUPPORTED_IDS)
def test_unsupported_placements_supply_no_bytes_even_with_a_zero_inset_srcrect(options):
    data = _document(**options)
    answers = _answers(data)
    assert answers['has_word_crop'] is True
    assert answers['visible_word_image'] is None
    assert answers['visible_word_relationship'] == (True, None)
    assert answers['lookup_bytes'] is None       # the call that feeds vision gets nothing


@pytest.mark.parametrize('options', UNSUPPORTED_OPTIONS, ids=UNSUPPORTED_IDS)
def test_unsupported_placements_never_reach_the_vision_model(monkeypatch, options):
    import ai
    import remediate_office
    monkeypatch.setattr(ai, 'vision_is_available', lambda: True)
    monkeypatch.setattr(ai, 'describe_image_structured',
                        lambda *a, **k: pytest.fail('unsupported placement must not call vision'))
    proposals, _ = remediate_office.alt_proposals_for_office(
        _document(**options), 'docx', context_file='unsupported.docx', scan_id='fixture-scan')
    assert not any(p.get('model_call_id') for p in proposals)


@pytest.mark.parametrize('options', UNSUPPORTED_OPTIONS, ids=UNSUPPORTED_IDS)
def test_unsupported_placements_never_ocr_the_full_raster(tmp_path, monkeypatch, options):
    import ocr
    import proposals
    path = tmp_path / 'unsupported.docx'
    path.write_bytes(_document(**options))
    monkeypatch.setattr(ocr, 'is_available', lambda: True)
    monkeypatch.setattr(proposals, 'criteria_enabled', lambda sc: sc == '1.4.5')
    monkeypatch.setattr(ocr, 'ocr_text',
                        lambda *a, **k: pytest.fail('full raster OCR must not run'))
    assert proposals.propose_images_of_text(path, '.docx', ai_enabled=False) == []


def test_a_real_crop_still_supplies_only_the_visible_region():
    data = _document({'t': '19861'})
    raw = _entries(data)['word/media/image1.png']
    assert has_word_crop(data, 'image 1') is True
    rid, image_bytes = _lookup(data)
    assert image_bytes != raw
    with Image.open(io.BytesIO(image_bytes)) as image:
        assert image.size == (800, 517)          # 646 minus the 19.861% hidden band
    assert visible_word_image(data, 'image 1')['crop'] == {'l': 0, 't': 19861, 'r': 0, 'b': 0}


def test_shared_placement_with_a_zero_inset_srcrect_describes_neither_copy(monkeypatch):
    # Identical bytes: python-docx emits ONE media part behind two drawings. Refusing it at
    # visible_word_image was never enough — the lookup is what feeds the model, and a
    # has_word_crop that answered "not cropped" sent the raster for BOTH placements.
    import ai
    import remediate_office
    data = _document(shared=True)
    assert len(set(_rids(data))) == 1 and len(_rids(data)) == 2
    assert _answers(data)['lookup_bytes'] is None
    assert _lookup(data, index=1) is None
    monkeypatch.setattr(ai, 'vision_is_available', lambda: True)
    monkeypatch.setattr(ai, 'describe_image_structured',
                        lambda *a, **k: pytest.fail('a shared picture must not be described'))
    remediate_office.alt_proposals_for_office(data, 'docx', context_file='shared.docx',
                                              scan_id='fixture-scan')


def test_empty_crop_draft_reaches_vision_and_an_approved_alt_lands_on_that_image(monkeypatch):
    import ai
    import apply_alt
    import remediate_office
    original = _document(pictures=2)
    before = bytes(original)
    entries = _entries(original)
    rid_one, rid_two = _rids(original)
    seen = []

    def describe(image, *args, **kwargs):
        seen.append(image)
        return {'alt': f'A striped instruction panel ({len(seen)}).', 'grounded': False,
                'model': 'fixture-model', 'ai_call_id': f'fixture-call-{len(seen)}',
                'source': 'vision'}

    monkeypatch.setattr(ai, 'vision_is_available', lambda: True)
    monkeypatch.setattr(ai, 'describe_image_structured', describe)
    proposals, evidence = remediate_office.alt_proposals_for_office(
        original, 'docx', context_file='empty-crop.docx', scan_id='fixture-scan')

    # The EXACT embedded rasters, not a re-encode and not a crop of them.
    assert seen == [entries['word/media/image1.png'], entries['word/media/image2.png']]
    assert [p['locator'] for p in proposals] == [f'{W_DOC}#{rid_one}', f'{W_DOC}#{rid_two}']
    assert [e['locator'] for e in evidence] == [f'{W_DOC}#{rid_one}', f'{W_DOC}#{rid_two}']
    assert all(p['proposed_value'].startswith('A striped instruction panel') for p in proposals)
    assert proposals[1]['model_call_id'] == 'fixture-call-2'
    assert all(p.get('thumb') for p in proposals)

    # A reviewer approves the SECOND image's draft. The locator the proposer minted must be
    # writable by the existing approved-value writer, onto that image and not the other.
    approved = proposals[1]['proposed_value']
    written, applied, unresolved = apply_alt.apply_alt_text(
        original, {proposals[1]['locator']: approved})
    assert unresolved == []
    assert [a['locator'] for a in applied] == [f'{W_DOC}#{rid_two}']
    assert _descr_by_rid(written) == {rid_one: '', rid_two: approved}
    assert _entries(written)['word/media/image2.png'] == entries['word/media/image2.png']
    assert original == before                  # the source document is never mutated in place


def test_empty_crop_no_longer_suppresses_the_images_of_text_draft(tmp_path, monkeypatch):
    # The same predicate gates the 1.4.5 lane, and the reproduced file still has 1.4.5
    # outstanding: propose_images_of_text skipped any image has_word_crop called cropped but
    # visible_word_image would not supply, so an empty srcRect lost the images-of-text
    # finding as well as the alt draft.
    import ocr
    import proposals
    path = tmp_path / 'empty-crop.docx'
    path.write_bytes(_document())
    monkeypatch.setattr(ocr, 'is_available', lambda: True)
    monkeypatch.setattr(proposals, 'criteria_enabled', lambda sc: sc == '1.4.5')
    seen = []

    def read(image, *args, **kwargs):
        with Image.open(io.BytesIO(image)) as decoded:
            seen.append(decoded.size)
        return 'Visible words only here in this uncropped instruction paragraph today safely'

    monkeypatch.setattr(ocr, 'ocr_text', read)
    result = proposals.propose_images_of_text(path, '.docx', ai_enabled=False)
    assert len(result) == 1
    assert result[0]['sc'] == '1.4.5'
    assert seen and set(seen) == {(800, 646)}       # the whole raster, nothing withheld
    # No crop, so no crop evidence and none of the crop-only caveats: telling a reviewer the
    # draft "transcribes only the visible Word crop" would be false for this picture, and it
    # would withdraw automatic picture replacement from an image that never lost it.
    assert result[0].get('visible_crop') is None
    assert 'visible Word crop' not in result[0]['rationale']
    assert 'Automatic picture replacement remains unavailable' not in result[0]['rationale']


def test_unnamed_images_draft_but_are_absent_from_the_shared_image_walk():
    """The remaining limitation, stated as behaviour rather than as a comment.

    `formats.office.images.undescribed_images` keys an image by its shape NAME and skips
    unnamed elements — "no addressable locator" — so the 1.1.1 detectors and the
    document-wide planner do not list the reproduced picture at all. The drafting lane
    walks `wp:docPr` directly and does reach it, which is why the empty-crop fix restores a
    usable draft; nothing here widens the name-keyed walk.
    """
    import formats.office.images as images
    unnamed = _entries(_document())
    named = _entries(_document(name='Picture 1'))
    assert images.undescribed_images(unnamed) == []
    assert images.undescribed_images(named) == [
        {'locator': f'{W_DOC}#Picture 1', 'part': W_DOC, 'name': 'Picture 1'}]
    # Both still resolve to their pixels, which is what the draft needs.
    assert _lookup(_document())[1] == unnamed['word/media/image1.png']
    assert _lookup(_document(name='Picture 1'))[1] == named['word/media/image1.png']
