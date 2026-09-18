"""ADR 0055 describe-instead-of-replace, proved end to end (WCAG 1.4.5 resolved, 1.1.1 written).

The other half of the pptx image-of-text story. #1715 clears 1.4.5 by DELETING the picture
(tests/test_remediation_verified_pptx_image_of_text.py). This is the lane for the image the
reviewer decides must STAY — the replacement writer refuses it, the styling carries meaning, or
it is a 1.4.9 chart that replacement would gut — where the honest outcome is to describe it.

WHAT MAKES THIS LANE DIFFERENT FROM EVERY OTHER ONE HERE: the criterion the reviewer resolves is
NOT the criterion the write clears. 1.4.5 is resolved by judgement — the picture stays, so it can
never clear on a re-scan — while the description is 1.1.1 content, written by the proven alt lane
and verified against 1.1.1. Getting that split wrong is not a cosmetic error: credited against
1.4.5, the row could never be marked applied, count_unapplied_approved_values would count it
forever, and the file would be permanently unpublishable. `test_the_deck_still_fails_1_4_5_and_
certifies_anyway` is the assertion that keeps the two apart.

WHAT THIS CLAIMS: the picture is still in the package, every placement of it now carries the
reviewer's description as alt text, a REAL re-scan agrees 1.1.1 has cleared, and the file
certifies with 1.4.5 recorded as resolved by judgement rather than fixed. NOT that the deck now
satisfies 1.4.5 — it does not, deliberately, and the reviewer said so.

Nothing but the blob store is patched: handlers._apply_approved_values runs the production seam
through proposals.verify_residual_scs to scanner.analyse_and_assess.
"""
from __future__ import annotations

import functools
import glob
import io
import sys
import tempfile
import zipfile
from pathlib import Path

import pytest
from hitl_viewed import viewed_fields

ACP = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ACP / "api"))

# NO PROVES_LANES DECLARATION, AND THE FILENAME IS PART OF THAT — note it is
# test_remediation_described_… and not test_remediation_verified_…, which is the namespace
# tests/test_capability_assisted_contract.py globs.
#
# That contract derives the applier registry from those modules' declarations and requires each
# (format, criterion) lane to be claimed by exactly ONE, so "pptx 1.1.1 is ASSISTED" always has a
# single findable proof. This module does a full round trip and credits a row — but the lane it
# exercises is ('pptx', '1.1.1'), which test_remediation_verified_pptx_alt.py already proves and
# owns. What is new here is a PATH INTO an existing lane, not a capability: describe-instead-of-
# replace produces ordinary 1.1.1 alt text, written by the same apply_alt_text and verified
# against the same criterion.
#
# So this file stays outside the claiming namespace rather than being admitted to it. Declaring
# the lane twice would tell the matrix it gained something it did not; declaring an empty set
# inside that namespace would make a module that simply FORGOT to declare look deliberate. The
# guard is right as it stands, and a collision with it was information, not an obstacle.
#
# The 1.4.5 half of the decision claims nothing either, and that is the sharper point: it is
# resolved by JUDGEMENT, so no write clears it and no round trip could prove it. #1665 recorded
# that lane as HUMAN and it stays HUMAN.

pytest.importorskip("pptx")
pytest.importorskip("PIL")

FILE = "keepme.pptx"
SID = "rv-pptx-described"

TITLE = "Open enrolment"
BODY = "Unrelated body copy that must survive the write."
LINES = ("Benefits at a glance",
         "Medical dental and vision cover",
         "Enrollment closes on Friday")

# What the reviewer writes: prose ABOUT the picture. Deliberately not the transcript — the
# distinction is the whole reason this lane exists, and _row_approved_values' proposed_value
# fallback is switched off for described rows because of it (store.queue_described_image_alt).
DESCRIPTION = ("A benefits summary card in the company's display face, listing medical, dental "
               "and vision cover, and noting that enrollment closes on Friday.")


def _ocr_ready() -> bool:
    import ocr
    return ocr.is_available()


needs_ocr = pytest.mark.skipif(not _ocr_ready(),
                               reason="tesseract/pytesseract unavailable — 1.4.5 cannot fire")


def _font():
    from PIL import ImageFont
    for pat in ("/usr/share/fonts/**/DejaVuSans.ttf", "/usr/share/fonts/**/DejaVuSans*.ttf",
                "/usr/share/fonts/**/*.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf"):
        hits = glob.glob(pat, recursive=True)
        if hits:
            return ImageFont.truetype(hits[0], 34)
    return ImageFont.load_default()


@functools.lru_cache(maxsize=None)
def _png(lines: tuple[str, ...], size: tuple[int, int] = (760, 260)) -> Path:
    """A picture of prose. Cached for the same reason the 1.4.5 module caches its own: tesseract
    is CPU-bound, the backend suite runs -n auto, and an OCR-heavy file can push OTHER tests'
    reads past ACP_OCR_TIMEOUT_S."""
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(im)
    font = _font()
    for i, line in enumerate(lines):
        d.text((24, 24 + i * 70), line, fill="black", font=font)
    p = Path(tempfile.mkdtemp()) / "text.png"
    im.save(p)
    return p


@functools.lru_cache(maxsize=None)
def _deck(on_slides: int = 2) -> bytes:
    """ONE image of text, placed on `on_slides` slides — one media part, one locator, N pictures.

    Two by default, and that is load-bearing rather than incidental: 'image N' names the media
    part, so a lane that described only the first placement would leave the second carrying its
    source filename and 1.1.1 would still fail. The first draft of resolve_media_locators did
    exactly that, and this deck is what caught it.
    """
    from pptx import Presentation
    from pptx.util import Inches

    img = _png(LINES)
    prs = Presentation()
    for n in range(on_slides):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"{TITLE} {n + 1}"
        slide.shapes.add_picture(str(img), Inches(0.5), Inches(2), Inches(3.5), Inches(1.2))
        box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
        box.text_frame.text = BODY
    out = Path(tempfile.mkdtemp()) / FILE
    prs.save(out)
    return out.read_bytes()


def _spill(data: bytes) -> Path:
    p = Path(tempfile.mkdtemp()) / FILE
    p.write_bytes(data)
    return p


def _assess(data: bytes) -> set[str]:
    """The SCs a REAL assessment reports — the same call the production re-verification makes."""
    from assessment_policy import _extract_sc
    from scanner import analyse_and_assess
    with tempfile.TemporaryDirectory() as d:
        (Path(d) / FILE).write_bytes(data)
        fd, _ = analyse_and_assess(Path(d), FILE, detect_pii=False)
    return {sc for i in (fd or {}).get("issues", []) if (sc := _extract_sc(i.get("wcag", "")))}


def _media(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [n for n in z.namelist() if n.startswith("ppt/media/")]


def _descrs(data: bytes) -> list[str]:
    """Every picture's alt text, read back through python-pptx — a reader with no part in the
    write, so this cannot pass by agreeing with the writer's own idea of the XML."""
    from pptx import Presentation
    prs = Presentation(str(_spill(data)))
    return [sh._element._nvXxPr.cNvPr.get("descr", "")
            for s in prs.slides for sh in s.shapes if sh.shape_type == 13]     # PICTURE


def _texts(data: bytes) -> list[str]:
    from pptx import Presentation
    prs = Presentation(str(_spill(data)))
    return [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]


def _locators(data: bytes) -> list[str]:
    """The locators the real 1.4.5 proposer mints — taken from the proposer, never hand-written,
    so this test breaks if the locator scheme drifts instead of silently testing a fiction."""
    import proposals
    return [p["locator"] for p in proposals.propose_images_of_text(_spill(data), ".pptx")]


class _Blob:
    """The only thing patched in this module. Stores bytes verbatim; decides nothing."""

    def __init__(self, data: bytes):
        self.data, self.uploads = data, []

    def download_remediated(self, owner, sid, f):
        return self.data

    def upload_remediated(self, owner, sid, f, data, mime):
        self.data = data
        self.uploads.append((f, mime))
        return "http://b/2"
    # The approved writer publishes digest-scoped and moves the pointer at commit.
    def upload_immutable_retry(self, owner, sid, f, data, mime):
        return self.upload_remediated(owner, sid, f, data, mime)


@pytest.fixture()
def store(monkeypatch):
    import store as store_mod
    monkeypatch.setattr(store_mod, "_SQLITE_PATH", Path(tempfile.mkdtemp()) / "rvd.db")
    return store_mod.Store()


@pytest.fixture(scope="module")
def deck() -> bytes:
    return _deck()


def _seed(store, deck: bytes, locator: str) -> str:
    """A scanned + remediated deck with one image-of-text card, pending."""
    store.init_scan_run(SID, "drive", 1, "2026-09-07T00:00:00Z", "rubric", "hash")
    store.save_file_result(SID, {
        "file": FILE, "engine": "office", "status": "pass", "score": 60, "compliant": 0,
        "skipped_rules": 0, "drive_file_id": "d1",
        "issues": [{"ruleId": "OCR_IMAGE_OF_TEXT", "wcag": "1.4.5 Images of Text",
                    "severity": "SERIOUS", "detail": f"embedded {locator} contains readable text"}],
    }, "2026-09-07T00:00:00Z")
    store.record_remediation(SID, FILE, drive_write_url="http://d/1", blob_url="http://b/1")
    return store.enqueue_proposals(SID, FILE, "1.4.5", [
        {"locator": locator, "before": "text baked into an image",
         "proposed_value": "Benefits at a glance …", "rationale": "r", "source": "OCR"}],
        rule_name="Images of Text")


def _decide(store, item_id: str, monkeypatch, *, resolution, values):
    """The reviewer's decision, through the PRODUCTION route — not by poking the store.

    The route is where the ADR 0055 wiring lives (validation, and the queue_described_image_alt
    call), so a test that wrote the rows directly would prove the store and skip the feature.
    """
    import core
    from routes.hitl import HitlUpdate, hitl_update
    monkeypatch.setattr(core, "store", store)
    return hitl_update(item_id, HitlUpdate(status="approved", **viewed_fields(item_id), resolution=resolution,
                                           approved_values=values), None)


def _run_lane(monkeypatch, store, blob):
    """The production handler, with the re-scan UNPATCHED."""
    import core
    import handlers
    monkeypatch.setattr(core, "store", store)
    monkeypatch.setitem(sys.modules, "blob", blob)
    handlers._apply_approved_values({"scan_id": SID, "file": FILE}, {})


# ---------------------------------------------------------------- the fixture is honest

@needs_ocr
def test_the_picture_is_really_an_image_of_text(deck):
    """Without this every assertion below is vacuous: if OCR cannot read the fixture, 1.4.5
    never fires and 'the reviewer kept an image of text' describes nothing."""
    assert "1.4.5" in _assess(deck)
    assert len(_media(deck)) == 1                    # one media part…
    assert len(_descrs(deck)) == 2                   # …placed on two slides


@needs_ocr
def test_the_proposer_mints_a_media_index_locator(deck):
    """The locator this lane has to translate, taken from the real proposer. apply_alt cannot
    read it at all, which is the gap ADR 0055's translation fills."""
    import apply_alt
    import apply_pptx_image_of_text as aoit
    locs = _locators(deck)
    assert locs == ["image 1"]
    assert apply_alt.parse_locator(locs[0]) is None
    assert aoit.is_media_index_locator(locs[0])


# ---------------------------------------------------------------- the round trip

@needs_ocr
def test_the_description_reaches_every_placement_and_1_1_1_clears(store, deck, monkeypatch):
    """The whole lane: decide → write → REAL re-scan → credit."""
    locator = _locators(deck)[0]
    item_id = _seed(store, deck, locator)
    _decide(store, item_id, monkeypatch,
            resolution=store.DESCRIBED_RESOLUTION, values=[DESCRIPTION])

    # The decision created the 1.1.1 obligation, approved and unapplied.
    rows = {r["rule_id"]: r for r in store.list_hitl_queue(scan_id=SID)}
    described = rows[f"1.1.1{store.DESCRIBED_RULE_SUFFIX}"]
    assert described["status"] == "approved" and not described.get("applied")
    assert store.approved_alt_values(SID, FILE) == {locator: DESCRIPTION}
    assert store.count_unapplied_approved_values(SID, FILE) == 1
    assert store.mark_file_compliant_if_reviewed(SID, FILE) is False   # nothing written yet

    blob = _Blob(deck)
    _run_lane(monkeypatch, store, blob)

    # Every placement carries the description — not just the first.
    assert _descrs(blob.data) == [DESCRIPTION, DESCRIPTION]
    # The picture STAYED. This is the reviewer's actual decision, and the one thing that
    # distinguishes this lane from the replacement lane.
    assert _media(blob.data) == _media(deck)
    # Unrelated content survived.
    assert BODY in _texts(blob.data)
    # A real re-scan agrees 1.1.1 has cleared, so the row was credited and the file certified.
    assert "1.1.1" not in _assess(blob.data)
    assert store.get_hitl_item(described["id"])["applied"]
    assert store.count_unapplied_approved_values(SID, FILE) == 0


@needs_ocr
def test_the_deck_still_fails_1_4_5_and_certifies_anyway(store, deck, monkeypatch):
    """The split that makes the lane possible, asserted in both directions.

    1.4.5 still FAILS on the written copy — the raster is right there, by the reviewer's own
    decision — and the file certifies regardless, because that finding was resolved by judgement
    and the only content it owed (the 1.1.1 alt text) was written and verified.

    If a future change credits the description against 1.4.5 instead, the first assertion still
    passes and the second turns red: the row could never be marked applied, so the file would be
    permanently unpublishable. That is the failure this test exists to catch.
    """
    locator = _locators(deck)[0]
    item_id = _seed(store, deck, locator)
    _decide(store, item_id, monkeypatch,
            resolution=store.DESCRIBED_RESOLUTION, values=[DESCRIPTION])
    blob = _Blob(deck)
    _run_lane(monkeypatch, store, blob)

    assert "1.4.5" in _assess(blob.data), "the picture was supposed to STAY"
    assert store.get_scan(SID) is not None
    assert store.mark_file_compliant_if_reviewed(SID, FILE) or _already_compliant(store)


def _already_compliant(store) -> bool:
    """mark_file_compliant_if_reviewed is idempotent and returns False once it has fired, so the
    assertion above accepts either 'it certified now' or 'it is already certified'."""
    with store._db.cursor() as cur:
        store._db.execute(cur, "SELECT compliant FROM file_records WHERE scan_id=%s AND file=%s",
                          (SID, FILE))
        return bool((store._db.fetchone(cur) or {}).get("compliant"))


# ---------------------------------------------------------------- the refusals

@needs_ocr
def test_a_described_decision_without_a_description_is_refused(store, deck, monkeypatch):
    """Keeping an image of text and describing it with nothing resolves nothing — it would leave
    the picture unreadable AND mark the finding closed. Refused at the boundary, before any row
    is touched, because after the fact the file would simply certify undescribed."""
    from fastapi import HTTPException
    item_id = _seed(store, deck, _locators(deck)[0])
    with pytest.raises(HTTPException) as e:
        _decide(store, item_id, monkeypatch,
                resolution=store.DESCRIBED_RESOLUTION, values=["   "])
    assert e.value.status_code == 422
    assert store.get_hitl_item(item_id)["status"] == "pending"       # nothing was written


@needs_ocr
def test_described_is_refused_on_a_criterion_it_cannot_mean(store, deck, monkeypatch):
    """'I kept the image and described it' is not an answer to a link-text finding."""
    from fastapi import HTTPException
    _seed(store, deck, _locators(deck)[0])
    other = store.enqueue_proposals(SID, FILE, "2.4.4", [
        {"locator": "https://example.com", "before": "click here",
         "proposed_value": "the benefits guide", "rationale": "r", "source": "ai"}],
        rule_name="Link Purpose")
    with pytest.raises(HTTPException) as e:
        _decide(store, other, monkeypatch,
                resolution=store.DESCRIBED_RESOLUTION, values=["a description"])
    assert e.value.status_code == 422
