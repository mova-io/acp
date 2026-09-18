"""R1 producers: a remediation diff carries `locator`/`page` only when the producer KNOWS them.

What "knows" means here, and what every test below pins:

* A location comes from the object the fixer actually edited — the paragraph it restyled, the
  element it wrote `descr=` on, the slide part it changed, the PDF page whose content stream it
  rewrote — never from the criterion, and never from a count of something else.
* An unknown location is an ABSENT key. Not `None`, not `""`, not a guess: absence is what the
  store and the report read as "not recorded".
* Office producers never emit a page. A Word page is a property of a layout engine, and a paragraph
  index is not one; a slide PART is not a slide number either (reading order lives in
  presentation.xml).
* A PDF page is emitted only when one real page is known: a pass that rewrote two pages writes one
  aggregate record, which has no single page, so it records none.

Synthetic fixtures only (python-docx / reportlab / hand-built XML); no AI, no network.
"""
from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "api"))

import remediate_office as office  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_WORD_TOKEN = re.compile(r"word:(?:p:[1-9]\d*(?::run:[1-9]\d*)?|table:[1-9]\d*:row:[1-9]\d*|document:outline)")


def _assert_no_page(diffs):
    assert all("page" not in d for d in diffs), [d for d in diffs if "page" in d]


def _assert_locator_is_known_or_absent(diffs):
    for d in diffs:
        if "locator" in d:
            assert isinstance(d["locator"], str) and d["locator"].strip(), d


# ── DOCX: structural fixes carry the same token the note already recorded ──────────────────────

def _docx_corrected(tmp_path, document, scope):
    source = tmp_path / "locations.docx"
    document.save(source)
    diffs: list = []
    path, applied, _ = office.remediate_office(source, ai_enabled=False, diffs=diffs,
                                               in_scope=lambda sc: sc in scope)
    assert path and applied
    return path, diffs


def test_docx_structural_diffs_carry_their_exact_word_token_and_never_a_page(tmp_path):
    from docx import Document
    from docx.shared import RGBColor, Pt
    doc = Document()
    doc.add_paragraph("Introduction text")
    doc.add_paragraph("First section", style="Heading 2")
    p = doc.add_paragraph()
    p.add_run("Ordinary black text ")
    p.add_run("Light text").font.color.rgb = RGBColor.from_string("CCCCCC")
    body = doc.add_paragraph("Body paragraph with normal sized content")
    pseudo = doc.add_paragraph().add_run("Section Overview")
    pseudo.bold, pseudo.font.size = True, Pt(22)
    for heading in ["First table", "Second table"]:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = heading
        table.cell(1, 0).text = "Value"
    assert body is not None
    _, diffs = _docx_corrected(tmp_path, doc, {"1.3.1", "1.4.3", "2.4.2", "3.1.1"})

    structural = [d for d in diffs if d["rule_id"] in {"1.3.1", "1.4.3"}]
    assert structural
    for d in structural:
        token = re.search(r"\[location:([^\]]+)\]$", d["note"])
        # The locator IS the note's token — the same object, now in a field that no note
        # truncation or display filter can lose.
        assert token and d.get("locator") == token.group(1), d
        assert _WORD_TOKEN.fullmatch(d["locator"])
    locators = sorted(d["locator"] for d in structural)
    assert "word:table:1:row:1" in locators and "word:table:2:row:1" in locators
    assert "word:p:3:run:2" in locators                       # the light run, not the paragraph
    # Document properties are the whole document, not a place in it.
    metadata = [d for d in diffs if d["rule_id"] in {"2.4.2", "3.1.1"}]
    assert metadata and all("locator" not in d for d in metadata)
    _assert_no_page(diffs)


def test_docx_outline_normalisation_records_the_outline_scope_not_one_paragraph(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Title", style="Heading 1")
    doc.add_paragraph("Section", style="Heading 1")
    _, diffs = _docx_corrected(tmp_path, doc, {"1.3.1"})
    diff = next(d for d in diffs if "demoted 1 to Heading 2" in d["after"])
    assert diff["locator"] == "word:document:outline"
    _assert_no_page(diffs)


def _docx_body(body: str) -> dict:
    xml = f'<w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>'
    return {"word/document.xml": xml.encode("utf-8")}


def _sdt(kind="date"):
    return (f"<w:sdt><w:sdtPr><w:{kind}/></w:sdtPr>"
            "<w:sdtContent><w:r><w:t> </w:t></w:r></w:sdtContent></w:sdt>")


def test_docx_inline_form_label_is_located_by_its_containing_paragraph():
    entries = _docx_body("<w:p><w:r><w:t>Intro</w:t></w:r></w:p>"
                         "<w:p><w:r><w:t>Date of birth:</w:t></w:r>" + _sdt() + "</w:p>")
    diffs: list = []
    office._remediate_docx_structure(entries, diffs, [], lambda sc: sc in {"3.3.2", "4.1.2"})
    labels = [d for d in diffs if d["rule_id"] in {"3.3.2", "4.1.2"}]
    assert {d["rule_id"] for d in labels} == {"3.3.2", "4.1.2"}
    assert {d.get("locator") for d in labels} == {"word:p:2"}
    _assert_no_page(diffs)


# ── DOCX images: `part#fragment` only when it resolves back to the element written ─────────────

def _png(colour):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (96, 72), colour).save(buf, format="PNG")
    return buf.getvalue()


def _docx_images(tmp_path, pictures, names, titles, filename="images.docx") -> Path:
    from docx import Document
    doc = Document()
    for picture in pictures:
        doc.add_picture(io.BytesIO(picture))
    raw = io.BytesIO()
    doc.save(raw)
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    xml = entries["word/document.xml"].decode()
    found = list(re.finditer(r"<wp:docPr\b[^>]*?/?>", xml))
    assert len(found) == len(pictures)
    out, last = [], 0
    for m, name, title in zip(found, names, titles):
        tag = m.group(0)
        tag = re.sub(r'\sname="[^"]*"', f' name="{name}"', tag)
        tag = re.sub(r'\sdescr="[^"]*"', "", tag)
        tag = tag.replace("<wp:docPr", f'<wp:docPr title="{title}"', 1)
        out.append(xml[last:m.start()] + tag)
        last = m.end()
    entries["word/document.xml"] = ("".join(out) + xml[last:]).encode()
    path = tmp_path / filename
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in entries.items():
            z.writestr(n, data)
    return path


def _alt_diffs(path):
    diffs: list = []
    out, applied, _ = office.remediate_office(path, ai_enabled=False, diffs=diffs,
                                              in_scope=lambda sc: sc == "1.1.1")
    assert out and applied
    return out, [d for d in diffs if d["rule_id"] == "1.1.1"]


def _descr_at(data: bytes, locator: str) -> str:
    """The descr of the element `locator` resolves to — via the approved-alt writer's own resolver."""
    from apply_alt import parse_locator, resolve_target, tag_for_part
    part, fragment = parse_locator(locator)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read(part).decode()
    at = resolve_target(xml, tag_for_part(part), fragment)
    assert at is not None, locator
    return re.search(r'\bdescr="([^"]*)"', xml[at:xml.index(">", at)]).group(1)


def test_docx_faithful_alt_diffs_name_the_element_that_was_written(tmp_path):
    path = _docx_images(tmp_path, [_png("blue"), _png("red")],
                        ["Revenue chart", "Staff photo"],
                        ["Quarterly revenue by region", "The support team at the 2026 offsite"])
    out, diffs = _alt_diffs(path)
    assert [d["after"] for d in diffs] == ["Quarterly revenue by region",
                                           "The support team at the 2026 offsite"]
    assert [d["locator"] for d in diffs] == ["word/document.xml#Revenue chart",
                                             "word/document.xml#Staff photo"]
    data = Path(out).read_bytes()
    for d in diffs:
        assert _descr_at(data, d["locator"]) == d["after"]
    _assert_no_page(diffs)


def test_duplicate_shape_name_falls_back_to_the_image_relationship(tmp_path):
    path = _docx_images(tmp_path, [_png("blue"), _png("red")], ["Chart", "Chart"],
                        ["First chart description", "Second chart description"])
    out, diffs = _alt_diffs(path)
    first, second = diffs
    assert first["locator"] == "word/document.xml#Chart"      # the name finds this one first
    # The second image's name resolves to the FIRST image, so it is refused; its own image
    # relationship is unique and resolves to it.
    assert re.fullmatch(r"word/document\.xml#rId\w+", second["locator"])
    data = Path(out).read_bytes()
    assert _descr_at(data, first["locator"]) == "First chart description"
    assert _descr_at(data, second["locator"]) == "Second chart description"


def test_no_fragment_that_resolves_to_the_element_means_no_locator(tmp_path):
    # Same pixels twice → one shared image part and relationship, and one shared name: no
    # fragment addresses the SECOND element alone, so its location is unknown and absent.
    same = _png("green")
    path = _docx_images(tmp_path, [same, same], ["Logo", "Logo"],
                        ["Company logo on the cover", "Company logo in the footer"])
    _, diffs = _alt_diffs(path)
    first, second = diffs
    assert first["locator"] == "word/document.xml#Logo"
    assert "locator" not in second
    _assert_no_page(diffs)


# ── PPTX / XLSX: the edited part, never a slide number ─────────────────────────────────────────

_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _slide_xml():
    low = ('<p:sp><p:nvSpPr/><p:spPr><a:xfrm><a:off x="0" y="900"/><a:ext cx="1" cy="1"/></a:xfrm>'
           '<a:solidFill><a:srgbClr val="F07D00"/></a:solidFill></p:spPr>'
           '<p:txBody><a:p><a:r><a:rPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr>'
           '<a:t>Low contrast label</a:t></a:r></a:p></p:txBody></p:sp>')
    top = ('<p:sp><p:nvSpPr/><p:spPr><a:xfrm><a:off x="0" y="10"/><a:ext cx="1" cy="1"/></a:xfrm>'
           '</p:spPr><p:txBody><a:p><a:r><a:rPr/><a:t>Top text</a:t></a:r></a:p></p:txBody></p:sp>')
    rows = "".join(f'<a:tr h="1"><a:tc><a:txBody><a:p><a:r><a:t>r{i}</a:t></a:r></a:p></a:txBody></a:tc></a:tr>'
                   for i in range(2))
    table = ('<p:graphicFrame><a:graphic><a:graphicData><a:tbl><a:tblPr firstRow="0"/>'
             f'<a:tblGrid><a:gridCol w="1"/></a:tblGrid>{rows}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>')
    return (f'<?xml version="1.0"?><p:sld xmlns:p="{_P}" xmlns:a="{_A}"><p:cSld><p:spTree>'
            f'<p:nvGrpSpPr/><p:grpSpPr/>{low}{top}{table}</p:spTree></p:cSld></p:sld>')


def test_pptx_slide_fixes_name_the_slide_part_they_edited():
    # Only slide7.xml exists: a producer that numbered slides by position would say "slide 1".
    entries = {"ppt/slides/slide7.xml": _slide_xml().encode()}
    diffs: list = []
    applied = office._remediate_pptx_slides(entries, diffs)
    assert applied
    assert {d["rule_id"] for d in diffs} >= {"2.4.2", "1.4.3", "1.3.1", "1.3.2"}
    assert {d.get("locator") for d in diffs} == {"ppt/slides/slide7.xml"}
    _assert_no_page(diffs)


def test_pptx_table_header_without_a_part_records_no_location():
    diffs: list = []
    _, n = office._pptx_mark_table_headers(_slide_xml(), diffs)
    assert n == 1 and "locator" not in diffs[0]


def test_xlsx_table_part_is_recorded_but_an_aggregate_unhide_is_not():
    ss = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    entries = {
        "xl/tables/table3.xml": f'<table xmlns="{ss}" id="3" ref="A1:B3" headerRowCount="0"/>'.encode(),
        "xl/worksheets/sheet1.xml": (f'<worksheet xmlns="{ss}"><sheetData>'
                                     '<row r="2" hidden="1"><c r="A2"><v>5</v></c></row>'
                                     '</sheetData></worksheet>').encode(),
    }
    diffs: list = []
    office._remediate_xlsx_structure(entries, diffs)
    table = next(d for d in diffs if d["rule_id"] == "1.3.1")
    hidden = next(d for d in diffs if d["rule_id"] == "1.3.2")
    assert table["locator"] == "xl/tables/table3.xml"
    assert "locator" not in hidden                # one record for every unhidden line in the file
    _assert_no_page(diffs)


def test_rec_never_stores_an_empty_or_non_string_locator():
    diffs: list = []
    for value in (None, "", "   ", 3, ["word:p:1"]):
        office._rec(diffs, "1.3.1", "b", "a", "n", locator=value)
    assert all(set(d) == {"rule_id", "before", "after", "note"} for d in diffs)
    assert office._part_locator("word/document.xml") is None
    assert office._part_locator("ppt/slides/slide12.xml") == "ppt/slides/slide12.xml"


# ── PDF: a page only when one real page is known ───────────────────────────────────────────────

pikepdf = pytest.importorskip("pikepdf")
pytest.importorskip("reportlab")

import remediate_pdf as rp  # noqa: E402
from engines import NO_PDF, PDF_OK  # noqa: E402


def _light_text_pdf(path: Path, pages: int) -> Path:
    from reportlab.lib.colors import Color
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path))
    for i in range(pages):
        c.setFillColor(Color(0.8, 0.8, 0.8))
        c.drawString(72, 650, f"Light grey text on page {i + 1}")
        c.showPage()
    c.save()
    return path


def _form_pdf(path: Path, field_pages: list[int], pages: int = 2, generic=False) -> Path:
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path))
    for page in range(1, pages + 1):
        if page in field_pages:
            name = f"Text{page}" if generic else f"Home Address {page}"
            c.acroForm.textfield(name=name, x=72, y=700, width=200, height=20, borderWidth=1)
        else:
            c.drawString(72, 720, "no fields here")
        c.showPage()
    c.save()
    return path


@pytest.mark.skipif(not PDF_OK, reason=NO_PDF)
def test_single_page_recolour_records_that_page_and_metadata_records_none(tmp_path):
    src = _light_text_pdf(tmp_path / "one.pdf", 1)
    diffs: list = []
    out, applied, _ = rp.remediate_pdf(src, ai_enabled=False, diffs=diffs,
                                       in_scope=lambda sc: sc in {"1.4.3", "3.1.1", "2.4.2"})
    assert out
    contrast = next(d for d in diffs if d["rule_id"] == "1.4.3")
    assert contrast["page"] == 1
    document_level = [d for d in diffs if d["rule_id"] in {"3.1.1", "2.4.2"}]
    assert document_level and all("page" not in d and "locator" not in d for d in document_level)


@pytest.mark.skipif(not PDF_OK, reason=NO_PDF)
def test_a_recolour_that_spans_pages_records_no_single_page(tmp_path):
    src = _light_text_pdf(tmp_path / "two.pdf", 2)
    diffs: list = []
    rp.remediate_pdf(src, ai_enabled=False, diffs=diffs, in_scope=lambda sc: sc == "1.4.3")
    contrast = next(d for d in diffs if d["rule_id"] == "1.4.3")
    assert "page" not in contrast


def test_recolour_reports_only_pages_it_actually_rewrote(tmp_path):
    from reportlab.lib.colors import Color
    from reportlab.pdfgen import canvas
    src = tmp_path / "mixed.pdf"
    c = canvas.Canvas(str(src))
    c.drawString(72, 650, "Black text passes")
    c.showPage()
    c.setFillColor(Color(0.8, 0.8, 0.8))
    c.drawString(72, 650, "Light grey text fails")
    c.showPage()
    c.save()
    pages: list = []
    with pikepdf.open(str(src)) as pdf:
        assert rp._fix_pdf_text_contrast(pdf, pages=pages) > 0
    assert pages == [2]


@pytest.mark.skipif(not PDF_OK, reason=NO_PDF)
def test_tab_order_records_the_one_page_whose_tabs_it_set(tmp_path):
    src = _form_pdf(tmp_path / "form.pdf", field_pages=[2])
    diffs: list = []
    rp.remediate_pdf(src, ai_enabled=False, diffs=diffs, in_scope=lambda sc: sc == "2.4.3")
    tabs = next(d for d in diffs if d["rule_id"] == "2.4.3")
    assert tabs["page"] == 2                    # page 2, not a default of 1


@pytest.mark.skipif(not PDF_OK, reason=NO_PDF)
def test_tab_order_across_pages_records_no_page(tmp_path):
    src = _form_pdf(tmp_path / "form2.pdf", field_pages=[1, 2])
    diffs: list = []
    rp.remediate_pdf(src, ai_enabled=False, diffs=diffs, in_scope=lambda sc: sc == "2.4.3")
    tabs = next(d for d in diffs if d["rule_id"] == "2.4.3")
    assert "page" not in tabs


def _figure_pdf(path: Path, *, page_index: int | None) -> Path:
    from reportlab.pdfgen import canvas
    raw = path.with_name("raw-" + path.name)
    c = canvas.Canvas(str(raw))
    for i in range(2):
        c.drawString(72, 720, f"Body text {i + 1}")
        c.rect(120, 480, 200, 30, fill=1)
        c.showPage()
    c.save()
    pdf = pikepdf.open(str(raw))
    fig = pikepdf.Dictionary(Type=pikepdf.Name("/StructElem"), S=pikepdf.Name("/Figure"), K=0)
    if page_index is not None:
        fig.Pg = pdf.pages[page_index].obj
    fig = pdf.make_indirect(fig)
    doc = pdf.make_indirect(pikepdf.Dictionary(Type=pikepdf.Name("/StructElem"), S=pikepdf.Name("/Document"),
                                               K=pikepdf.Array([fig])))
    fig.P = doc
    root = pdf.make_indirect(pikepdf.Dictionary(Type=pikepdf.Name("/StructTreeRoot"), K=pikepdf.Array([doc])))
    doc.P = root
    pdf.Root.StructTreeRoot = root
    pdf.Root.MarkInfo = pikepdf.Dictionary(Marked=True)
    pdf.save(str(path))
    pdf.close()
    return path


def test_approved_figure_alt_row_carries_the_figures_own_page(tmp_path):
    src = _figure_pdf(tmp_path / "fig.pdf", page_index=1)
    _, applied, unresolved = rp.apply_pdf_figure_alt(src.read_bytes(), {"pdf:fig:2:0": "Sales by quarter"})
    assert not unresolved
    assert applied[0]["locator"] == "pdf:fig:2:0" and applied[0]["page"] == 2


def test_approved_figure_alt_without_a_page_reference_records_no_page(tmp_path):
    src = _figure_pdf(tmp_path / "nopg.pdf", page_index=None)
    _, applied, unresolved = rp.apply_pdf_figure_alt(src.read_bytes(), {"pdf:fig:?:0": "Sales by quarter"})
    assert not unresolved and applied
    assert "page" not in applied[0]


def test_approved_field_name_row_carries_the_fields_own_page(tmp_path):
    src = _form_pdf(tmp_path / "fields.pdf", field_pages=[2], generic=True)
    _, applied, unresolved = rp.apply_pdf_field_name(src.read_bytes(), {"pdf:field:2:0": "Home address"})
    assert not unresolved
    assert applied[0]["page"] == 2 and applied[0]["locator"] == "pdf:field:2:0"


# ── unverified_changes: the reverify path keeps what the writer recorded ───────────────────────

def test_known_location_copies_only_real_values():
    from unverified_changes import _known_location
    assert _known_location({"locator": "pdf:fig:2:0", "page": 2}) == {"locator": "pdf:fig:2:0", "page": 2}
    for bad in ({"locator": ""}, {"locator": None}, {"page": 0}, {"page": -1}, {"page": True},
                {"page": "2"}, {"page": 2.0}, {}):
        assert _known_location(bad) == {}
    # A page is never parsed out of locator text; the writer has to have recorded it.
    assert _known_location({"locator": "pdf:fig:2:0"}) == {"locator": "pdf:fig:2:0"}


def test_rerecorded_rows_keep_stored_locations_and_drop_reconstructed_ones():
    from unverified_changes import _as_recorded
    stored = {"rule_id": "1.1.1", "before": "", "after": "x", "note": "n", "locator": "word:p:4", "seq": 0}
    assert _as_recorded(stored)["locator"] == "word:p:4"
    assert _as_recorded({**stored, "location_source": "recorded", "page": 3})["page"] == 3
    legacy = _as_recorded({**stored, "location_source": "legacy_note", "page": 3})
    assert "locator" not in legacy and "page" not in legacy
    assert legacy["note"] == "n"                  # the note (the legacy evidence) is untouched


@pytest.fixture
def store(monkeypatch, tmp_path):
    import store as st
    monkeypatch.setattr(st, "_SQLITE_PATH", tmp_path / "locations.db")
    return st.Store()


def _seed_saved_unverified(store, changes):
    from hashlib import sha256
    from test_remediation_contribution import seed_exact_writer
    import remediation_contribution as c
    _run, item = seed_exact_writer(store)
    data = b"exact saved corrected fixture"
    artifact = sha256(data).hexdigest()
    with store._db.cursor() as cur:
        store._db.execute(cur, "UPDATE file_records SET corrected_sha256=%s WHERE scan_id='scan'", (artifact,))
        store._db.execute(cur, "UPDATE hitl_queue SET applied=1 WHERE id=%s", (item,))
    store.log_decision("system", "apply.saved_unverified", scan_id="scan", file="a.docx", rule_id="1.1.1",
                       detail=c.encoded(dict(artifact_sha256=artifact, source_sha256="a" * 64, item_ids=[item],
                                             baseline_residual=["1.1.1"], changes=changes,
                                             verification="not_verified")))
    return data


def test_reverified_write_passes_the_writers_locator_to_the_store(store, monkeypatch):
    from proposals import Verification
    from unverified_changes import record_verification
    data = _seed_saved_unverified(store, [dict(locator="image", before="", after="A tree")])
    recorded = []
    real = store.record_remediation_diffs
    monkeypatch.setattr(store, "record_remediation_diffs",
                        lambda scan_id, file, diffs: (recorded.append([dict(d) for d in diffs]),
                                                      real(scan_id, file, diffs))[1])
    assert record_verification(store, "scan", "a.docx", data, Verification(True, set())) == 1
    (entries,) = recorded
    (entry,) = entries
    assert entry["locator"] == "image"
    assert "page" not in entry                                 # a Word writer recorded none
    assert entry["note"] == "AI applied; exact saved copy subsequently verified · image"


# ── the real store round trip (Stream D's remediation_diff.locator/page columns) ───────────────

def test_reverify_round_trip_keeps_stored_locations_and_leaves_legacy_ones_legacy(store):
    from proposals import Verification
    from unverified_changes import record_verification
    data = _seed_saved_unverified(store, [dict(locator="image", before="", after="A tree")])
    # Two rows recorded earlier: one whose writer stored its location, one from before the
    # columns existed whose location only survives in its reviewer note.
    store.record_remediation_diffs("scan", "a.docx", [
        {"rule_id": "1.3.1", "before": "b", "after": "a", "note": "n", "locator": "word:p:4"},
        {"rule_id": "2.4.4", "before": "b", "after": "a",
         "note": "approved by a reviewer · word/document.xml#rIdLegacy"},
    ])
    assert record_verification(store, "scan", "a.docx", data, Verification(True, set())) == 1
    rows = {r["rule_id"]: r for r in store.get_remediation_diffs("scan", "a.docx")}
    assert (rows["1.1.1"]["locator"], rows["1.1.1"]["page"], rows["1.1.1"]["location_source"]) == (
        "image", None, "recorded")
    assert (rows["1.3.1"]["locator"], rows["1.3.1"]["location_source"]) == ("word:p:4", "recorded")
    # Re-recording did not promote the note-derived location into a stored one.
    assert (rows["2.4.4"]["locator"], rows["2.4.4"]["location_source"]) == (
        "word/document.xml#rIdLegacy", "legacy_note")
    with store._db.cursor() as cur:
        store._db.execute(cur, "SELECT locator FROM remediation_diff WHERE scan_id='scan' AND rule_id='2.4.4'")
        assert store._db.fetchall(cur)[0]["locator"] is None


def test_office_producer_locations_round_trip_and_no_word_page_is_stored(tmp_path, store):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Title", style="Heading 2")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text, table.cell(1, 0).text = "Heading", "Value"
    _, diffs = _docx_corrected(tmp_path, doc, {"1.3.1", "2.4.2"})
    store.record_remediation_diffs("scan", "a.docx", diffs)
    rows = store.get_remediation_diffs("scan", "a.docx")
    assert sorted(r["locator"] for r in rows if r["rule_id"] == "1.3.1") == ["word:p:1", "word:table:1:row:1"]
    title = next(r for r in rows if r["rule_id"] == "2.4.2")
    assert title["locator"] is None and title["location_source"] is None
    assert all(r["page"] is None for r in rows)


@pytest.mark.skipif(not PDF_OK, reason=NO_PDF)
def test_pdf_producer_page_round_trips_through_the_store(tmp_path, store):
    src = _light_text_pdf(tmp_path / "one.pdf", 1)
    diffs: list = []
    rp.remediate_pdf(src, ai_enabled=False, diffs=diffs, in_scope=lambda sc: sc in {"1.4.3", "3.1.1"})
    store.record_remediation_diffs("scan", "one.pdf", diffs)
    rows = {r["rule_id"]: r for r in store.get_remediation_diffs("scan", "one.pdf")}
    assert (rows["1.4.3"]["page"], rows["1.4.3"]["locator"], rows["1.4.3"]["location_source"]) == (
        1, None, "recorded")
    assert rows["3.1.1"]["page"] is None and rows["3.1.1"]["location_source"] is None
