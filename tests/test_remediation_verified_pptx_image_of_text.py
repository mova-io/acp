"""The 1.4.5 pptx image-of-text lane, proved end to end (WCAG 1.4.5 Images of Text).

Same bar as every lane before it: the original deck trips the finding, an approval changes the
saved deck, a REAL re-scan verifies it, unrelated content survives, and a broken engine earns no
credit. Nothing but the blob store is patched — `handlers._apply_approved_values` runs the
production seam through `proposals.verify_residual_scs` to `scanner.analyse_and_assess`.

WHAT MAKES THIS LANE DIFFERENT FROM EVERY OTHER ONE HERE: it DELETES content. The other lanes
add a description, a name, a language mark. 1.4.5 asks for real text INSTEAD of a picture of
text, so the fix is to replace the picture with a text box and remove the image. #1665
downgraded the old descr-writing lane to HUMAN because writing alt text left the raster in place
and the criterion kept failing; this is the fix that actually clears it.

THE TRAP, MEASURED BEFORE THE WRITER WAS WRITTEN, and pinned below by
`test_removing_only_the_picture_does_not_clear_it`: `ocr._ooxml_images` walks the ZIP NAMELIST
for `ppt/media/*` rasters and never opens a slide. Deleting the `<p:pic>` element therefore does
NOT clear the finding — tesseract still reads the orphaned media part and the re-scan re-fires.
Only removing the media part clears it. That test fails if a future refactor "tidies up" the
writer into leaving the bytes behind, which would silently restore the exact defect #1665 found.

WHY THE APPROVED VALUE IS THE PROPOSER'S OWN DRAFT HERE. Unlike the vision lanes, this
proposer is deterministic — `propose_images_of_text` returns the tesseract transcript — so the
test drives the whole chain including the draft, and the value written is the value a reviewer
would actually see on the card.

WHAT THIS CLAIMS: the words are real text now, selectable and announced by a screen reader, and
the raster is gone from the package. NOT that the slide looks the same, and not that the
transcript is perfect — OCR misreads things, which is exactly why a human approves each one and
why the lane is `assisted` rather than `auto`.
"""
from __future__ import annotations

import functools
import glob
import io
import sys
import tempfile
import zipfile
from pathlib import Path

import pytest
from hitl_viewed import approve_bound

ACP = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ACP / "api"))

# The (format, criterion) lanes this module PROVES end to end — read by
# tests/test_capability_assisted_contract.py, which derives the applier registry from
# these declarations instead of a hand-written list. A literal set, so it can be read
# without importing this module.
PROVES_LANES = {("pptx", "1.4.5")}

pytest.importorskip("pptx")
pytest.importorskip("PIL")

FILE = "benefits.pptx"
SID = "rv-pptx-145"

TITLE = "Open enrolment"
BODY = "Unrelated body copy that must survive the write."
# Rendered into the picture, and therefore also what OCR should read back out. Deliberately
# more than _MIN_WORDS (10) real words and no numerals, so the 1.4.5 band takes it and
# _looks_like_chart does not mistake it for a chart.
LINES = ("Benefits at a glance",
         "Medical dental and vision cover",
         "Enrollment closes on Friday")


def _ocr_ready() -> bool:
    import ocr
    return ocr.is_available()


needs_ocr = pytest.mark.skipif(not _ocr_ready(),
                               reason="tesseract/pytesseract unavailable — 1.4.5 cannot fire")


def _font():
    from PIL import ImageFont
    for pat in ("/usr/share/fonts/**/DejaVuSans.ttf", "/usr/share/fonts/**/DejaVuSans*.ttf",
                "/usr/share/fonts/**/*.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf"):
        hits = glob.glob(pat, recursive=True)
        if hits:
            return ImageFont.truetype(hits[0], 34)
    return ImageFont.load_default()


@functools.lru_cache(maxsize=None)
def _png(lines: tuple[str, ...], size: tuple[int, int] = (760, 260)) -> Path:
    """A picture of prose. Big enough to clear _MIN_PIXELS (20000) and legible enough that
    tesseract reads the words back — a fixture OCR cannot read would make every assertion
    below vacuous, which `test_the_picture_is_really_an_image_of_text` guards.

    CACHED, and the decks below are too. Tesseract is CPU-bound and the backend suite runs
    `-n auto` on a small runner, so an OCR-heavy file does not just cost its own time: it can
    push OTHER tests' reads past ACP_OCR_TIMEOUT_S, which surfaces as an OCR_IMAGE_UNREAD
    advisory in a test that was asserting a clean read. Observed exactly that on a loaded box.
    Building each distinct image and deck once keeps this file's contribution bounded.
    """
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(im)
    font = _font()
    for i, line in enumerate(lines):
        d.text((24, 24 + i * 70), line, fill="black", font=font)
    p = Path(tempfile.mkdtemp()) / "text.png"
    im.save(p)
    return p


@functools.lru_cache(maxsize=None)
def _deck(*, pictures: int = 1, on_slides: int = 1, titled: bool = True) -> bytes:
    """`on_slides` slides, each carrying `pictures` copies of the SAME image-of-text.

    One image reused across slides is one media part and one locator — which is why the lane
    replaces every placement or none: the part is deleted once.
    """
    from pptx import Presentation
    from pptx.util import Inches

    img = _png(LINES)
    prs = Presentation()
    for n in range(on_slides):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        if titled:
            slide.shapes.title.text = f"{TITLE} {n + 1}"
        for k in range(pictures):
            slide.shapes.add_picture(str(img), Inches(0.5 + 4 * k), Inches(2), Inches(3.5), Inches(1.2))
        box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
        box.text_frame.text = BODY
    out = Path(tempfile.mkdtemp()) / FILE
    prs.save(out)
    return out.read_bytes()


@functools.lru_cache(maxsize=None)
def _two_images_deck() -> bytes:
    """Two DIFFERENT images of text, so a one-of-two approval leaves the criterion failing."""
    from pptx import Presentation
    from pptx.util import Inches

    a = _png(LINES)
    b = _png(("Contact the benefits team", "Questions answered every Tuesday",
              "Ask about dependent cover"))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = TITLE
    slide.shapes.add_picture(str(a), Inches(0.5), Inches(2), Inches(3.5), Inches(1.2))
    slide.shapes.add_picture(str(b), Inches(4.5), Inches(2), Inches(3.5), Inches(1.2))
    out = Path(tempfile.mkdtemp()) / FILE
    prs.save(out)
    return out.read_bytes()


def _spill(data: bytes) -> Path:
    p = Path(tempfile.mkdtemp()) / FILE
    p.write_bytes(data)
    return p


def _assess(data: bytes) -> set[str]:
    """The SCs a REAL assessment reports — the same call the production re-verification makes."""
    from assessment_policy import _extract_sc
    from scanner import analyse_and_assess
    with tempfile.TemporaryDirectory() as d:
        (Path(d) / FILE).write_bytes(data)
        fd, _ = analyse_and_assess(Path(d), FILE, detect_pii=False)
    return {sc for i in (fd or {}).get("issues", []) if (sc := _extract_sc(i.get("wcag", "")))}


def _first_party_145(data: bytes) -> list[str]:
    """The 1.4.5 findings the OCR detector itself reports, asked directly rather than through a
    full scan — the .NET analyser, where built, adds its own findings and a scan-level count
    would depend on which engines happen to be installed."""
    import ocr
    return [f["ruleId"] for f in ocr.images_of_text(_spill(data), ".pptx")
            if f.get("ruleId") == "OCR_IMAGE_OF_TEXT"]


def _media(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [n for n in z.namelist() if n.startswith("ppt/media/")]


def _texts(data: bytes) -> list[str]:
    """Every shape's text, read back by python-pptx — a reader with no part in the write."""
    from pptx import Presentation
    prs = Presentation(str(_spill(data)))
    return [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]


def _proposed(data: bytes) -> dict[str, str]:
    """{locator: transcript} exactly as the reviewer sees it on the card."""
    import proposals
    return {p["locator"]: p["proposed_value"]
            for p in proposals.propose_images_of_text(_spill(data), ".pptx")}


class _Blob:
    """The only thing patched in this module. Stores bytes verbatim; decides nothing."""

    def __init__(self, data: bytes):
        self.data, self.uploads = data, []

    def download_remediated(self, owner, sid, f):
        return self.data

    def upload_remediated(self, owner, sid, f, data, mime):
        self.data = data
        self.uploads.append((f, mime))
        return "http://b/2"
    # The approved writer publishes digest-scoped and moves the pointer at commit.
    def upload_immutable_retry(self, owner, sid, f, data, mime):
        return self.upload_remediated(owner, sid, f, data, mime)


@pytest.fixture()
def store(monkeypatch):
    import store as store_mod
    monkeypatch.setattr(store_mod, "_SQLITE_PATH", Path(tempfile.mkdtemp()) / "rv.db")
    return store_mod.Store()


def _seed(store, values: dict[str, str], rule_id: str = "1.4.5") -> str:
    """A scanned + remediated deck with one image-of-text card, and `values` approved on it."""
    store.init_scan_run(SID, "drive", 1, "2026-09-07T00:00:00Z", "rubric", "hash")
    store.save_file_result(SID, {
        "file": FILE, "engine": "office", "status": "pass", "score": 60, "compliant": 0,
        "skipped_rules": 0, "drive_file_id": "d1",
        "issues": [{"ruleId": "OCR_IMAGE_OF_TEXT", "wcag": "1.4.5 Images of Text",
                    "severity": "SERIOUS", "detail": f"embedded {loc} contains readable text"}
                   for loc in values],
    }, "2026-09-07T00:00:00Z")
    store.record_remediation(SID, FILE, drive_write_url="http://d/1", blob_url="http://b/1")
    item_id = store.enqueue_proposals(SID, FILE, rule_id, [
        {"locator": loc, "before": "text baked into an image", "proposed_value": "",
         "rationale": "r", "source": "OCR"} for loc in values], rule_name="Images of Text")
    approve_bound(store, item_id, list(values.values()))
    return item_id


def _run_lane(monkeypatch, store, blob):
    """The production handler, with the re-scan UNPATCHED."""
    import core
    import handlers
    monkeypatch.setattr(core, "store", store)
    monkeypatch.setitem(sys.modules, "blob", blob)
    handlers._apply_approved_values({"scan_id": SID, "file": FILE}, {})


@pytest.fixture(scope="module")
def deck() -> bytes:
    return _deck()


# ── 1. the finding, and that the fixture is really what it claims ─────────────

@needs_ocr
def test_a_real_assessment_reports_1_4_5_on_a_picture_of_prose(deck):
    assert "1.4.5" in _assess(deck)


@needs_ocr
def test_the_picture_is_really_an_image_of_text(deck):
    """The fixture's own bite check. If tesseract cannot read this picture, every 'no longer
    reports 1.4.5' assertion below would pass on a deck the detector never flagged."""
    assert _first_party_145(deck) == ["OCR_IMAGE_OF_TEXT"]
    transcript = " ".join(_proposed(deck)["image 1"].split()).lower()
    assert "benefits" in transcript and "enrollment" in transcript, transcript


@needs_ocr
def test_a_deck_with_no_pictures_is_not_flagged():
    """The control. Without it, a detector that flagged every deck would satisfy the test above
    and the whole file would be measuring nothing."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = TITLE
    s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text_frame.text = " ".join(LINES)
    out = Path(tempfile.mkdtemp()) / FILE
    prs.save(out)
    assert "1.4.5" not in _assess(out.read_bytes())


@needs_ocr
def test_removing_only_the_picture_does_not_clear_it(deck):
    """THE TRAP THIS LANE TURNS ON, measured. `ocr._ooxml_images` reads ppt/media/* from the zip
    and never opens a slide, so a writer that removed the <p:pic> and left the media part would
    change the deck, pass every "the text box is there" assertion, and still fail the criterion
    it was credited against — the exact shape of the defect #1665 found in the descr lane.
    """
    import re
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(deck)) as zin, \
         zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zin.infolist():
            content = zin.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml":
                content = re.sub(rb"<p:pic>.*?</p:pic>", b"", content, flags=re.S)
            zo.writestr(info, content)
    stripped = out.getvalue()
    assert _media(stripped), "the media part was removed too; this control proves nothing"
    assert "1.4.5" in _assess(stripped)


# ── 2. approval → write → re-scan → credit, through the real path ─────────────

@pytest.fixture()
def applied(store, monkeypatch, deck):
    blob = _Blob(deck)
    _seed(store, {"image 1": " ".join(LINES)})
    _run_lane(monkeypatch, store, blob)
    return blob, store


@needs_ocr
def test_the_saved_deck_carries_the_text_as_a_real_shape(applied):
    blob, _ = applied
    assert any("Benefits at a glance" in t for t in _texts(blob.data)), _texts(blob.data)


@needs_ocr
def test_the_image_is_gone_from_the_package(applied):
    """Not merely unreferenced — removed. This is what clears the criterion."""
    blob, _ = applied
    assert _media(blob.data) == []


@needs_ocr
def test_the_image_relationship_is_gone_too(applied):
    blob, _ = applied
    with zipfile.ZipFile(io.BytesIO(blob.data)) as z:
        rels = z.read("ppt/slides/_rels/slide1.xml.rels").decode()
    assert "/image" not in rels, rels


@needs_ocr
def test_unrelated_content_survives(applied):
    blob, _ = applied
    texts = _texts(blob.data)
    assert f"{TITLE} 1" in texts and BODY in texts


@needs_ocr
def test_the_deck_still_opens(applied):
    """Through python-pptx, which had no part in writing the change."""
    from pptx import Presentation
    blob, _ = applied
    assert zipfile.ZipFile(_spill(blob.data)).testzip() is None
    assert len(Presentation(str(_spill(blob.data))).slides) == 1


@needs_ocr
def test_a_second_real_assessment_no_longer_reports_1_4_5(applied):
    """THE claim: a fresh assessment of the SAVED bytes, not the writer's return value."""
    blob, _ = applied
    assert _first_party_145(blob.data) == []
    assert "1.4.5" not in _assess(blob.data)


@needs_ocr
def test_the_row_is_credited_and_the_copy_is_stored(applied):
    blob, store = applied
    assert store.count_unapplied_approved_values(SID, FILE) == 0
    assert blob.uploads


@needs_ocr
def test_the_write_introduces_no_new_failure(applied, deck):
    """Since #1712 a fix that breaks another criterion blocks certification, and this is the one
    lane that both DELETES content and INSERTS a shape — the two edits most able to disturb
    something else. Measured rather than assumed.

    1.1.1 and 1.4.9 clear here too, and that is not the lane overreaching: the picture they were
    about is gone, so the criteria it failed go with it. The assertion is one-directional for
    that reason — nothing NEW may appear, and 1.4.5 must be among what left.
    """
    blob, _ = applied
    before, after = _assess(deck), _assess(blob.data)
    assert after - before == set(), f"the write introduced {sorted(after - before)}"
    assert "1.4.5" in before - after


@needs_ocr
def test_one_image_on_two_slides_is_replaced_on_both(store, monkeypatch):
    """One media part, two placements, one locator. The part is deleted once, so every picture
    of it must become a text box or the deck would reference a part that is not there."""
    data = _deck(on_slides=2)
    assert len(_media(data)) == 1, "python-pptx stopped de-duplicating the image; fixture invalid"
    blob = _Blob(data)
    _seed(store, {"image 1": " ".join(LINES)})
    _run_lane(monkeypatch, store, blob)

    assert _media(blob.data) == []
    assert sum("Benefits at a glance" in t for t in _texts(blob.data)) == 2
    assert "1.4.5" not in _assess(blob.data)
    assert store.count_unapplied_approved_values(SID, FILE) == 0


# ── 3. where the lane must NOT credit ─────────────────────────────────────────

@needs_ocr
def test_a_partial_write_is_not_credited_because_the_criterion_still_fails(store, monkeypatch):
    """Two different images of text, one approved. The write succeeds and the deck genuinely
    improves — and 1.4.5 still fails, because the other picture is still a picture.

    This is the control that separates "the writer wrote something" from "the criterion
    cleared". A lane crediting on the write would mark the file compliant here and publish a
    deck that still fails the criterion it was certified against.
    """
    from apply_pptx_image_replacement import apply_pptx_image_replacement
    data = _two_images_deck()
    assert len(_media(data)) == 2, "expected two distinct images"
    assert len(_first_party_145(data)) == 2, "both pictures must trip 1.4.5 for this control"

    written, ap, _ = apply_pptx_image_replacement(data, {"image 1": " ".join(LINES)})
    assert ap, "the writer refused the value, so this control is not about crediting"
    assert "1.4.5" in _assess(written), (
        "replacing one of two images cleared the criterion — this control cannot distinguish a "
        "withheld credit from a cleared one")

    blob = _Blob(data)
    _seed(store, {"image 1": " ".join(LINES)})
    _run_lane(monkeypatch, store, blob)

    assert store.count_unapplied_approved_values(SID, FILE) == 1, (
        "the value was credited even though the criterion still fails on re-scan")
    assert store.mark_file_compliant_if_reviewed(SID, FILE) is False
    assert not blob.uploads, "an uncleared write was published as the corrected copy"
    assert blob.data == data


@needs_ocr
def test_replacing_both_images_does_clear_it(store, monkeypatch):
    """The other half of the pair, so the control above is known to be about COUNT rather than
    about the lane being broken for multi-image decks."""
    data = _two_images_deck()
    blob = _Blob(data)
    _seed(store, {"image 1": " ".join(LINES), "image 2": "Contact the benefits team"})
    _run_lane(monkeypatch, store, blob)

    assert _media(blob.data) == []
    assert "1.4.5" not in _assess(blob.data)
    assert store.count_unapplied_approved_values(SID, FILE) == 0
    assert blob.uploads


@needs_ocr
def test_a_grouped_picture_is_refused_rather_than_misplaced(store, monkeypatch, deck):
    """A picture inside a <p:grpSp> is positioned relative to the group's own transform, so a
    text box at those coordinates lands somewhere else on the slide. The writer refuses, and a
    refusal must leave the deck alone and credit nothing."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(deck)) as zin, \
         zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zin.infolist():
            content = zin.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml":
                xml = content.decode()
                xml = xml.replace("<p:pic>", "<p:grpSp><p:pic>", 1)
                xml = xml.replace("</p:pic>", "</p:pic></p:grpSp>", 1)
                content = xml.encode()
            zo.writestr(info, content)
    data = out.getvalue()

    blob = _Blob(data)
    _seed(store, {"image 1": " ".join(LINES)})
    _run_lane(monkeypatch, store, blob)

    assert blob.data == data, "a refused replacement still rewrote the deck"
    assert _media(blob.data), "the image was deleted despite the refusal"
    assert store.count_unapplied_approved_values(SID, FILE) == 1
    assert not blob.uploads


@needs_ocr
def test_an_approval_aimed_at_an_image_that_is_not_there_is_not_credited(store, monkeypatch, deck):
    blob = _Blob(deck)
    _seed(store, {"image 9": " ".join(LINES)})
    _run_lane(monkeypatch, store, blob)

    assert blob.data == deck
    assert store.count_unapplied_approved_values(SID, FILE) == 1
    assert not blob.uploads


@needs_ocr
def test_a_1_4_9_approval_does_not_reach_this_writer(store, monkeypatch, deck):
    """1.4.9 is AAA and exempts nothing, so its rows can be charts — and this writer DELETES the
    image. The lane reads ("1.4.5",) only, so a 1.4.9 approval leaves the deck untouched rather
    than destroying a chart nobody agreed to lose."""
    blob = _Blob(deck)
    _seed(store, {"image 1": " ".join(LINES)}, rule_id="1.4.9")
    _run_lane(monkeypatch, store, blob)

    assert blob.data == deck, "a 1.4.9 approval replaced an image"
    assert _media(blob.data), "a 1.4.9 approval deleted an image"
    assert not blob.uploads


# ── 4. a broken engine earns nothing ──────────────────────────────────────────

@needs_ocr
@pytest.mark.parametrize("name,script,timeout", [
    ("cannot be launched", None, None),
    ("exits non-zero", "#!/bin/sh\necho boom >&2\nexit 9\n", None),
    ("hangs past the timeout", "#!/bin/sh\nsleep 30\n", "2"),
])
def test_a_broken_office_analyser_never_credits_this_lane_either(monkeypatch, deck, name,
                                                                 script, timeout):
    """Re-asserted per lane rather than assumed to inherit: the fail-open #1058 closed lived in
    ONE shared seam, so a regression there takes every lane at once. 1.4.5 comes from the OCR
    pass, pure Python running after the .NET call, so the residual is a real set even with no
    analyser at all."""
    import stat as _stat

    import scanner
    if script is None:
        monkeypatch.setattr(scanner, "DOTNET", "/nonexistent/dotnet", raising=False)
    else:
        fake = Path(tempfile.mkdtemp()) / "dotnet"
        fake.write_text(script)
        fake.chmod(fake.stat().st_mode | _stat.S_IEXEC)
        monkeypatch.setattr(scanner, "DOTNET", str(fake), raising=False)
    if timeout:
        monkeypatch.setenv("ACP_OFFICE_CLI_TIMEOUT", timeout)

    from proposals import verify_residual_scs
    residual = verify_residual_scs(deck, FILE)
    assert residual is not None, (
        f"an office CLI that {name} made the re-scan return None — every approved value on this "
        f"lane would be credited on a scan that never happened")
    assert "1.4.5" in residual
