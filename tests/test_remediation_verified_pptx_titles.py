"""The 2.4.6 pptx slide-title lane, proved end to end (WCAG 2.4.6 Headings and Labels).

Same bar as every lane before it: the original deck trips the finding, an approval changes the
saved deck, a REAL re-scan verifies it, unrelated content survives, and a broken engine earns no
credit. Nothing but the blob store is patched — `handlers._apply_approved_values` runs the
production seam through `proposals.verify_residual_scs` to `scanner.analyse_and_assess`.

WHY THIS LANE WAS THE LAST OF EIGHTEEN TO GET A PROOF. It was ASSISTED in the capability table
and listed in the hand-written applier registry from the day it shipped, on the strength of a
writer unit test (hand-built zip, no detector, no re-scan) and a proposer test. Nothing ran the
chain. When the registry was made to DERIVE from round-trip fixtures instead, this was the one
(format, criterion) pair with nothing to derive from — and writing the proof found a hole:
the detector counts a `ctrTitle` placeholder (the Title Slide layout) as the slide's title,
while the writer matched only `type="title"`. An approved title on a Title Slide was refused as
unresolved and never credited. The writer now accepts both, and this file proves both.

WHY THE APPROVED VALUE IS AUTHORED HERE RATHER THAN PROPOSED. `propose_slide_titles` drafts from
a text model, and a lane that can only be proved where a model happens to be available is not
proved. A reviewer typing the title is the real workflow when no model is configured, and
`handlers` reads the approved value the same way whichever produced it.

WHERE THE DETECTOR'S GATE IS: `office_structure.pptx_checks` fires PPTX_TITLE_EMPTY once per
slide whose layout HAS a title placeholder that carries no text. A slide whose layout has no
title slot (Blank) is never a finding — that is a design choice, not a defect — so the writer
refusing such a slide is correct, and the control for it asserts a refusal, not a write.

WHAT THIS CLAIMS: the slides now carry an announceable title, the deck survived, and ACP's own
criterion stops firing. NOT that the titles are the ones a human would choose — naming is a
judgement, which is why the lane is `assisted`.
"""
from __future__ import annotations

import io
import sys
import tempfile
import zipfile
from pathlib import Path

import pytest
from hitl_viewed import approve_bound

ACP = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ACP / "api"))

# The (format, criterion) lanes this module PROVES end to end — read by
# tests/test_capability_assisted_contract.py, which derives the applier registry from
# these declarations instead of a hand-written list. A literal set, so it can be read
# without importing this module.
PROVES_LANES = {("pptx", "2.4.6")}

pytest.importorskip("pptx")

FILE = "deck.pptx"
SID = "rv-pptx-246"

TITLE_ONLY, TITLE_AND_CONTENT, TITLE_SLIDE, BLANK = 5, 1, 0, 6   # python-pptx default layouts

APPROVED = "Q3 findings by severity"
SECOND = "Remediation plan & owners"          # an ampersand, so the write must escape it
BODY = "Unrelated body copy that must survive the write."
KEPT = "Already titled"


def _deck(slides: list[tuple[int, str | None]]) -> bytes:
    """One slide per (layout index, title-or-None). None leaves the layout's title placeholder
    empty — the detector's condition. Every slide also gets body text so the proposer's gate
    (a slide with content to name) holds too."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for layout, title in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[layout])
        if title is not None and slide.shapes.title is not None:
            slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
        box.text_frame.text = BODY
    out = Path(tempfile.mkdtemp()) / FILE
    prs.save(out)
    return out.read_bytes()


def _assess(data: bytes) -> set[str]:
    """The SCs a REAL assessment reports — the same call the production re-verification makes."""
    from assessment_policy import _extract_sc
    from scanner import analyse_and_assess
    with tempfile.TemporaryDirectory() as d:
        (Path(d) / FILE).write_bytes(data)
        fd, _ = analyse_and_assess(Path(d), FILE, detect_pii=False)
    return {sc for i in (fd or {}).get("issues", []) if (sc := _extract_sc(i.get("wcag", "")))}


def _spill(data: bytes) -> Path:
    p = Path(tempfile.mkdtemp()) / FILE
    p.write_bytes(data)
    return p


def _first_party_empty_titles(data: bytes) -> list[str]:
    """Where the FIRST-PARTY detector says a title is empty ("Slide N"), asked of the detector
    rather than of a full scan — the .NET analyser, where built, reports 2.4.6 in its own shape
    and a scan-level assertion would depend on which engines happen to be installed."""
    from office_structure import pptx_checks
    return [f["location"] for f in pptx_checks(_spill(data)) if f.get("ruleId") == "PPTX_TITLE_EMPTY"]


def _slide_xml(data: bytes, num: int) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(f"ppt/slides/slide{num}.xml").decode("utf-8")


def _titles(data: bytes) -> list[str]:
    """Slide titles as python-pptx reads them back — a reader that had no part in the write."""
    from pptx import Presentation
    prs = Presentation(str(_spill(data)))
    return [(s.shapes.title.text if s.shapes.title is not None else None) for s in prs.slides]


class _Blob:
    """The only thing patched in this module. Stores bytes verbatim; decides nothing."""

    def __init__(self, data: bytes):
        self.data, self.uploads = data, []

    def download_remediated(self, owner, sid, f):
        return self.data

    def upload_remediated(self, owner, sid, f, data, mime):
        self.data = data
        self.uploads.append((f, mime))
        return "http://b/2"
    # The approved writer publishes digest-scoped and moves the pointer at commit.
    def upload_immutable_retry(self, owner, sid, f, data, mime):
        return self.upload_remediated(owner, sid, f, data, mime)


@pytest.fixture()
def store(monkeypatch):
    import store as store_mod
    monkeypatch.setattr(store_mod, "_SQLITE_PATH", Path(tempfile.mkdtemp()) / "rv.db")
    return store_mod.Store()


def _seed(store, values: dict[str, str]) -> str:
    """A scanned + remediated deck with one 2.4.6 card, and `values` approved on it."""
    store.init_scan_run(SID, "drive", 1, "2026-09-07T00:00:00Z", "rubric", "hash")
    store.save_file_result(SID, {
        "file": FILE, "engine": "office", "status": "pass", "score": 60, "compliant": 0,
        "skipped_rules": 0, "drive_file_id": "d1",
        "issues": [{"ruleId": "PPTX_TITLE_EMPTY", "wcag": "2.4.6 Headings and Labels",
                    "severity": "MODERATE", "location": loc.title()} for loc in values],
    }, "2026-09-07T00:00:00Z")
    store.record_remediation(SID, FILE, drive_write_url="http://d/1", blob_url="http://b/1")
    item_id = store.enqueue_proposals(SID, FILE, "2.4.6", [
        {"locator": loc, "before": "(title placeholder left empty)", "proposed_value": "",
         "rationale": "r", "source": "reviewer"} for loc in values],
        rule_name="Headings and Labels")
    approve_bound(store, item_id, list(values.values()))
    return item_id


def _run_lane(monkeypatch, store, blob):
    """The production handler, with the re-scan UNPATCHED."""
    import core
    import handlers
    monkeypatch.setattr(core, "store", store)
    monkeypatch.setitem(sys.modules, "blob", blob)
    handlers._apply_approved_values({"scan_id": SID, "file": FILE}, {})


@pytest.fixture(scope="module")
def untitled() -> bytes:
    """Two slides on two layouts whose title placeholders are left empty, and one that is titled."""
    return _deck([(TITLE_ONLY, None), (TITLE_AND_CONTENT, None), (TITLE_ONLY, KEPT)])


# ── 1. the finding and its gate ───────────────────────────────────────────────

def test_a_real_assessment_reports_2_4_6_on_an_empty_title_placeholder(untitled):
    assert "2.4.6" in _assess(untitled)


def test_the_first_party_detector_names_exactly_the_empty_slides(untitled):
    assert _first_party_empty_titles(untitled) == ["Slide 1", "Slide 2"]


def test_a_titled_deck_is_not_flagged():
    """The control. Without it, a detector that flagged every slide would satisfy the test
    above and the whole file would be measuring nothing."""
    deck = _deck([(TITLE_ONLY, APPROVED), (TITLE_AND_CONTENT, SECOND)])
    assert _first_party_empty_titles(deck) == []
    assert "2.4.6" not in _assess(deck)


def test_a_layout_with_no_title_slot_is_not_a_finding():
    """A Blank slide has no title placeholder at all. That is a design choice, not a defect, so
    the detector stays silent — and, below, the writer refuses to invent a title shape for it."""
    assert _first_party_empty_titles(_deck([(BLANK, None)])) == []


def test_a_title_slide_layout_is_the_same_finding():
    """The Title Slide layout's centred title is `ctrTitle`, not `title`. The detector counts it
    as the slide's title; until this proof the writer did not, and an approval aimed at it was
    refused. Pinned here so the two cannot drift apart again."""
    deck = _deck([(TITLE_SLIDE, None)])
    assert 'type="ctrTitle"' in _slide_xml(deck, 1)
    assert _first_party_empty_titles(deck) == ["Slide 1"]


# ── 2. approval → write → re-scan → credit, through the real path ─────────────

@pytest.fixture()
def applied(store, monkeypatch, untitled):
    blob = _Blob(untitled)
    _seed(store, {"slide 1": APPROVED, "slide 2": SECOND})
    _run_lane(monkeypatch, store, blob)
    return blob, store


def test_the_saved_deck_carries_the_approved_titles(applied):
    blob, _ = applied
    assert _titles(blob.data) == [APPROVED, SECOND, KEPT]


def test_the_ampersand_is_escaped_in_the_xml_and_read_back_verbatim(applied):
    blob, _ = applied
    assert "Remediation plan &amp; owners" in _slide_xml(blob.data, 2)
    assert _titles(blob.data)[1] == SECOND


def test_unrelated_content_and_the_titled_slide_survive(applied):
    blob, _ = applied
    for num in (1, 2, 3):
        assert BODY in _slide_xml(blob.data, num)
    assert _titles(blob.data)[2] == KEPT


def test_the_deck_still_opens(applied):
    blob, _ = applied
    assert zipfile.ZipFile(_spill(blob.data)).testzip() is None
    assert len(_titles(blob.data)) == 3


def test_a_second_real_assessment_no_longer_reports_2_4_6(applied):
    """THE claim: a fresh assessment of the SAVED bytes, not the writer's return value."""
    blob, _ = applied
    assert _first_party_empty_titles(blob.data) == []
    assert "2.4.6" not in _assess(blob.data)


def test_the_row_is_credited_and_the_copy_is_stored(applied):
    blob, store = applied
    assert store.count_unapplied_approved_values(SID, FILE) == 0
    assert blob.uploads


def test_a_title_slide_round_trips_too(store, monkeypatch):
    """The hole this proof found, closed and held closed through the LANE: a Title Slide's
    empty `ctrTitle` is written, re-scanned clean and credited like any other title."""
    deck = _deck([(TITLE_SLIDE, None)])
    blob = _Blob(deck)
    _seed(store, {"slide 1": APPROVED})
    _run_lane(monkeypatch, store, blob)

    assert _titles(blob.data) == [APPROVED]
    assert _first_party_empty_titles(blob.data) == []
    assert "2.4.6" not in _assess(blob.data)
    assert store.count_unapplied_approved_values(SID, FILE) == 0
    assert blob.uploads


# ── 3. where the lane must NOT credit ─────────────────────────────────────────

def test_a_partial_write_is_not_credited_because_the_criterion_still_fails(store, monkeypatch,
                                                                            untitled):
    """Two empty titles, one approved. The write succeeds and the deck genuinely improves — and
    2.4.6 still fails, because the other slide is still nameless.

    This is the control that separates "the writer wrote something" from "the criterion
    cleared". A lane crediting on the write would mark the file compliant here and publish a
    deck that still fails the criterion it was certified against.
    """
    from apply_pptx_slide_titles import apply_pptx_slide_titles
    written, ap, _ = apply_pptx_slide_titles(untitled, {"slide 1": APPROVED})
    assert ap, "the writer refused the value, so this control is not about crediting"
    assert _first_party_empty_titles(written) == ["Slide 2"], (
        "titling one of two slides cleared the criterion — this control cannot distinguish a "
        "withheld credit from a cleared one")

    blob = _Blob(untitled)
    _seed(store, {"slide 1": APPROVED})
    _run_lane(monkeypatch, store, blob)

    assert store.count_unapplied_approved_values(SID, FILE) == 1, (
        "the value was credited even though the criterion still fails on re-scan")
    assert store.mark_file_compliant_if_reviewed(SID, FILE) is False
    assert not blob.uploads, "an uncleared write was published as the corrected copy"
    assert blob.data == untitled


def test_an_approval_aimed_at_a_slide_with_no_title_slot_is_refused_not_credited(store,
                                                                                  monkeypatch):
    """A Blank slide is not a finding, so nothing should ever be approved for it — but a
    reviewer can type anything. The writer must refuse (no shape to fill, and inventing one
    is layout authoring, not remediation), and refusal must not credit."""
    deck = _deck([(BLANK, None)])
    blob = _Blob(deck)
    _seed(store, {"slide 1": APPROVED})
    _run_lane(monkeypatch, store, blob)

    assert blob.data == deck
    assert store.count_unapplied_approved_values(SID, FILE) == 1
    assert not blob.uploads


def test_an_approval_aimed_at_a_slide_that_is_not_there_is_not_credited(store, monkeypatch,
                                                                         untitled):
    blob = _Blob(untitled)
    _seed(store, {"slide 9": APPROVED})
    _run_lane(monkeypatch, store, blob)

    assert blob.data == untitled
    assert store.count_unapplied_approved_values(SID, FILE) == 1
    assert not blob.uploads


def test_an_already_titled_deck_is_left_alone(store, monkeypatch):
    deck = _deck([(TITLE_ONLY, APPROVED)])
    assert "2.4.6" not in _assess(deck)
    blob = _Blob(deck)
    _seed(store, {})
    _run_lane(monkeypatch, store, blob)
    assert blob.data == deck and not blob.uploads


# ── 4. a broken engine earns nothing ──────────────────────────────────────────

@pytest.mark.parametrize("name,script,timeout", [
    ("cannot be launched", None, None),
    ("exits non-zero", "#!/bin/sh\necho boom >&2\nexit 9\n", None),
    ("hangs past the timeout", "#!/bin/sh\nsleep 30\n", "2"),
])
def test_a_broken_office_analyser_never_credits_this_lane_either(monkeypatch, untitled, name,
                                                                 script, timeout):
    """Re-asserted per lane rather than assumed to inherit: the fail-open #1058 closed lived in
    ONE shared seam, so a regression there takes every lane at once. 2.4.6 on pptx comes from
    office_structure, pure Python running after the .NET call, so the residual is a real set
    even with no analyser at all."""
    import stat as _stat

    import scanner
    if script is None:
        monkeypatch.setattr(scanner, "DOTNET", "/nonexistent/dotnet", raising=False)
    else:
        fake = Path(tempfile.mkdtemp()) / "dotnet"
        fake.write_text(script)
        fake.chmod(fake.stat().st_mode | _stat.S_IEXEC)
        monkeypatch.setattr(scanner, "DOTNET", str(fake), raising=False)
    if timeout:
        monkeypatch.setenv("ACP_OFFICE_CLI_TIMEOUT", timeout)

    from proposals import verify_residual_scs
    residual = verify_residual_scs(untitled, FILE)
    assert residual is not None, (
        f"an office CLI that {name} made the re-scan return None — every approved value on this "
        f"lane would be credited on a scan that never happened")
    assert "2.4.6" in residual
