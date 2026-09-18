"""A review row whose target a DIFFERENT verified fix removed stops being work — and only then.

THE PRODUCTION SHAPE (synthetic reproduction; no customer content). A Word document with one body
picture carried a PENDING 1.1.1 row (drafted alt text, finding `docx:drawing:1:paragraph:N`) and
an APPROVED, applied 1.4.5 row that replaced the same picture with its OCR transcript. The
corrected copy has no drawing left and a complete fresh assessment reports nothing — yet the 1.1.1
row stayed pending forever and the ledger kept one finding `awaiting_review`, because 1.1.1 reads
REVIEW on every format and `_superseded_items` only retracts on PASS/NOT_EVALUATED.

Fixtures are real OOXML packages built by python-docx and edited by the production 1.4.5 writer
(`apply_office_image_replacement`); the store is a real SQLite Store with a real finding ledger.
The end-to-end proof through the real approved writer and the real Office analyser lives in
tests/test_remediation_verified_target_replacement.py.
"""
from __future__ import annotations

import io
import json
import sys
from hashlib import sha256

import pytest

import review_target_reconciliation as rtr
from apply_office_image_replacement import apply_office_image_replacement

SID = 'trm-scan'
FILE = 'synthetic-target-replacement.docx'
OWNER = 'owner@example.com'
TEXT = 'Opening hours\nMonday to Friday'
C1_KEYS = {'removed_by_item_id', 'removed_by_rule_id', 'targets', 'finding_ids',
           'corrected_artifact_sha256', 'source_artifact_sha256', 'verified_at', 'assessment',
           'assessment_status', 'skipped_rules'}


def picture(color):
    from PIL import Image
    out = io.BytesIO()
    Image.new('RGB', (120, 60), color).save(out, 'PNG')
    return out.getvalue()


def docx(colors=('red',)):
    from docx import Document
    doc = Document()
    doc.add_heading('Synthetic report', 1)
    doc.add_paragraph('Synthetic body copy.')
    for color in colors:
        doc.add_picture(io.BytesIO(picture(color)))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def sha(data):
    return sha256(data).hexdigest()


class Blob:
    """Stores bytes verbatim; decides nothing."""

    def __init__(self, remediated, source=None):
        self.data, self.source, self.uploads = remediated, source, []

    def enabled(self):
        return True

    def download_remediated(self, owner, sid, f):
        return self.data

    def download_source(self, owner, sid, f, checksum=None):
        return self.source

    def upload_remediated(self, owner, sid, f, data, mime):
        self.data = data
        self.uploads.append(f)
        return 'http://b/new'


class World:
    def __init__(self, store, *, colors=('red',), alt_targets=('word/document.xml#Picture 1',),
                 alt_count=None, alt_status='pending', replaced=('image 1',)):
        self.store = store
        self.source = docx(colors)
        self.corrected = self.source
        if replaced:
            self.corrected, applied, unresolved = apply_office_image_replacement(
                self.source, 'docx', {loc: TEXT for loc in replaced})
            assert applied and not unresolved
        with store._db.cursor() as cur:
            store._db.execute(cur,
                "INSERT INTO scan_runs(id,source,status,workflow_id,workflow_revision,owner_email) "
                "VALUES(%s,'local','done',%s,1,%s)", (SID, SID, OWNER))
            store._db.execute(cur,
                "INSERT INTO file_records(scan_id,file,checksum) VALUES(%s,%s,'c1')", (SID, FILE))
            store._db.execute(cur,
                "INSERT INTO scan_rule_traces(scan_id,file,rule_id,rule_name,plain_name,level,"
                "fix_mode,outcome,finding_count) VALUES(%s,%s,'1.1.1','Non-text Content','Alt',"
                "'A','ai','REVIEW',%s)", (SID, FILE, len(colors)))
            store._db.execute(cur,
                "INSERT INTO scan_rule_traces(scan_id,file,rule_id,rule_name,plain_name,level,"
                "fix_mode,outcome,finding_count) VALUES(%s,%s,'1.4.5','Images of Text','Text',"
                "'AA','ai','FAIL',1)", (SID, FILE))
            for i, _ in enumerate(colors):
                # The Office analyser's own location shape: docx:drawing:{docPr id}:paragraph:{i}
                store._db.execute(cur,
                    "INSERT INTO issue_records(scan_id,file,rule_id,wcag,severity,detail,location) "
                    "VALUES(%s,%s,'DOCX-ALT-001','SC_1_1_1','CRITICAL','missing alt',%s)",
                    (SID, FILE, f'docx:drawing:{i + 1}:paragraph:{i + 2}'))
        execution = store.enqueue_stage_batch(SID, 'remediate', 'remediate_file',
            [{'scan_id': SID, 'file': FILE}], snapshot_id=SID, request_fingerprint=SID)
        self.batch = execution['batch_id']
        store.seed_finding_dispositions(SID, self.batch)
        store.record_remediation(SID, FILE, blob_url='http://b/1', corrected_sha256=sha(self.corrected))
        # The 1.4.5 replacement: approved, applied, and its finding verified by the writer's diff.
        self.fixer = store.enqueue_proposals(SID, FILE, '1.4.5', [
            {'locator': loc, 'before': 'image of text', 'proposed_value': TEXT, 'source': 'OCR'}
            for loc in (replaced or ('image 1',))], rule_name='Images of Text')
        if replaced:
            with store._db.cursor() as cur:
                store._db.execute(cur, "UPDATE hitl_queue SET status='approved',applied=1 WHERE id=%s",
                                  (self.fixer,))
            store.record_remediation_diffs(SID, FILE, [
                {'rule_id': '1.4.5', 'before': 'image of text', 'after': TEXT}])
        # The 1.1.1 row with its AI draft(s).
        self.alt = store.enqueue_proposals(SID, FILE, '1.1.1', [
            {'locator': loc, 'before': '', 'proposed_value': f'Synthetic banner {n}', 'source': 'vision'}
            for n, loc in enumerate(alt_targets)], rule_name='Non-text Content',
            finding_count=alt_count)
        with store._db.cursor() as cur:
            store._db.execute(cur, "UPDATE hitl_queue SET proposal_snapshot_ids=%s WHERE id=%s",
                              (json.dumps([f'snap-{n}' for n in range(len(alt_targets))]), self.alt))
        if alt_status == 'approved':
            store.complete_hitl_decision(self.alt, 'approved', None, None, resolution=None,
                                         approved_values=[None] * len(alt_targets),
                                         actor=OWNER, detail=None)

    def evidence(self, **changes):
        from verification_identity import scope_identity
        base = {'assessment': rtr.SAVED_EVIDENCE, 'artifact_sha256': sha(self.corrected),
                'assessment_ok': True, 'assessment_status': 'analysed', 'skipped_rules': 0,
                'errors': [], 'remaining_issues': [], 'remaining_criteria': [],
                'assessment_scope': scope_identity(self.store.get_scan_scope(SID)),
                'assessment_reused': False, 'verified_at': '2026-09-17T12:00:00+00:00'}
        base.update(changes)
        return {k: v for k, v in base.items() if v is not DROP}

    def reconcile(self, **evidence_changes):
        return rtr.reconcile(self.store, SID, FILE, source=self.source, corrected=self.corrected,
                             evidence=self.evidence(**evidence_changes), source_kind='assessed_source')

    def findings(self, rule='1.1.1'):
        return [f for f in self.store.list_finding_dispositions(SID, self.batch) if f['rule_id'] == rule]

    def lines(self):
        return [d for d in self.store.list_decisions(scan_id=SID) if d['action'] == rtr.ACTION]

    def row(self, **kwargs):
        rows = self.store.list_hitl_queue(scan_id=SID, **kwargs)
        return next((r for r in rows if r['id'] == self.alt), None)


DROP = object()


@pytest.fixture
def store(isolated_store):
    return isolated_store


# ── the production case ────────────────────────────────────────────────────────────────────

def test_replacement_retires_the_pending_alt_row_with_evidence_and_nothing_else(store):
    w = World(store)
    before_row = store.get_hitl_item(w.alt)
    assert [f['disposition'] for f in w.findings()] == ['awaiting_review']
    assert w.findings()[0]['instance_key'] == 'docx:drawing:1:paragraph:2'
    assert w.row() is not None                      # visible work before

    result = w.reconcile()

    assert [s['item_id'] for s in result['superseded']] == [w.alt]
    # The DB row is untouched: still pending, same version, no alt text, no approval.
    after_row = store.get_hitl_item(w.alt)
    for key in ('status', 'decision_version', 'proposals', 'approved_value', 'applied',
                'reviewed_at', 'approved_proposal_snapshot_ids'):
        assert after_row.get(key) == before_row.get(key), key
    assert after_row['status'] == 'pending'
    assert not any(p.get('approved_value') for p in after_row['proposals'])
    # Hidden from the work list, explained in the audit view.
    assert w.row() is None
    audit = w.row(include_superseded=True)
    assert audit['superseded'] is True
    assert audit['superseded_reason'] == 'target_removed_by_verified_fix'
    evidence = audit['superseded_evidence']
    assert set(evidence) == C1_KEYS
    assert evidence['removed_by_item_id'] == w.fixer and evidence['removed_by_rule_id'] == '1.4.5'
    assert evidence['targets'] == ['word/document.xml#Picture 1']
    assert evidence['corrected_artifact_sha256'] == sha(w.corrected)
    assert evidence['source_artifact_sha256'] == sha(w.source)
    assert evidence['assessment'] == 'release.corrected_copy_assessed'
    assert evidence['assessment_status'] == 'analysed' and evidence['skipped_rules'] == 0
    # Exactly that finding, with the evidence id, via the ledger's own transition.
    [finding] = w.findings()
    assert finding['disposition'] == 'superseded_by_reassessment'
    [line] = w.lines()
    assert finding['fix_evidence_ids'] == [f"target_removed:{line['id']}"]
    assert evidence['finding_ids'] == [finding['finding_id']]
    assert line['actor'] == 'acp'
    detail = json.loads(line['detail'])
    assert detail['item_id'] == w.alt and detail['decision_version'] == 0
    assert detail['proposal_snapshot_ids'] == ['snap-0'] and detail['batch_id'] == w.batch
    # No other finding moved; 1.4.5 stays resolved_verified; nothing certified.
    assert [f['disposition'] for f in w.findings('1.4.5')] == ['resolved_verified']
    assert not (store.get_file_record(SID, FILE) or {}).get('compliant')
    with store._db.cursor() as cur:
        store._db.execute(cur, "SELECT COUNT(*) AS n FROM hitl_events WHERE item_id=%s", (w.alt,))
        assert store._db.fetchone(cur)['n'] == 0
        store._db.execute(cur, "SELECT COUNT(*) AS n FROM jobs WHERE scan_id=%s", (SID,))
        assert store._db.fetchone(cur)['n'] <= 1     # only the remediate batch job


def test_replay_is_idempotent(store):
    w = World(store)
    first = w.reconcile()
    second = w.reconcile()
    assert first['superseded'] and not second['superseded']
    assert [u['item_id'] for u in second['unchanged']] == [w.alt]
    assert len(w.lines()) == 1
    events = store.finding_disposition_events(SID, w.batch, w.findings()[0]['finding_id'])
    assert [e['to_disposition'] for e in events].count('superseded_by_reassessment') == 1


# ── what must NOT be retired ───────────────────────────────────────────────────────────────

def test_other_image_persisting_keeps_the_row_and_both_findings(store):
    """One 1.1.1 row covers both pictures; only the first was replaced. Partial = untouched."""
    w = World(store, colors=('red', 'blue'),
              alt_targets=('word/document.xml#Picture 1', 'word/document.xml#Picture 2'))
    before = [(f['finding_id'], f['disposition']) for f in w.findings()]
    result = w.reconcile(remaining_issues=[{'wcag': 'SC_1_1_1', 'location': 'docx:drawing:2:paragraph:3'}],
                         remaining_criteria=['1.1.1'])
    assert not result['superseded']
    assert {'item_id': w.alt, 'reason': 'target_remains'} in result['skipped']
    assert [(f['finding_id'], f['disposition']) for f in w.findings()] == before
    assert w.lines() == [] and w.row() is not None


def test_ordinal_shift_of_a_surviving_picture_neither_blocks_nor_leaks(store):
    """Picture 1 is replaced, so Picture 2 becomes the FIRST drawing of the corrected part.
    Its remaining 1.1.1 issue must not be mistaken for the removed target (same ordinal), and its
    finding must stay open."""
    w = World(store, colors=('red', 'blue'), alt_targets=('word/document.xml#Picture 1',),
              alt_count=1)
    after = rtr.placements(w.corrected)
    assert [(p.ordinal, p.docpr_id) for p in after.items] == [(0, '2')]     # the shift is real
    result = w.reconcile(remaining_issues=[{'wcag': 'SC_1_1_1', 'location': 'docx:drawing:2:paragraph:3'}],
                         remaining_criteria=['1.1.1'])
    assert [s['item_id'] for s in result['superseded']] == [w.alt], result
    by_key = {f['instance_key']: f['disposition'] for f in w.findings()}
    assert by_key == {'docx:drawing:1:paragraph:2': 'superseded_by_reassessment',
                      'docx:drawing:2:paragraph:3': None}


def test_a_drawing_that_survives_without_its_relationship_is_not_removed(store):
    """Removal is read off the DRAWING, not the relationships: a picture whose rel (and name)
    changed but whose w:drawing still stands in the part is still there."""
    import re
    import zipfile
    w = World(store, replaced=())
    parts = {}
    with zipfile.ZipFile(io.BytesIO(w.source)) as z:
        for name in z.namelist():
            parts[name] = z.read(name)
    parts['word/document.xml'] = parts['word/document.xml'].replace(b'name="Picture 1"', b'name="Renamed"')
    parts['word/_rels/document.xml.rels'] = re.sub(
        rb'<Relationship [^>]*Id="rId9"[^>]*/>', b'', parts['word/_rels/document.xml.rels'])
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w') as z:
        for name, data in parts.items():
            z.writestr(name, data)
    w.corrected = out.getvalue()
    assert rtr.resolve('word/document.xml#Picture 1', rtr.placements(w.corrected)) is None
    store.record_remediation(SID, FILE, blob_url='http://b/9', corrected_sha256=sha(w.corrected))
    with store._db.cursor() as cur:
        store._db.execute(cur, "UPDATE hitl_queue SET status='approved',applied=1 WHERE id=%s", (w.fixer,))
    store.record_remediation_diffs(SID, FILE, [{'rule_id': '1.4.5', 'before': 'i', 'after': TEXT}])
    result = w.reconcile()
    assert {'item_id': w.alt, 'reason': 'target_remains'} in result['skipped'], result
    assert w.lines() == []


def test_a_remaining_issue_that_locates_the_target_blocks(store):
    w = World(store)
    result = w.reconcile(remaining_issues=[{'wcag': 'SC_1_1_1', 'location': 'docx:drawing:1:paragraph:2'}])
    assert not result['superseded'] and w.lines() == []


def test_an_unlocated_remaining_issue_of_the_criterion_blocks(store):
    w = World(store)
    result = w.reconcile(remaining_issues=[{'wcag': 'SC_1_1_1'}])
    assert {'item_id': w.alt, 'reason': 'unlocated_remaining_issue'} in result['skipped']
    assert w.lines() == []


@pytest.mark.parametrize('changes,reason', [
    ({'assessment_ok': False}, 'verification_failed'),
    ({'assessment_status': 'error'}, 'verification_incomplete'),
    ({'assessment_status': 'uncertain'}, 'verification_incomplete'),
    ({'skipped_rules': 1}, 'verification_partial'),
    ({'errors': [{'rule': 'x'}]}, 'verification_errors'),
    ({'artifact_sha256': 'f' * 64}, 'verification_digest_mismatch'),
    ({'assessment_reused': True}, 'verification_reused'),
    ({'assessment': 'something.else'}, 'verification_kind_unknown'),
    # MISSING is unknown, never zero/empty:
    ({'skipped_rules': DROP}, 'verification_skipped_rules_unknown'),
    ({'skipped_rules': None}, 'verification_partial'),
    ({'errors': DROP}, 'verification_errors_unknown'),
    ({'errors': None}, 'verification_errors_unknown'),
    ({'remaining_issues': DROP}, 'verification_incomplete'),
    ({'assessment_status': DROP}, 'verification_incomplete'),
    ({'assessment_ok': DROP}, 'verification_failed'),
    ({'assessment_scope': {'1.4.5': ['docx']}}, 'scope_changed'),
], ids=lambda v: v if isinstance(v, str) else None)
def test_incomplete_or_mismatched_evidence_changes_nothing(store, changes, reason):
    w = World(store)
    result = w.reconcile(**changes)
    assert result['superseded'] == [] and result['skipped'][0]['reason'] == reason
    assert w.lines() == [] and w.findings()[0]['disposition'] == 'awaiting_review'
    assert w.row() is not None


def test_missing_evidence_changes_nothing(store):
    w = World(store)
    result = rtr.reconcile(store, SID, FILE, source=w.source, corrected=w.corrected, evidence=None,
                           source_kind='assessed_source')
    assert result['skipped'][0]['reason'] == 'verification_missing' and w.lines() == []


def test_assessed_source_digest_mismatch_changes_nothing(store):
    w = World(store)
    with store._db.cursor() as cur:
        store._db.execute(cur,
            "INSERT INTO remediation_contribution_proposals(owner_id,scan_id,run_id,proposal_id,"
            "proposal_sha256,source_sha256,assessment_revision,file,rule_id,item_id,finding_ids_json,"
            "created_at) VALUES(%s,%s,%s,'p1','x',%s,'r',%s,'1.1.1',%s,'[]','t')",
            (OWNER, SID, w.batch, 'e' * 64, FILE, w.alt))
    assert w.reconcile()['skipped'][0]['reason'] == 'source_digest_mismatch'
    assert w.lines() == []


def test_artifact_not_current_changes_nothing(store):
    w = World(store)
    store.record_remediation(SID, FILE, blob_url='http://b/2', corrected_sha256='d' * 64)
    assert w.reconcile()['skipped'][0]['reason'] == 'artifact_not_current'


def test_no_verified_removing_fix_changes_nothing(store):
    """The picture is gone but the 1.4.5 row was never applied: an unexplained removal."""
    w = World(store)
    with store._db.cursor() as cur:
        store._db.execute(cur, "UPDATE hitl_queue SET applied=0 WHERE id=%s", (w.fixer,))
    result = w.reconcile()
    assert {'item_id': w.alt, 'reason': 'no_verified_removing_fix'} in result['skipped']
    assert w.lines() == []


def test_criterion_outside_the_assessed_scope_changes_nothing(store, monkeypatch):
    w = World(store)
    scope = {'1.4.5': frozenset({'docx'})}
    monkeypatch.setattr(store, 'get_scan_scope', lambda *a, **k: scope)
    result = w.reconcile(assessment_scope={'1.4.5': ['docx']})
    assert {'item_id': w.alt, 'reason': 'criterion_not_assessed'} in result['skipped']


# ── binding: scope/sha re-checked UNDER the commit lock; lock order ────────────────────────

def test_scope_change_between_prepare_and_commit_refuses(store, monkeypatch):
    w = World(store)
    real = rtr._scope
    calls = []

    def scope(st, sid, f, *, refresh=False):
        calls.append(refresh)
        if len(calls) >= 2:                       # the re-check under the lock
            return {'1.1.1': frozenset({'docx'})}, {'1.1.1': ['docx']}
        return real(st, sid, f, refresh=refresh)
    monkeypatch.setattr(rtr, '_scope', scope)
    result = w.reconcile()
    assert result['superseded'] == []
    assert result['skipped'][-1]['reason'] == 'artifact_changed_before_commit'
    assert w.lines() == [] and w.findings()[0]['disposition'] == 'awaiting_review'


def test_sha_change_between_prepare_and_commit_refuses(store, monkeypatch):
    w = World(store)
    monkeypatch.setattr(rtr, 'lock_file', lambda *a: 'c' * 64)
    assert w.reconcile()['skipped'][-1]['reason'] == 'artifact_changed_before_commit'
    assert w.lines() == []


def test_reconciliation_locks_review_rows_before_the_file(store, monkeypatch):
    """SEQUENCE EVIDENCE ONLY (no disposable PostgreSQL in this environment): the order the
    locks are REQUESTED in is row(s) then file — the approved writer's documented order."""
    w = World(store)
    order = []
    real_rows, real_file = store._get_hitl_item_for_decision, rtr.lock_file
    monkeypatch.setattr(store, '_get_hitl_item_for_decision',
                        lambda i: (order.append(('row', i)), real_rows(i))[1])
    monkeypatch.setattr(rtr, 'lock_file', lambda *a: (order.append(('file', a[2])), real_file(*a))[1])
    assert w.reconcile()['superseded']
    assert order == [('row', w.alt), ('file', FILE)]


# ── stale approvals, retries and resurrection ──────────────────────────────────────────────

def test_approved_row_retired_later_is_refused_by_retry_writer_and_approval(store, monkeypatch):
    w = World(store, alt_status='approved')
    item = store.get_hitl_item(w.alt)
    assert item['status'] == 'approved' and not item.get('applied')
    binding, refusal = store.approved_write_binding(item)
    assert refusal is None                          # a retry WOULD have been admitted before
    stale_payload = {'scan_id': SID, 'file': FILE, 'item_id': w.alt, 'approved_binding': binding}

    assert w.reconcile()['superseded']

    audit = w.row(include_superseded=True)
    assert audit['superseded'] and audit['approval_recheck_required'] is False
    assert store.get_hitl_item(w.alt)['status'] == 'approved'       # approval kept on record
    # Retry admission refuses before any job exists.
    with store._db.cursor() as cur:
        store._db.execute(cur, "SELECT COUNT(*) AS n FROM jobs WHERE type='apply_approved_values'")
        jobs_before = store._db.fetchone(cur)['n']
    result = store.retry_approved_write(w.alt)
    assert result['accepted'] is False and result['reason'] == store.RETRY_TARGET_REMOVED
    with store._db.cursor() as cur:
        store._db.execute(cur, "SELECT COUNT(*) AS n FROM jobs WHERE type='apply_approved_values'")
        assert store._db.fetchone(cur)['n'] == jobs_before
    # A retry job enqueued BEFORE reconciliation is refused by the writer under its lock.
    import core
    import handlers
    from worker import FatalJobError
    blob = Blob(w.corrected, w.source)
    monkeypatch.setattr(core, 'store', store)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    with pytest.raises(FatalJobError, match='removed the content'):
        handlers._apply_approved_values(stale_payload, {})
    # The file-wide job enqueued by the approval writes nothing for it.
    handlers._apply_approved_values({'scan_id': SID, 'file': FILE}, {})
    assert blob.uploads == [] and blob.data == w.corrected
    assert not store.get_hitl_item(w.alt).get('applied')
    # Not merely "unresolved and harmless": the retired row never reached the writer at all.
    assert not [d for d in store.list_decisions(scan_id=SID)
                if d['action'] in ('apply.unresolved', 'apply.unverified')]


def test_legacy_and_batch_approval_of_a_retired_row_are_refused(store):
    w = World(store)
    w.reconcile()
    for status in ('approved', 'rejected', 'skipped'):
        with pytest.raises(ValueError, match='stale proposal selection'):
            store.complete_hitl_decision(w.alt, status, None, None, resolution=None,
                                         approved_values=None, actor=OWNER, detail=None)
    with pytest.raises(ValueError, match='stale proposal selection'):
        store.complete_hitl_decision(w.alt, 'approved', None, None, resolution=None,
                                     approved_values=['Synthetic banner 0'], actor=OWNER,
                                     detail=None, request_id='r1', expected_version=0,
                                     expected_proposal_snapshot_ids=['snap-0'],
                                     expected_source_revision=store.remediation_source_revision(SID))
    assert store.get_hitl_item(w.alt)['status'] == 'pending'
    with store._db.cursor() as cur:
        store._db.execute(cur, "SELECT COUNT(*) AS n FROM jobs WHERE type='apply_approved_values'")
        assert store._db.fetchone(cur)['n'] == 0


def test_the_route_maps_the_refusal_to_409(store, monkeypatch):
    from fastapi import HTTPException
    import core
    from routes import hitl as routes
    w = World(store)
    w.reconcile()
    monkeypatch.setattr(core, 'store', store)
    with pytest.raises(HTTPException) as exc:
        routes.hitl_update(w.alt, routes.HitlUpdate(status='approved'), None)
    assert exc.value.status_code == 409


def test_pending_resync_does_not_resurrect_the_finding(store):
    w = World(store)
    w.reconcile()
    store.sync_hitl_finding_dispositions(w.alt, 'pending')
    store.queue_hitl_review_for_file(SID, FILE, [{'rule_id': '1.1.1', 'finding_count': 1}],
                                     batch_id=w.batch)
    store.record_remediation_diffs(SID, FILE, [{'rule_id': '1.1.1', 'before': '', 'after': 'x'}])
    assert [f['disposition'] for f in w.findings()] == ['superseded_by_reassessment']
    assert w.row() is None


def test_writer_refuses_a_row_retired_while_it_was_writing(store, monkeypatch):
    """The race: the job read the approved 1.1.1 row as writable, wrote it, and a reconciliation
    committed before the job stored its copy. Nothing is uploaded or credited; retryable."""
    w = World(store, alt_status='approved', replaced=())          # picture still present
    import core
    import handlers
    from proposals import Verification
    blob = Blob(w.corrected, w.source)
    monkeypatch.setattr(core, 'store', store)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    monkeypatch.setattr(handlers, '_verify_residual', lambda data, name, scan_id=None: Verification(
        True, set(), assessment={'status': 'analysed', 'issues': [], 'errors': [], 'skipped_rules': 0},
        artifact_sha256=sha(data)))
    calls = []

    def removed(st, sid, f):
        calls.append(1)
        return set() if len(calls) == 1 else {w.alt}
    monkeypatch.setattr(rtr, 'removed_item_ids', removed)
    with pytest.raises(RuntimeError, match='review target removed'):
        handlers._apply_approved_values({'scan_id': SID, 'file': FILE}, {})
    assert blob.uploads == []
    assert not store.get_hitl_item(w.alt).get('applied')


def test_writer_takes_row_locks_before_the_file_lock(store, monkeypatch):
    """SEQUENCE EVIDENCE ONLY: the approved writer's commit requests review-row locks first."""
    w = World(store, alt_status='approved', replaced=())
    import core
    import handlers
    from proposals import Verification
    blob = Blob(w.corrected, w.source)
    monkeypatch.setattr(core, 'store', store)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    monkeypatch.setattr(handlers, '_verify_residual', lambda data, name, scan_id=None: Verification(
        True, set(), assessment={'status': 'analysed', 'issues': [], 'errors': [], 'skipped_rules': 0},
        artifact_sha256=sha(data)))
    order = []
    real_rows, real_file = rtr.lock_rows, rtr.lock_file
    monkeypatch.setattr(rtr, 'lock_rows', lambda st, ids: (order.append(('rows', sorted(ids))), real_rows(st, ids))[1])
    monkeypatch.setattr(rtr, 'lock_file', lambda st, s, f: (order.append(('file', f)), real_file(st, s, f))[1])
    handlers._apply_approved_values({'scan_id': SID, 'file': FILE}, {})
    assert blob.uploads == [FILE]
    assert order[:2] == [('rows', [w.alt]), ('file', FILE)]


# ── lifecycle: evidence lapses with the artifact, and so does the ledger ───────────────────

def test_new_artifact_restoring_the_target_reopens_row_and_ledger_without_a_sync(store):
    w = World(store)
    execution = store.get_stage_execution(w.batch)
    snapshot = lambda: store._stage_domain_reconciliation(execution, {})  # noqa: E731
    w.reconcile()
    assert w.row() is None
    assert snapshot()['buckets']['superseded'] == 1
    assert snapshot()['unresolved_findings'] == []

    # A new corrected artifact (sha b) that carries the picture again — e.g. a rebuilt copy.
    store.record_remediation(SID, FILE, blob_url='http://b/3', corrected_sha256=sha(w.source))

    assert w.row() is not None                                   # ordinary work again
    assert w.row(include_superseded=True).get('superseded') is False
    [finding] = w.findings()
    assert finding['disposition'] == 'awaiting_review'           # the ledger agrees
    assert finding['fix_evidence_ids'] == [f'target_removal_lapsed:{sha(w.source)}']
    view = snapshot()
    assert view['buckets']['superseded'] == 0 and view['buckets']['awaiting_review'] == 1
    assert [f['finding_id'] for f in view['unresolved_findings']] == [finding['finding_id']]
    # No hitl sync/decision was involved.
    assert [d['action'] for d in store.list_decisions(scan_id=SID)
            if d['action'].startswith('hitl.') and d['action'] != rtr.ACTION] == []


def test_same_sha_rewrite_does_not_reopen(store):
    w = World(store)
    w.reconcile()
    store.record_remediation(SID, FILE, blob_url='http://b/1', corrected_sha256=sha(w.corrected))
    assert w.findings()[0]['disposition'] == 'superseded_by_reassessment' and w.row() is None


# ── snapshot C2 ────────────────────────────────────────────────────────────────────────────

def test_snapshot_names_each_unresolved_finding_and_its_row(store):
    w = World(store)
    execution = store.get_stage_execution(w.batch)
    view = store._stage_domain_reconciliation(execution, {})
    [entry] = view['unresolved_findings']
    assert entry == {'finding_id': w.findings()[0]['finding_id'], 'file': FILE, 'rule_id': '1.1.1',
                     'rule_name': 'Non-text Content', 'disposition': 'awaiting_review',
                     'review_item_id': w.alt}
    assert view['unresolved_findings_total'] == 1 and view['unresolved_findings_truncated'] is False
    for key in ('unit', 'total', 'accounted', 'buckets', 'exact'):
        assert key in view


def test_snapshot_truncation_says_so_and_keeps_the_count(store, monkeypatch):
    w = World(store, colors=('red', 'blue'),
              alt_targets=('word/document.xml#Picture 1', 'word/document.xml#Picture 2'))
    monkeypatch.setattr(store, 'UNRESOLVED_FINDINGS_CAP', 1)
    view = store._stage_domain_reconciliation(store.get_stage_execution(w.batch), {})
    assert len(view['unresolved_findings']) == 1
    assert view['unresolved_findings_total'] == 2 and view['unresolved_findings_truncated'] is True


# ── trigger (c): the completed production run, from persisted evidence ─────────────────────

def _persist(store, w, actor=OWNER, **changes):
    record = store.get_file_record(SID, FILE)
    evidence = w.evidence(**changes)
    evidence.pop('assessment')
    evidence.pop('verified_at')
    evidence['remediated_at'] = record['remediated_at']
    store.log_decision(actor, 'release.corrected_copy_assessed', scan_id=SID, file=FILE,
                       detail=json.dumps(evidence))


def _completed(store, w):
    with store._db.cursor() as cur:
        store._db.execute(cur, "UPDATE stage_executions SET state='succeeded' WHERE execution_id=%s",
                          (w.batch,))
        store._db.execute(cur, "UPDATE jobs SET status='done' WHERE scan_id=%s", (SID,))


def test_completed_run_is_reconciled_from_persisted_saved_copy_evidence(store, monkeypatch):
    w = World(store)
    _persist(store, w)
    _completed(store, w)
    blob = Blob(w.corrected, w.source)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    report = {}
    store.reconcile_completed_remediation_reviews(SID, report=report)
    assert [s['item_id'] for s in report[FILE]['superseded']] == [w.alt]
    audit = w.row(include_superseded=True)
    assert audit['superseded_evidence']['assessment'] == 'release.corrected_copy_assessed'
    assert blob.uploads == [] and blob.data == w.corrected       # no document write
    assert store.get_hitl_item(w.alt)['status'] == 'pending'


def test_the_auto_route_reports_what_it_reconciled(store, monkeypatch):
    import core
    from routes import hitl as routes
    w = World(store)
    _persist(store, w)
    _completed(store, w)
    monkeypatch.setitem(sys.modules, 'blob', Blob(w.corrected, w.source))
    monkeypatch.setattr(core, 'store', store)
    monkeypatch.setattr(core, 'fire_webhook', lambda *a, **k: None)

    class Req:
        class state:
            user_email = None
    response = routes.hitl_auto_queue(SID, Req())
    assert response['target_removals']['superseded'] == [w.alt]


@pytest.mark.parametrize('actor,changes', [('system', {}), (OWNER, {'skipped_rules': 2}),
                                           (OWNER, {'artifact_sha256': 'a' * 64})])
def test_persisted_evidence_must_be_the_owners_complete_check_of_this_sha(store, monkeypatch, actor, changes):
    w = World(store)
    _persist(store, w, actor=actor, **changes)
    _completed(store, w)
    monkeypatch.setitem(sys.modules, 'blob', Blob(w.corrected, w.source))
    report = {}
    store.reconcile_completed_remediation_reviews(SID, report=report)
    assert not report[FILE]['superseded'] and w.lines() == []


def test_persisted_path_without_cached_source_changes_nothing(store, monkeypatch):
    w = World(store)
    _persist(store, w)
    _completed(store, w)
    monkeypatch.setitem(sys.modules, 'blob', Blob(w.corrected, None))
    report = {}
    store.reconcile_completed_remediation_reviews(SID, report=report)
    assert report[FILE]['skipped'][0]['reason'] == 'bytes_unavailable' and w.lines() == []


def test_automatic_queue_puts_no_marker_on_a_superseded_row(store):
    from automatic_review_queue import annotate
    row = {'id': 'x', 'scan_id': SID, 'status': 'pending', 'superseded': True, 'proposals': [{}],
           'automatic_approval': {'responsibility': 'human'}}
    [out] = annotate(store, [row], OWNER)
    assert 'automatic_approval' not in out


# ── follow-up 2: repeated polls read no stored document; the narrow reconcile-targets route ──

class CountingBlob(Blob):
    def __init__(self, remediated, source=None):
        super().__init__(remediated, source)
        self.reads = []

    def download_remediated(self, owner, sid, f):
        self.reads.append(('remediated', f))
        return super().download_remediated(owner, sid, f)

    def download_source(self, owner, sid, f, checksum=None):
        self.reads.append(('source', f))
        return super().download_source(owner, sid, f, checksum)


def test_repeated_calls_after_retirement_read_no_stored_document(store, monkeypatch):
    """The retired row keeps its audit status 'pending'. Eligibility alone would re-download
    both documents on every poll; current evidence must answer from metadata."""
    w = World(store)
    _persist(store, w)
    blob = CountingBlob(w.corrected, w.source)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    first = rtr.reconcile_scan(store, SID)
    assert [s['item_id'] for s in first[FILE]['superseded']] == [w.alt]
    reads = len(blob.reads)
    assert reads >= 2                                   # the one real check read both
    for _ in range(3):
        again = rtr.reconcile_scan(store, SID)
        assert again[FILE]['superseded'] == []
        assert [u['item_id'] for u in again[FILE]['unchanged']] == [w.alt]
    assert len(blob.reads) == reads, blob.reads[reads:]
    assert len(w.lines()) == 1


def test_a_mixed_file_is_still_checked_and_then_memoized(store, monkeypatch):
    """One row retired, one genuinely open: the file is processed (the open row is decided on
    the bytes), and the identical next call answers from memory until an input changes."""
    w = World(store)
    w.reconcile()                                           # the 1.1.1 row is retired
    other = store.enqueue_proposals(SID, FILE, '1.4.9', [
        {'locator': 'image 1', 'before': '', 'proposed_value': TEXT, 'source': 'OCR'}],
        rule_name='Images of Text (No Exception)')
    _persist(store, w)
    blob = CountingBlob(w.corrected, w.source)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    first = rtr.reconcile_scan(store, SID)[FILE]
    assert len(blob.reads) >= 2                             # processed, not skipped as retired
    assert any(s['item_id'] == other for s in first['skipped'])
    assert any(u['item_id'] == w.alt for u in first['unchanged'])
    reads = len(blob.reads)
    assert rtr.reconcile_scan(store, SID)[FILE] == first
    assert len(blob.reads) == reads                         # memoized: nothing changed
    with store._db.cursor() as cur:                         # a binding change re-evaluates
        store._db.execute(cur, "UPDATE hitl_queue SET proposal_snapshot_ids=%s WHERE id=%s",
                          (json.dumps(['snap-new']), other))
    rtr.reconcile_scan(store, SID)
    assert len(blob.reads) > reads


def test_whole_scan_calls_are_bounded(store, monkeypatch):
    w = World(store)
    _persist(store, w)
    blob = CountingBlob(w.corrected, w.source)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    result = rtr.reconcile_scan(store, SID, max_files=0)
    assert result[FILE]['skipped'] == [{'item_id': None, 'reason': 'deferred_bounded'}]
    assert blob.reads == [] and w.lines() == []


def _client(store, monkeypatch, owner=OWNER):
    import core
    from app import app
    from fastapi.testclient import TestClient
    monkeypatch.setattr(core, 'store', store)
    monkeypatch.setattr(core, 'ACCESS_CODE', '')
    monkeypatch.setattr(core, 'GOOGLE_CLIENT_ID', 'test-client-id')
    monkeypatch.setattr(core, 'E2E_KEY', None)
    monkeypatch.setattr(core, 'OWNER_EMAIL', owner)
    monkeypatch.setattr(core, 'verify_gis_token', lambda t: t)
    monkeypatch.setattr(core, 'email_allowed', lambda e: True)
    return TestClient(app)


def test_reconcile_targets_route_is_owner_scoped_narrow_and_idempotent(store, monkeypatch):
    w = World(store)
    _persist(store, w)
    blob = CountingBlob(w.corrected, w.source)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    client = _client(store, monkeypatch)
    url = f'/hitl/queue/{SID}/reconcile-targets'
    # No routing, no AI, no writer: any of these being reached is a failure.
    for name in ('queue_hitl_items', 'reconcile_completed_remediation_reviews',
                 'queue_hitl_review_for_file', 'enqueue_job', 'enqueue_proposals'):
        monkeypatch.setattr(store, name, lambda *a, _n=name, **k: pytest.fail(f'{_n} must not run'))
    rows_before = {r['id'] for r in store.list_hitl_queue(scan_id=SID, include_superseded=True)}

    assert client.post(url).status_code == 401
    assert client.post(url, headers={'Authorization': 'Bearer other@example.com'}).status_code == 404
    assert w.lines() == [] and blob.reads == []

    response = client.post(url, headers={'Authorization': f'Bearer {OWNER}'})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body == {'scan_id': SID, 'superseded_count': 1,
                    'files': [{'file': FILE, 'superseded': [w.alt], 'unchanged': [], 'skipped': []}]}
    assert len(w.lines()) == 1
    assert store.get_hitl_item(w.alt)['status'] == 'pending'
    assert blob.uploads == []                                   # no document write

    reads = len(blob.reads)
    second = client.post(url, headers={'Authorization': f'Bearer {OWNER}'}).json()
    assert second['superseded_count'] == 0
    assert second['files'] == [{'file': FILE, 'superseded': [], 'unchanged': [w.alt],
                                'skipped': [{'item_id': None, 'reason': 'nothing_to_reconcile'}]}]
    assert len(w.lines()) == 1 and len(blob.reads) == reads
    assert {r['id'] for r in store.list_hitl_queue(scan_id=SID, include_superseded=True)} == rows_before


def test_reconcile_targets_requires_the_review_capability():
    from workspace_capability_map import required_capabilities
    # Keyed by the route TEMPLATE, as the middleware looks it up.
    assert required_capabilities('POST', '/hitl/queue/{scan_id}/reconcile-targets') == {'remediate.review'}


def test_auto_route_still_routes_and_reconciles(store, monkeypatch):
    import core
    from routes import hitl as routes
    w = World(store)
    _persist(store, w)
    _completed(store, w)
    monkeypatch.setitem(sys.modules, 'blob', Blob(w.corrected, w.source))
    monkeypatch.setattr(core, 'store', store)
    monkeypatch.setattr(core, 'fire_webhook', lambda *a, **k: None)
    called = []
    real = store.queue_hitl_items
    monkeypatch.setattr(store, 'queue_hitl_items', lambda sid: (called.append(sid), real(sid))[1])

    class Req:
        class state:
            user_email = None
    response = routes.hitl_auto_queue(SID, Req())
    assert called == [SID] and response['target_removals']['superseded'] == [w.alt]


# ── follow-up 3: every lapse of the binding reopens the ledger, not only a new artifact ─────

def _consistent_open(store, w):
    """Row visible as work again AND the ledger agrees, read through the real surfaces."""
    assert w.row() is not None
    assert w.row(include_superseded=True).get('superseded') is False
    [finding] = w.findings()
    assert finding['disposition'] == 'awaiting_review'
    view = store._stage_domain_reconciliation(store.get_stage_execution(w.batch), {})
    assert view['buckets']['superseded'] == 0
    assert view['unresolved_findings_total'] == 1
    assert [f['finding_id'] for f in view['unresolved_findings']] == [finding['finding_id']]
    events = store.finding_disposition_events(SID, w.batch, finding['finding_id'])
    retired_at = max(i for i, e in enumerate(events) if e['to_disposition'] == 'superseded_by_reassessment')
    reopen = events[retired_at + 1]                    # append-only, and it says why
    assert reopen['to_disposition'] == 'awaiting_review'
    assert reopen['fix_evidence_ids'][0].startswith('target_removal_lapsed:')
    return finding


def test_a_real_proposal_refresh_reopens_the_retired_finding(store):
    """Same corrected sha, new proposal from a later remediation pass (enqueue_proposals is the
    production refresh): the binding lapses, so the row is work again and the ledger must say so."""
    w = World(store)
    w.reconcile()
    assert w.row() is None and w.findings()[0]['disposition'] == 'superseded_by_reassessment'
    sha_before = store.get_file_record(SID, FILE)['corrected_sha256']

    store.enqueue_proposals(SID, FILE, '1.1.1', [
        {'locator': 'word/document.xml#Picture 1', 'before': '',
         'proposed_value': 'A refreshed synthetic draft', 'source': 'vision'}],
        rule_name='Non-text Content')

    assert store.get_file_record(SID, FILE)['corrected_sha256'] == sha_before
    _consistent_open(store, w)


def test_an_identical_refresh_keeps_the_retirement(store):
    """A refresh that changes nothing the evidence is bound to is not a lapse."""
    w = World(store)
    w.reconcile()
    row = store.get_hitl_item(w.alt)
    store.sync_hitl_finding_dispositions(w.alt, row['status'])
    assert w.row() is None and w.findings()[0]['disposition'] == 'superseded_by_reassessment'


def test_a_real_decision_version_change_reopens_the_retired_finding(store):
    """The one decision still accepted on a retired row — re-saving it as pending with a note —
    bumps decision_version through complete_hitl_decision; the ledger follows the row."""
    w = World(store)
    w.reconcile()
    store.complete_hitl_decision(w.alt, 'pending', 'looked again', None, resolution=None,
                                 approved_values=None, actor=OWNER, detail=None)
    assert store.get_hitl_item(w.alt)['decision_version'] == 1
    _consistent_open(store, w)


def test_the_sync_path_reopens_even_when_its_group_limit_would_skip_the_finding(store):
    """sync projects `finding_count` rows in instance-key order. With a second, already-open
    finding for the row sorting first, the group projection alone would never reach the
    retired one — the explicit lapse check must."""
    w = World(store)
    w.reconcile()
    [retired] = w.findings()
    with store._db.cursor() as cur:
        store._db.execute(cur,
            "INSERT INTO finding_disposition(scan_id,batch_id,finding_id,workflow_id,snapshot_id,"
            "document_id,file,rule_id,instance_key,assessment_status,disposition,review_item_id,"
            "fix_evidence_ids,verified_at,revision,created_at,updated_at) "
            "VALUES(%s,%s,'aaa-open',%s,%s,'doc',%s,'1.1.1','aaa-first','review','awaiting_review',"
            "%s,'[]',NULL,1,'t','t')", (SID, w.batch, SID, SID, FILE, w.alt))
        store._db.execute(cur, "UPDATE hitl_queue SET proposal_snapshot_ids=%s WHERE id=%s",
                          (json.dumps(['snap-lapsed']), w.alt))
    store.sync_hitl_finding_dispositions(w.alt, 'pending')
    by_id = {f['finding_id']: f['disposition'] for f in w.findings()}
    assert by_id[retired['finding_id']] == 'awaiting_review'


def test_reconcile_targets_reopens_a_lapsed_retirement_even_without_new_proof(store, monkeypatch):
    """The copy changed by a path that did not reopen (simulated by a direct pointer update) and
    no saved-copy check exists for the new copy: the call must still reopen the finding under
    its locks, and must not claim anything retired."""
    w = World(store)
    _persist(store, w)
    blob = CountingBlob(w.corrected, w.source)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    rtr.reconcile_scan(store, SID)
    assert w.findings()[0]['disposition'] == 'superseded_by_reassessment'
    with store._db.cursor() as cur:
        store._db.execute(cur, "UPDATE file_records SET corrected_sha256=%s WHERE scan_id=%s AND file=%s",
                          ('b' * 64, SID, FILE))
    assert w.findings()[0]['disposition'] == 'superseded_by_reassessment'      # stale until asked
    order = []
    real_rows, real_file = rtr.lock_rows, rtr.lock_file
    monkeypatch.setattr(rtr, 'lock_rows', lambda st, ids: (order.append('rows'), real_rows(st, ids))[1])
    monkeypatch.setattr(rtr, 'lock_file', lambda st, s, f: (order.append('file'), real_file(st, s, f))[1])
    reads = len(blob.reads)

    body = _client(store, monkeypatch).post(f'/hitl/queue/{SID}/reconcile-targets',
                                            headers={'Authorization': f'Bearer {OWNER}'}).json()

    assert body['superseded_count'] == 0
    assert body['files'][0]['skipped'] == [{'item_id': None, 'reason': 'verification_missing'}]
    assert order[:2] == ['rows', 'file']
    assert len(blob.reads) == reads                                             # no proof, no read
    _consistent_open(store, w)


def test_nothing_lapsed_takes_no_write_lock(store, monkeypatch):
    """Polling a file whose retirement is still current must not take the write locks."""
    w = World(store)
    _persist(store, w)
    monkeypatch.setitem(sys.modules, 'blob', CountingBlob(w.corrected, w.source))
    rtr.reconcile_scan(store, SID)
    taken = []
    monkeypatch.setattr(rtr, 'lock_file', lambda *a: taken.append(a) or None)
    rtr.reconcile_scan(store, SID)
    assert taken == []
